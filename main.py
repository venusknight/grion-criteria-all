# -*- coding: utf-8 -*-
import re
import csv
import shutil
import sys
import sqlite3
import hashlib
import uuid
from pathlib import Path
from datetime import datetime, date

import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog

try:
    from PIL import Image, ImageTk
except ImportError:
    Image = None
    ImageTk = None

APP_NAME = "grION criteria"
APP_VERSION = "1.0"
DB_NAME = "grion_criteria.db"
LEGACY_DB_NAME = "grion_sythesis.db"

if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent

DB_PATH = BASE_DIR / DB_NAME
LEGACY_DB_PATH = BASE_DIR / LEGACY_DB_NAME

if not DB_PATH.exists() and LEGACY_DB_PATH.exists():
    shutil.copy2(LEGACY_DB_PATH, DB_PATH)

ASSETS_DIR = BASE_DIR / "assets"
ITEM_IMAGES_DIR = ASSETS_DIR / "imagens_itens"
PRODUCT_IMAGES_DIR = ASSETS_DIR / "imagens_produtos"
PROFILE_IMAGES_DIR = ASSETS_DIR / "imagens_perfil"
ATTACHMENTS_DIR = ASSETS_DIR / "anexos"
EXPORTS_DIR = BASE_DIR / "exports"
BACKUPS_DIR = BASE_DIR / "backups"
PDFS_DIR = BASE_DIR / "pdfs"
LOGO_PATH = ASSETS_DIR / "logo_empresa.png"
LEGACY_LOGO_PATH = BASE_DIR / "logo_empresa.png"

COLOR_BG = "#1E293B"
COLOR_TOPBAR = "#0F172A"
COLOR_SIDEBAR = "#121A2E"
COLOR_PANEL = "#151E31"
COLOR_PANEL_2 = "#182234"
COLOR_SURFACE = "#1E293B"
COLOR_TEXT = "#E8F1F8"
COLOR_MUTED = "#64748B"
COLOR_ACCENT = "#60A5FA"
COLOR_CYAN = "#22D3EE"
COLOR_BORDER = "#334155"
COLOR_NAV_ACTIVE = "#152954"
COLOR_BUTTON = "#1F2A44"
COLOR_BUTTON_HOVER = "#24365A"
COLOR_PRIMARY = "#2563EB"
COLOR_PRIMARY_HOVER = "#1D4ED8"
COLOR_SUCCESS = "#16A34A"
COLOR_WARNING = "#F59E0B"
COLOR_DANGER = "#EF4444"

def draw_round_rect(canvas, x1, y1, x2, y2, radius, fill, outline="", width=1):
    radius = min(radius, (x2 - x1) / 2, (y2 - y1) / 2)
    points = [
        x1 + radius, y1,
        x2 - radius, y1,
        x2, y1,
        x2, y1 + radius,
        x2, y2 - radius,
        x2, y2,
        x2 - radius, y2,
        x1 + radius, y2,
        x1, y2,
        x1, y2 - radius,
        x1, y1 + radius,
        x1, y1,
    ]
    return canvas.create_polygon(points, smooth=True, splinesteps=18, fill=fill, outline=outline, width=width)

MODULE_PERMISSIONS = [
    ("dashboard", "Dashboard"),
    ("produtos", "Produtos"),
    ("parceiros", "Parceiros"),
    ("compras", "Compras"),
    ("vendas", "Vendas"),
    ("faturamento", "Faturamento"),
    ("movimentacoes", "Movimentações"),
    ("funcionarios", "Funcionários"),
    ("relatorios", "Relatórios"),
    ("cadastros_aux", "Cadastros auxiliares"),
    ("comercial", "Comercial"),
    ("propostas", "Propostas"),
    ("contratos", "Pedidos e contratos"),
    ("projetos", "Projetos e serviços"),
    ("escalas", "Escalas"),
    ("horas", "Apontamento de horas"),
    ("rh", "RH"),
    ("viagens", "Viagens"),
    ("recebimento", "Recebimento"),
    ("equipamentos", "Equipamentos"),
    ("financeiro", "Financeiro"),
    ("fiscal", "Fiscal"),
    ("operacional", "Relatórios operacionais"),
    ("acessos", "Acessos"),
    ("configuracoes", "Configurações"),
    ("sistema", "Sistema"),
    ("perfil", "Perfil"),
]

FINAL_DOCUMENT_STATUSES = ("Recebido", "Recebida", "Faturado", "Faturado internamente", "Cancelado")
PENDING_DOCUMENT_COLUMNS = ["id", "numero", "tipo_documento", "parceiro", "data_emissao", "status", "total"]
LOW_STOCK_COLUMNS = ["codigo", "nome", "estoque_atual", "estoque_minimo", "status"]

ROLE_LABELS = {
    "admin": "Administrador geral",
    "gestor": "Gestor",
    "compras": "Compras",
    "vendas": "Vendas",
    "estoque": "Estoque",
    "financeiro": "Financeiro",
    "colaborador": "Colaborador",
    "client": "Cliente",
    "viewer": "Visualização",
}

ROLE_DEFAULT_MODULES = {
    "admin": {code for code, _ in MODULE_PERMISSIONS},
    "gestor": {"dashboard", "produtos", "parceiros", "compras", "vendas", "faturamento", "movimentacoes", "funcionarios", "relatorios", "cadastros_aux", "comercial", "propostas", "contratos", "projetos", "escalas", "horas", "rh", "viagens", "recebimento", "equipamentos", "financeiro", "fiscal", "operacional", "perfil"},
    "compras": {"dashboard", "produtos", "parceiros", "compras", "recebimento", "financeiro", "relatorios", "perfil"},
    "vendas": {"dashboard", "produtos", "parceiros", "vendas", "comercial", "propostas", "contratos", "faturamento", "relatorios", "perfil"},
    "estoque": {"dashboard", "produtos", "movimentacoes", "recebimento", "equipamentos", "relatorios", "perfil"},
    "financeiro": {"dashboard", "compras", "faturamento", "financeiro", "fiscal", "relatorios", "perfil"},
    "colaborador": {"dashboard", "produtos", "parceiros", "projetos", "horas", "operacional", "perfil"},
    "client": {"dashboard", "vendas", "perfil"},
    "viewer": {"dashboard", "relatorios", "perfil"},
}

AUXILIARY_CATEGORIES = [
    "Contato por empresa",
    "Cargo/especialidade",
    "Valor por hora",
    "Centro de custo",
    "Banco/conta",
    "Forma de pagamento",
    "Condição comercial",
    "Imposto",
    "Planta/local",
    "Documento cadastral",
]

ERP_MODULE_CONFIGS = {
    "comercial": {
        "title": "Comercial e clientes",
        "prefix": "COM",
        "categories": ["Prospecção", "Cadastro como fornecedor", "Contato", "Oportunidade", "Cotação", "Escopo solicitado", "Visita técnica", "Negociação", "Proposta pendente", "Perda/cancelamento"],
        "statuses": ["Aberto", "Em andamento", "Pendente cliente", "Aprovado", "Perdido", "Cancelado"],
    },
    "contratos": {
        "title": "Pedidos e contratos",
        "prefix": "CON",
        "categories": ["OC cliente", "Contrato", "Aditivo", "Prorrogação", "Alerta sem OC", "Alteração contratual"],
        "statuses": ["Em análise", "Ativo", "Saldo aberto", "Concluído", "Vencido", "Cancelado"],
    },
    "projetos": {
        "title": "Projetos e serviços",
        "prefix": "PRJ",
        "categories": ["Projeto", "Ordem de serviço", "Etapa", "Pendência", "Alteração de escopo", "Encerramento técnico", "Encerramento financeiro"],
        "statuses": ["Planejado", "Em andamento", "Pendente", "Atrasado", "Concluído", "Cancelado"],
    },
    "escalas": {
        "title": "Planejamento de equipe e escalas",
        "prefix": "ESC",
        "categories": ["Disponibilidade", "Escala diária", "Escala semanal", "Escala mensal", "Folga", "DSR", "Férias", "Falta", "Atestado", "Substituição"],
        "statuses": ["Planejado", "Confirmado", "Conflito", "Executado", "Cancelado"],
    },
    "horas": {
        "title": "Apontamento de horas",
        "prefix": "HRS",
        "categories": ["Entrada/saída", "Intervalo", "Hora normal", "Extra 50%", "Extra 100%", "Domingo/feriado", "Deslocamento", "Falta", "Atraso", "Atestado", "Aprovação"],
        "statuses": ["Lançado", "Em aprovação", "Aprovado", "Reprovado", "Fechado"],
    },
    "rh": {
        "title": "Funcionários e RH",
        "prefix": "RHU",
        "categories": ["Documento", "Vencimento", "Férias", "Atestado", "Afastamento", "Treinamento", "Benefício", "Adiantamento", "Alteração salarial", "Prestador"],
        "statuses": ["Ativo", "Pendente", "Vencendo", "Vencido", "Concluído", "Cancelado"],
    },
    "viagens": {
        "title": "Viagens e serviços externos",
        "prefix": "VIA",
        "categories": ["Solicitação", "Aprovação", "Hotel", "Passagem", "Veículo", "Alimentação", "Combustível", "Pedágio", "Adiantamento", "Despesa", "Prestação de contas", "Reembolso"],
        "statuses": ["Solicitado", "Aprovado", "Em viagem", "Prestando contas", "Concluído", "Reprovado", "Cancelado"],
    },
    "recebimento": {
        "title": "Recebimento e notas de entrada",
        "prefix": "REC",
        "categories": ["Recebimento", "Nota fiscal", "Conferência", "Divergência", "Devolução", "Recebimento parcial", "Entrada em estoque", "XML/DANFE"],
        "statuses": ["Recebido", "Em conferência", "Divergente", "Bloqueado", "Liberado", "Devolvido", "Cancelado"],
    },
    "equipamentos": {
        "title": "Ferramentas, máquinas e equipamentos",
        "prefix": "EQP",
        "categories": ["Ferramenta", "Máquina", "Equipamento", "Retirada", "Devolução", "Manutenção preventiva", "Manutenção corretiva", "Calibração", "Avaria/perda", "Aluguel"],
        "statuses": ["Disponível", "Em uso", "Em manutenção", "Calibração vencendo", "Vencido", "Perdido", "Baixado"],
    },
    "fiscal": {
        "title": "Fiscal e documentos",
        "prefix": "FIS",
        "categories": ["NF-e entrada", "NF-e saída", "NFS-e", "XML", "DANFE", "Cancelamento", "Carta de correção", "Retenção", "Erro fiscal", "Competência mensal"],
        "statuses": ["Recebido", "Emitido", "Em conferência", "Pendente contabilidade", "Com erro", "Cancelado", "Arquivado"],
    },
    "operacional": {
        "title": "Relatórios operacionais",
        "prefix": "ROP",
        "categories": ["Relatório diário", "Atividade executada", "Funcionários presentes", "Foto/evidência", "Problema", "Material utilizado", "Pendência", "Aprovação cliente", "Relatório final", "Termo de aceite"],
        "statuses": ["Aberto", "Em preenchimento", "Aguardando aprovação", "Aprovado", "Reprovado", "Concluído"],
    },
}

def normalize_digits(value: str) -> str:
    return re.sub(r"\D", "", value or "")

def hash_password(password: str) -> str:
    return hashlib.sha256((password or "").encode("utf-8")).hexdigest()

def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")

def today_iso() -> str:
    return date.today().isoformat()

def today_br() -> str:
    return format_date_br(today_iso())

def format_document(doc: str) -> str:
    digits = normalize_digits(doc)

    if len(digits) == 11:
        return f"{digits[:3]}.{digits[3:6]}.{digits[6:9]}-{digits[9:]}"

    if len(digits) == 14:
        return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"

    return doc or ""

def parse_float_br(value) -> float:
    text = str(value or "").strip()

    if not text:
        return 0.0

    text = text.replace(" ", "")

    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")

    try:
        return float(text)
    except ValueError:
        raise ValueError("Valor inválido.")

def format_date_br(value: str) -> str:
    value = (value or "").strip()

    if not value:
        return ""

    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            parsed = datetime.strptime(value, fmt).date()
            return parsed.strftime("%d/%m/%Y")
        except ValueError:
            pass

    digits = normalize_digits(value)

    if len(digits) == 8:
        try:
            if 1900 <= int(digits[:4]) <= 2099:
                parsed = datetime.strptime(digits, "%Y%m%d").date()
            else:
                parsed = datetime.strptime(digits, "%d%m%Y").date()
            return parsed.strftime("%d/%m/%Y")
        except ValueError:
            return value

    return value

def validate_date(value: str, field_name: str, allow_empty=True) -> str:
    value = (value or "").strip()

    if not value:
        if allow_empty:
            return ""
        raise ValueError(f"Informe {field_name}.")

    formats = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y")

    for fmt in formats:
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass

    digits = normalize_digits(value)

    if len(digits) == 8:
        try:
            if 1900 <= int(digits[:4]) <= 2099:
                return datetime.strptime(digits, "%Y%m%d").date().isoformat()
            return datetime.strptime(digits, "%d%m%Y").date().isoformat()
        except ValueError:
            pass

    raise ValueError(f"{field_name} inválida. Use o formato DD/MM/AAAA.")

def days_until(date_str: str):
    date_str = (date_str or "").strip()

    if not date_str:
        return None

    try:
        return (datetime.strptime(date_str, "%Y-%m-%d").date() - date.today()).days
    except ValueError:
        return None

def ensure_app_dirs():
    for path in (ASSETS_DIR, ITEM_IMAGES_DIR, PRODUCT_IMAGES_DIR, PROFILE_IMAGES_DIR, ATTACHMENTS_DIR, EXPORTS_DIR, BACKUPS_DIR, PDFS_DIR):
        path.mkdir(parents=True, exist_ok=True)

def relative_to_base(path: Path) -> str:
    try:
        return path.resolve().relative_to(BASE_DIR.resolve()).as_posix()
    except ValueError:
        return str(path)

def resolve_local_path(path_value: str):
    path_value = (path_value or "").strip()

    if not path_value:
        return None

    path = Path(path_value)

    if not path.is_absolute():
        path = BASE_DIR / path

    return path

def copy_item_image(source_path: str, item_codigo: str) -> str:
    source = Path(source_path)

    if not source.exists():
        raise ValueError("Imagem selecionada nao foi encontrada.")

    suffix = source.suffix.lower()

    if suffix not in (".png", ".jpg", ".jpeg"):
        raise ValueError("Use uma imagem PNG, JPG ou JPEG.")

    ensure_app_dirs()
    safe_code = re.sub(r"[^A-Za-z0-9_-]", "_", item_codigo or "item")
    destination = ITEM_IMAGES_DIR / f"{safe_code}_{uuid.uuid4().hex[:8]}{suffix}"
    shutil.copy2(source, destination)

    return relative_to_base(destination)

def copy_product_image(source_path: str, product_code: str) -> str:
    source = Path(source_path)

    if not source.exists():
        raise ValueError("Imagem selecionada não foi encontrada.")

    suffix = source.suffix.lower()

    if suffix not in (".png", ".jpg", ".jpeg"):
        raise ValueError("Use uma imagem PNG, JPG ou JPEG.")

    ensure_app_dirs()
    safe_code = re.sub(r"[^A-Za-z0-9_-]", "_", product_code or "produto")
    destination = PRODUCT_IMAGES_DIR / f"{safe_code}_{uuid.uuid4().hex[:8]}{suffix}"
    shutil.copy2(source, destination)

    return relative_to_base(destination)

def copy_profile_image(source_path: str, user_id) -> str:
    source = Path(source_path)

    if not source.exists():
        raise ValueError("Imagem selecionada não foi encontrada.")

    suffix = source.suffix.lower()

    if suffix not in (".png", ".jpg", ".jpeg"):
        raise ValueError("Use uma imagem PNG, JPG ou JPEG.")

    ensure_app_dirs()
    safe_code = re.sub(r"[^A-Za-z0-9_-]", "_", str(user_id or "perfil"))
    destination = PROFILE_IMAGES_DIR / f"{safe_code}_{uuid.uuid4().hex[:8]}{suffix}"
    shutil.copy2(source, destination)

    return relative_to_base(destination)

def copy_attachment(source_path: str, prefix: str = "anexo") -> str:
    source = Path(source_path)

    if not source.exists():
        raise ValueError("Arquivo selecionado não foi encontrado.")

    ensure_app_dirs()
    safe_prefix = re.sub(r"[^A-Za-z0-9_-]", "_", prefix or "anexo")
    destination = ATTACHMENTS_DIR / f"{safe_prefix}_{uuid.uuid4().hex[:8]}{source.suffix.lower()}"
    shutil.copy2(source, destination)
    return relative_to_base(destination)

def load_tk_image(path_value: str, size=(120, 120)):
    path = resolve_local_path(path_value)

    if not path or not path.exists() or Image is None or ImageTk is None:
        return None

    try:
        image = Image.open(path)
        image.thumbnail(size)
        return ImageTk.PhotoImage(image)
    except Exception:
        return None

def get_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def table_columns(conn, table_name: str):
    rows = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
    return [row["name"] for row in rows]

def add_column_if_missing(conn, table_name: str, column_name: str, column_sql: str):
    columns = table_columns(conn, table_name)

    if column_name not in columns:
        conn.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_sql}")

def init_database():
    ensure_app_dirs()
    conn = get_conn()
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            email TEXT UNIQUE,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'client',
            status TEXT NOT NULL DEFAULT 'pending',
            client_code TEXT,
            created_at TEXT NOT NULL
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS clients (
            codigo TEXT PRIMARY KEY,
            tipo_pessoa TEXT NOT NULL,
            nome_razao TEXT NOT NULL,
            documento TEXT UNIQUE NOT NULL,
            telefone TEXT,
            email TEXT,
            responsavel TEXT,
            endereco TEXT,
            user_id INTEGER,
            created_at TEXT NOT NULL,
            FOREIGN KEY (user_id) REFERENCES users(id)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS employees (
            codigo TEXT PRIMARY KEY,
            nome TEXT NOT NULL,
            documento TEXT UNIQUE,
            cargo TEXT,
            telefone TEXT,
            email TEXT,
            status TEXT NOT NULL DEFAULT 'Ativo',
            observacoes TEXT,
            created_at TEXT NOT NULL
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS items (
            codigo TEXT PRIMARY KEY,
            nome TEXT NOT NULL,
            categoria TEXT NOT NULL,
            unidade TEXT NOT NULL,
            ca TEXT,
            valor_unitario REAL DEFAULT 0,
            image_path TEXT,
            observacoes TEXT,
            created_at TEXT NOT NULL
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS movements (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            item_codigo TEXT NOT NULL,
            item_nome TEXT NOT NULL,
            item_categoria TEXT,
            item_unidade TEXT,
            tipo TEXT NOT NULL,
            quantidade INTEGER NOT NULL,
            data_movimentacao TEXT NOT NULL,
            validade TEXT,
            funcionario_codigo TEXT,
            funcionario_nome TEXT,
            service_codigo TEXT,
            destino TEXT,
            observacoes TEXT,
            created_at TEXT NOT NULL
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS services (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo_servico TEXT UNIQUE NOT NULL,
            cliente_codigo TEXT,
            nome_servico TEXT NOT NULL,
            descricao TEXT,
            valor REAL DEFAULT 0,
            prazo TEXT,
            status TEXT NOT NULL,
            observacoes TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY (cliente_codigo) REFERENCES clients(codigo)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS service_employees (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo_servico TEXT NOT NULL,
            funcionario_codigo TEXT NOT NULL,
            created_at TEXT NOT NULL,
            UNIQUE(codigo_servico, funcionario_codigo)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS service_required_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo_servico TEXT NOT NULL,
            item_codigo TEXT NOT NULL,
            created_at TEXT NOT NULL,
            UNIQUE(codigo_servico, item_codigo)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS roles (
            code TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            description TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS permissions (
            code TEXT PRIMARY KEY,
            module TEXT NOT NULL,
            description TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS role_permissions (
            role_code TEXT NOT NULL,
            permission_code TEXT NOT NULL,
            allowed INTEGER NOT NULL DEFAULT 1,
            PRIMARY KEY (role_code, permission_code)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS user_permissions (
            user_id INTEGER NOT NULL,
            permission_code TEXT NOT NULL,
            allowed INTEGER NOT NULL DEFAULT 1,
            PRIMARY KEY (user_id, permission_code),
            FOREIGN KEY (user_id) REFERENCES users(id)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS dashboard_preferences (
            user_id INTEGER,
            card_code TEXT NOT NULL,
            visible INTEGER NOT NULL DEFAULT 1,
            display_order INTEGER NOT NULL DEFAULT 0,
            PRIMARY KEY (user_id, card_code)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS products (
            codigo TEXT PRIMARY KEY,
            nome TEXT NOT NULL,
            descricao TEXT,
            categoria TEXT,
            tipo_produto TEXT,
            unidade TEXT,
            ncm TEXT,
            codigo_barras TEXT,
            valor_custo REAL DEFAULT 0,
            valor_venda REAL DEFAULT 0,
            estoque_atual REAL DEFAULT 0,
            estoque_minimo REAL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'Ativo',
            image_path TEXT,
            observacoes TEXT,
            cfop TEXT,
            unidade_tributavel TEXT,
            origem_produto TEXT,
            cst_csosn TEXT,
            aliquota_icms REAL DEFAULT 0,
            created_at TEXT NOT NULL,
            updated_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS partners (
            codigo TEXT PRIMARY KEY,
            tipo_pessoa TEXT NOT NULL,
            documento TEXT UNIQUE,
            nome_razao TEXT NOT NULL,
            nome_fantasia TEXT,
            inscricao_estadual TEXT,
            inscricao_municipal TEXT,
            telefone TEXT,
            celular TEXT,
            email TEXT,
            endereco TEXT,
            numero TEXT,
            bairro TEXT,
            cidade TEXT,
            estado TEXT,
            cep TEXT,
            pais TEXT,
            responsavel TEXT,
            tipo_parceiro TEXT,
            status TEXT NOT NULL DEFAULT 'Ativo',
            observacoes TEXT,
            regime_tributario TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS document_types (
            code TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            direction TEXT NOT NULL,
            prefix TEXT NOT NULL,
            category TEXT NOT NULL
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            numero TEXT UNIQUE NOT NULL,
            tipo_code TEXT NOT NULL,
            direcao TEXT NOT NULL,
            partner_code TEXT,
            data_emissao TEXT NOT NULL,
            data_validade TEXT,
            prazo_entrega TEXT,
            data_prevista_entrega TEXT,
            data_faturamento TEXT,
            vendedor_responsavel TEXT,
            condicao_pagamento TEXT,
            forma_pagamento TEXT,
            status TEXT NOT NULL,
            subtotal REAL DEFAULT 0,
            desconto_total REAL DEFAULT 0,
            acrescimo_total REAL DEFAULT 0,
            total REAL DEFAULT 0,
            observacoes TEXT,
            fiscal_status TEXT DEFAULT 'Não fiscal',
            fiscal_natureza_operacao TEXT,
            fiscal_finalidade TEXT,
            fiscal_chave_acesso TEXT,
            fiscal_protocolo TEXT,
            fiscal_xml_path TEXT,
            fiscal_danfe_path TEXT,
            fiscal_autorizado_em TEXT,
            fiscal_retorno TEXT,
            origem_documento_id INTEGER,
            created_at TEXT NOT NULL,
            updated_at TEXT,
            FOREIGN KEY (partner_code) REFERENCES partners(codigo),
            FOREIGN KEY (tipo_code) REFERENCES document_types(code)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS document_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER NOT NULL,
            product_code TEXT,
            produto_codigo TEXT,
            produto_nome TEXT NOT NULL,
            descricao TEXT,
            quantidade REAL NOT NULL,
            unidade TEXT,
            valor_unitario REAL DEFAULT 0,
            desconto REAL DEFAULT 0,
            acrescimo REAL DEFAULT 0,
            total REAL DEFAULT 0,
            ncm TEXT,
            tipo_produto TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY (document_id) REFERENCES documents(id) ON DELETE CASCADE,
            FOREIGN KEY (product_code) REFERENCES products(codigo)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS business_movements (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_code TEXT,
            document_id INTEGER,
            numero_documento TEXT,
            tipo TEXT NOT NULL,
            quantidade REAL NOT NULL,
            data_movimentacao TEXT NOT NULL,
            parceiro_codigo TEXT,
            parceiro_nome TEXT,
            observacoes TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY (product_code) REFERENCES products(codigo),
            FOREIGN KEY (document_id) REFERENCES documents(id)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS accounts_payable (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document_id INTEGER UNIQUE,
            numero_documento TEXT,
            partner_code TEXT,
            descricao TEXT,
            valor REAL DEFAULT 0,
            data_emissao TEXT,
            data_vencimento TEXT,
            status TEXT NOT NULL DEFAULT 'Em aberto',
            observacoes TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT,
            FOREIGN KEY (document_id) REFERENCES documents(id) ON DELETE CASCADE,
            FOREIGN KEY (partner_code) REFERENCES partners(codigo)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS auxiliary_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo TEXT UNIQUE NOT NULL,
            categoria TEXT NOT NULL,
            nome TEXT NOT NULL,
            partner_code TEXT,
            cidade TEXT,
            valor REAL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'Ativo',
            descricao TEXT,
            attachment_path TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT,
            FOREIGN KEY (partner_code) REFERENCES partners(codigo)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS erp_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            module_code TEXT NOT NULL,
            codigo TEXT UNIQUE NOT NULL,
            categoria TEXT,
            titulo TEXT NOT NULL,
            partner_code TEXT,
            product_code TEXT,
            employee_code TEXT,
            project_code TEXT,
            cost_center TEXT,
            data_inicio TEXT,
            data_fim TEXT,
            status TEXT,
            valor_previsto REAL DEFAULT 0,
            valor_real REAL DEFAULT 0,
            prioridade TEXT,
            responsavel TEXT,
            descricao TEXT,
            observacoes TEXT,
            attachment_path TEXT,
            created_by INTEGER,
            created_at TEXT NOT NULL,
            updated_at TEXT,
            FOREIGN KEY (partner_code) REFERENCES partners(codigo),
            FOREIGN KEY (product_code) REFERENCES products(codigo),
            FOREIGN KEY (created_by) REFERENCES users(id)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS commercial_proposals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo TEXT UNIQUE NOT NULL,
            partner_code TEXT,
            projeto TEXT,
            revisao INTEGER DEFAULT 0,
            data_emissao TEXT,
            validade TEXT,
            prazo TEXT,
            pagamento TEXT,
            condicoes TEXT,
            horas_normais REAL DEFAULT 0,
            horas_50 REAL DEFAULT 0,
            horas_100 REAL DEFAULT 0,
            valor_hora REAL DEFAULT 0,
            custos_viagem REAL DEFAULT 0,
            custos_materiais REAL DEFAULT 0,
            custos_terceiros REAL DEFAULT 0,
            impostos_percentual REAL DEFAULT 0,
            margem_percentual REAL DEFAULT 0,
            custo_total REAL DEFAULT 0,
            preco_venda REAL DEFAULT 0,
            lucro_previsto REAL DEFAULT 0,
            incluso TEXT,
            nao_incluso TEXT,
            status TEXT NOT NULL DEFAULT 'Em elaboração',
            motivo_status TEXT,
            attachment_path TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT,
            FOREIGN KEY (partner_code) REFERENCES partners(codigo)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS financial_entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo TEXT UNIQUE NOT NULL,
            tipo TEXT NOT NULL,
            descricao TEXT NOT NULL,
            partner_code TEXT,
            document_id INTEGER,
            project_code TEXT,
            cost_center TEXT,
            valor REAL DEFAULT 0,
            data_emissao TEXT,
            data_vencimento TEXT,
            data_baixa TEXT,
            status TEXT NOT NULL DEFAULT 'Aberto',
            forma_pagamento TEXT,
            conta_bancaria TEXT,
            observacoes TEXT,
            attachment_path TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT,
            FOREIGN KEY (partner_code) REFERENCES partners(codigo),
            FOREIGN KEY (document_id) REFERENCES documents(id)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS audit_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            action TEXT NOT NULL,
            entity TEXT NOT NULL,
            entity_id TEXT,
            details TEXT,
            created_at TEXT NOT NULL,
            FOREIGN KEY (user_id) REFERENCES users(id)
        )
    """)

    add_column_if_missing(conn, "items", "valor_unitario", "valor_unitario REAL DEFAULT 0")
    add_column_if_missing(conn, "items", "image_path", "image_path TEXT")
    add_column_if_missing(conn, "clients", "contato", "contato TEXT")
    add_column_if_missing(conn, "clients", "documentos", "documentos TEXT")
    add_column_if_missing(conn, "clients", "observacoes", "observacoes TEXT")
    add_column_if_missing(conn, "movements", "funcionario_codigo", "funcionario_codigo TEXT")
    add_column_if_missing(conn, "movements", "funcionario_nome", "funcionario_nome TEXT")
    add_column_if_missing(conn, "movements", "service_codigo", "service_codigo TEXT")
    add_column_if_missing(conn, "users", "display_name", "display_name TEXT")
    add_column_if_missing(conn, "users", "phone", "phone TEXT")
    add_column_if_missing(conn, "users", "document", "document TEXT")
    add_column_if_missing(conn, "users", "profile_photo", "profile_photo TEXT")
    add_column_if_missing(conn, "users", "updated_at", "updated_at TEXT")

    for role_code, role_name in ROLE_LABELS.items():
        cur.execute("""
            INSERT OR IGNORE INTO roles (code, name, description)
            VALUES (?, ?, ?)
        """, (
            role_code,
            role_name,
            "Perfil padrao do sistema"
        ))

    for permission_code, module_name in MODULE_PERMISSIONS:
        cur.execute("""
            INSERT OR IGNORE INTO permissions (code, module, description)
            VALUES (?, ?, ?)
        """, (
            permission_code,
            module_name,
            f"Acesso ao modulo {module_name}"
        ))

    for role_code, module_codes in ROLE_DEFAULT_MODULES.items():
        for permission_code, _module_name in MODULE_PERMISSIONS:
            cur.execute("""
                INSERT OR IGNORE INTO role_permissions (role_code, permission_code, allowed)
                VALUES (?, ?, ?)
            """, (
                role_code,
                permission_code,
                1 if permission_code in module_codes else 0
            ))

    document_types = [
        ("OCM", "Orçamento de compra", "compra", "OCM", "orcamento_compra"),
        ("OC", "Ordem de compra", "compra", "OC", "ordem_compra"),
        ("OV", "Orçamento de venda", "venda", "OV", "orcamento_venda"),
        ("PV", "Pedido de venda", "venda", "PV", "pedido_venda"),
        ("FT", "Faturamento", "venda", "FT", "faturamento"),
    ]

    for code, name, direction, prefix, category in document_types:
        cur.execute("""
            INSERT OR IGNORE INTO document_types (code, name, direction, prefix, category)
            VALUES (?, ?, ?, ?, ?)
        """, (code, name, direction, prefix, category))

    admin = cur.execute(
        "SELECT id FROM users WHERE username = ?",
        ("admin",)
    ).fetchone()

    if not admin:
        cur.execute("""
            INSERT INTO users (
                username,
                email,
                password_hash,
                role,
                status,
                client_code,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (
            "admin",
            "",
            hash_password("admin123"),
            "admin",
            "approved",
            None,
            now_iso()
        ))
    else:
        cur.execute("""
            UPDATE users
            SET password_hash = ?,
                role = ?,
                status = ?
            WHERE username = ?
        """, (
            hash_password("admin123"),
            "admin",
            "approved",
            "admin"
        ))

    conn.commit()
    conn.close()

def authenticate_user(login: str, password: str):
    conn = get_conn()
    cur = conn.cursor()

    user = cur.execute("""
        SELECT *
        FROM users
        WHERE username = ? OR email = ?
        LIMIT 1
    """, (login, login)).fetchone()

    conn.close()

    if not user:
        return None

    if user["status"] == "pending":
        return "pending"

    if user["status"] == "blocked":
        return "blocked"

    if user["password_hash"] == hash_password(password):
        return dict(user)

    return None

def verify_user_password(user_id, password: str) -> bool:
    conn = get_conn()
    row = conn.execute(
        "SELECT password_hash FROM users WHERE id = ?",
        (user_id,)
    ).fetchone()
    conn.close()

    return bool(row and row["password_hash"] == hash_password(password))

def get_user_by_id(user_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    conn.close()
    return dict(row) if row else None

def update_user_profile(user_id, username, email, display_name, phone, document, photo_source=None, current_photo=""):
    username = (username or "").strip()
    email = (email or "").strip()
    display_name = (display_name or "").strip()
    phone = (phone or "").strip()
    document = (document or "").strip()

    if not user_id:
        raise ValueError("Usuário não identificado.")

    if not username:
        raise ValueError("Informe o login.")

    conn = get_conn()
    cur = conn.cursor()

    if cur.execute("SELECT id FROM users WHERE username = ? AND id <> ?", (username, user_id)).fetchone():
        conn.close()
        raise ValueError("Este login já está em uso.")

    if email and cur.execute("SELECT id FROM users WHERE email = ? AND id <> ?", (email, user_id)).fetchone():
        conn.close()
        raise ValueError("Este e-mail já está em uso.")

    photo_path = current_photo or ""

    if photo_source:
        photo_path = copy_profile_image(photo_source, user_id)

    cur.execute("""
        UPDATE users
        SET username = ?,
            email = ?,
            display_name = ?,
            phone = ?,
            document = ?,
            profile_photo = ?,
            updated_at = ?
        WHERE id = ?
    """, (
        username,
        email,
        display_name,
        phone,
        document,
        photo_path,
        now_iso(),
        user_id
    ))

    conn.commit()
    conn.close()
    return get_user_by_id(user_id)

def clear_operational_data():
    tables = [
        "audit_logs",
        "financial_entries",
        "commercial_proposals",
        "erp_records",
        "auxiliary_records",
        "accounts_payable",
        "document_items",
        "business_movements",
        "documents",
        "products",
        "partners",
        "service_required_items",
        "service_employees",
        "movements",
        "services",
        "items",
        "employees",
        "clients",
    ]
    reset_sequences = [
        "audit_logs",
        "financial_entries",
        "commercial_proposals",
        "erp_records",
        "auxiliary_records",
        "accounts_payable",
        "document_items",
        "business_movements",
        "documents",
        "service_required_items",
        "service_employees",
        "movements",
        "services",
    ]

    conn = get_conn()

    try:
        for table in tables:
            conn.execute(f"DELETE FROM {table}")

        conn.execute("UPDATE users SET client_code = NULL WHERE username <> 'admin'")
        conn.executemany(
            "DELETE FROM sqlite_sequence WHERE name = ?",
            [(table,) for table in reset_sequences]
        )
        conn.commit()
    finally:
        conn.close()

def save_database_backup(backup_path):
    ensure_app_dirs()
    backup_path = Path(backup_path)
    shutil.copy2(DB_PATH, backup_path)
    return backup_path

def reset_database_to_factory():
    if DB_PATH.exists():
        DB_PATH.unlink()

    init_database()

def get_next_code(table_name: str, column_name: str, prefix: str) -> str:
    conn = get_conn()
    cur = conn.cursor()

    row = cur.execute(f"SELECT COUNT(*) AS qtd FROM {table_name}").fetchone()
    number = int(row["qtd"] or 0) + 1

    while True:
        code = f"{prefix}-{number:05d}"
        exists = cur.execute(
            f"SELECT {column_name} FROM {table_name} WHERE {column_name} = ?",
            (code,)
        ).fetchone()

        if not exists:
            conn.close()
            return code

        number += 1

def create_client_user(
    nome_razao: str,
    documento: str,
    telefone: str,
    email: str,
    username: str,
    password: str,
    tipo_pessoa: str = "Jurídica",
    responsavel: str = "",
    endereco: str = "",
    status: str = "pending"
):
    nome_razao = (nome_razao or "").strip()
    documento = normalize_digits(documento)
    telefone = (telefone or "").strip()
    email = (email or "").strip()
    username = (username or "").strip()
    password = password or ""

    if not nome_razao:
        raise ValueError("Informe o nome ou razão social.")

    if len(documento) not in (11, 14):
        raise ValueError("Informe um CPF ou CNPJ válido.")

    if not username:
        raise ValueError("Informe o login.")

    if len(password) < 4:
        raise ValueError("A senha deve ter pelo menos 4 caracteres.")

    conn = get_conn()
    cur = conn.cursor()

    if cur.execute("SELECT id FROM users WHERE username = ?", (username,)).fetchone():
        conn.close()
        raise ValueError("Este login já está em uso.")

    if email and cur.execute("SELECT id FROM users WHERE email = ?", (email,)).fetchone():
        conn.close()
        raise ValueError("Este e-mail já está em uso.")

    if cur.execute("SELECT codigo FROM clients WHERE documento = ?", (documento,)).fetchone():
        conn.close()
        raise ValueError("Este CPF/CNPJ já está cadastrado.")

    now = now_iso()

    cur.execute("""
        INSERT INTO users (
            username,
            email,
            password_hash,
            role,
            status,
            client_code,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (
        username,
        email,
        hash_password(password),
        "client",
        status,
        None,
        now
    ))

    user_id = cur.lastrowid
    codigo_cliente = f"CLI-{user_id:05d}"

    cur.execute("""
        INSERT INTO clients (
            codigo,
            tipo_pessoa,
            nome_razao,
            documento,
            telefone,
            email,
            responsavel,
            endereco,
            user_id,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        codigo_cliente,
        tipo_pessoa,
        nome_razao,
        documento,
        telefone,
        email,
        responsavel,
        endereco,
        user_id,
        now
    ))

    cur.execute(
        "UPDATE users SET client_code = ? WHERE id = ?",
        (codigo_cliente, user_id)
    )

    conn.commit()
    conn.close()

    return codigo_cliente

def get_pending_users():
    conn = get_conn()

    rows = conn.execute("""
        SELECT
            u.id,
            u.username,
            u.email,
            u.status,
            u.created_at,
            c.codigo,
            c.nome_razao,
            c.documento,
            c.telefone
        FROM users u
        LEFT JOIN clients c ON c.user_id = u.id
        WHERE u.status = 'pending'
        ORDER BY u.created_at DESC
    """).fetchall()

    conn.close()
    return [dict(row) for row in rows]

def update_user_status(user_id, status: str):
    conn = get_conn()

    conn.execute(
        "UPDATE users SET status = ? WHERE id = ?",
        (status, user_id)
    )

    conn.commit()
    conn.close()

def role_label(role: str) -> str:
    return ROLE_LABELS.get(role, ROLE_LABELS["client"])

def status_label(status: str) -> str:
    if status == "approved":
        return "Aprovado"

    if status == "blocked":
        return "Bloqueado"

    return "Pendente"

def role_value(label: str) -> str:
    legacy = {
        "Administrador": "admin",
        "Cliente": "client",
    }

    if label in legacy:
        return legacy[label]

    for role_code, role_name in ROLE_LABELS.items():
        if label == role_name:
            return role_code

    return "client"

def status_value(label: str) -> str:
    if label == "Aprovado":
        return "approved"

    if label == "Bloqueado":
        return "blocked"

    return "pending"

def get_users_access():
    conn = get_conn()

    rows = conn.execute("""
        SELECT
            u.id,
            u.username,
            u.email,
            u.role,
            u.status,
            COALESCE(u.client_code, c.codigo, '') AS client_code,
            c.nome_razao,
            c.documento,
            u.created_at
        FROM users u
        LEFT JOIN clients c ON c.user_id = u.id
        ORDER BY u.username
    """).fetchall()

    conn.close()

    users = []

    for row in rows:
        item = dict(row)
        item["perfil"] = role_label(item.get("role", "client"))
        item["status_acesso"] = status_label(item.get("status", "pending"))
        item["documento"] = format_document(item.get("documento", ""))
        users.append(item)

    return users

def update_user_access(user_id, role: str, status: str):
    if role not in ROLE_LABELS:
        raise ValueError("Perfil de acesso inválido.")

    if status not in ("approved", "pending", "blocked"):
        raise ValueError("Status de acesso inválido.")

    conn = get_conn()

    conn.execute("""
        UPDATE users
        SET role = ?,
            status = ?
        WHERE id = ?
    """, (
        role,
        status,
        user_id
    ))

    conn.commit()
    conn.close()

def user_can_access_module(user: dict, module_code: str) -> bool:
    role = (user or {}).get("role", "client")

    if role == "admin":
        return True

    user_id = (user or {}).get("id")

    if user_id:
        conn = get_conn()
        row = conn.execute("""
            SELECT allowed
            FROM user_permissions
            WHERE user_id = ?
              AND permission_code = ?
        """, (user_id, module_code)).fetchone()
        conn.close()

        if row is not None:
            return bool(row["allowed"])

    return module_code in ROLE_DEFAULT_MODULES.get(role, ROLE_DEFAULT_MODULES["client"])

def get_clients():
    conn = get_conn()

    rows = conn.execute("""
        SELECT
            c.codigo,
            c.tipo_pessoa,
            c.nome_razao,
            c.documento,
            c.telefone,
            c.email,
            c.responsavel,
            c.endereco,
            c.contato,
            c.documentos,
            c.observacoes,
            u.username,
            u.status
        FROM clients c
        LEFT JOIN users u ON u.id = c.user_id
        ORDER BY c.nome_razao
    """).fetchall()

    conn.close()

    clients = []
    for row in rows:
        item = dict(row)
        item["documento"] = format_document(item.get("documento", ""))
        clients.append(item)

    return clients

def get_client_by_code(codigo):
    conn = get_conn()

    row = conn.execute("""
        SELECT
            c.*,
            u.username,
            u.status
        FROM clients c
        LEFT JOIN users u ON u.id = c.user_id
        WHERE c.codigo = ?
    """, (codigo,)).fetchone()

    conn.close()

    if not row:
        return None

    client = dict(row)
    client["documento"] = format_document(client.get("documento", ""))
    return client

def get_companies():
    companies = []

    for client in get_clients():
        kind = str(client.get("tipo_pessoa", "")).lower()
        document = normalize_digits(client.get("documento", ""))

        if "jur" in kind or len(document) == 14:
            companies.append(client)

    return companies

def create_employee(nome, documento, cargo, telefone, email, status, observacoes):
    nome = (nome or "").strip()
    documento = normalize_digits(documento)

    if not nome:
        raise ValueError("Informe o nome do funcionário.")

    conn = get_conn()
    cur = conn.cursor()

    if documento:
        exists = cur.execute(
            "SELECT codigo FROM employees WHERE documento = ?",
            (documento,)
        ).fetchone()

        if exists:
            conn.close()
            raise ValueError("Já existe funcionário com este documento.")

    conn.close()

    codigo = get_next_code("employees", "codigo", "FUN")

    conn = get_conn()
    conn.execute("""
        INSERT INTO employees (
            codigo,
            nome,
            documento,
            cargo,
            telefone,
            email,
            status,
            observacoes,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        codigo,
        nome,
        documento,
        (cargo or "").strip(),
        (telefone or "").strip(),
        (email or "").strip(),
        status or "Ativo",
        (observacoes or "").strip(),
        now_iso()
    ))

    conn.commit()
    conn.close()

    return codigo

def get_employees(only_active=False):
    conn = get_conn()

    if only_active:
        rows = conn.execute("""
            SELECT *
            FROM employees
            WHERE status = 'Ativo'
            ORDER BY nome
        """).fetchall()
    else:
        rows = conn.execute("""
            SELECT *
            FROM employees
            ORDER BY nome
        """).fetchall()

    conn.close()

    employees = []

    for row in rows:
        item = dict(row)
        item["documento"] = format_document(item.get("documento", ""))
        employees.append(item)

    return employees

def get_employee_by_code(codigo):
    conn = get_conn()

    row = conn.execute(
        "SELECT * FROM employees WHERE codigo = ?",
        (codigo,)
    ).fetchone()

    conn.close()

    return dict(row) if row else None

def get_employee_options():
    employees = get_employees(only_active=True)
    return [f"{employee['codigo']} - {employee['nome']}" for employee in employees] or [""]

def get_items():
    conn = get_conn()

    rows = conn.execute("""
        SELECT
            i.*,
            COALESCE(balance.saldo, 0) AS saldo,
            COALESCE((
                SELECT m.validade
                FROM movements m
                WHERE m.item_codigo = i.codigo
                  AND TRIM(COALESCE(m.validade, '')) <> ''
                ORDER BY m.data_movimentacao DESC, m.id DESC
                LIMIT 1
            ), '') AS validade
        FROM items i
        LEFT JOIN (
            SELECT item_codigo, SUM(quantidade) AS saldo
            FROM movements
            GROUP BY item_codigo
        ) balance ON balance.item_codigo = i.codigo
        ORDER BY i.nome
    """).fetchall()

    conn.close()

    items = []
    for row in rows:
        item = dict(row)
        item["saldo"] = int(item.get("saldo") or 0)
        item["valor_unitario"] = f"{float(item.get('valor_unitario') or 0):.2f}"
        item["imagem"] = "Sim" if item.get("image_path") else "Nao"
        items.append(item)

    return items

def get_item_by_code(codigo: str):
    conn = get_conn()

    row = conn.execute("""
        SELECT
            i.*,
            COALESCE((
                SELECT SUM(m.quantidade)
                FROM movements m
                WHERE m.item_codigo = i.codigo
            ), 0) AS saldo,
            COALESCE((
                SELECT m.validade
                FROM movements m
                WHERE m.item_codigo = i.codigo
                  AND TRIM(COALESCE(m.validade, '')) <> ''
                ORDER BY m.data_movimentacao DESC, m.id DESC
                LIMIT 1
            ), '') AS validade
        FROM items i
        WHERE i.codigo = ?
    """, (codigo,)).fetchone()

    conn.close()

    if not row:
        return None

    item = dict(row)
    item["saldo"] = int(item.get("saldo") or 0)
    item["valor_unitario"] = float(item.get("valor_unitario") or 0)
    item["imagem"] = "Sim" if item.get("image_path") else "Nao"

    return item

def get_item_balance(codigo: str) -> int:
    conn = get_conn()

    row = conn.execute("""
        SELECT COALESCE(SUM(quantidade), 0) AS saldo
        FROM movements
        WHERE item_codigo = ?
    """, (codigo,)).fetchone()

    conn.close()

    return int(row["saldo"] or 0)

def get_last_validity(codigo: str) -> str:
    conn = get_conn()

    row = conn.execute("""
        SELECT validade
        FROM movements
        WHERE item_codigo = ?
          AND TRIM(COALESCE(validade, '')) <> ''
        ORDER BY data_movimentacao DESC, id DESC
        LIMIT 1
    """, (codigo,)).fetchone()

    conn.close()

    return row["validade"] if row else ""

def get_employee_item_balance(funcionario_codigo: str, item_codigo: str) -> int:
    conn = get_conn()

    row = conn.execute("""
        SELECT COALESCE(SUM(-quantidade), 0) AS saldo_funcionario
        FROM movements
        WHERE funcionario_codigo = ?
          AND item_codigo = ?
    """, (funcionario_codigo, item_codigo)).fetchone()

    conn.close()

    saldo = int(row["saldo_funcionario"] or 0)
    return max(saldo, 0)

def get_employee_items():
    conn = get_conn()

    rows = conn.execute("""
        SELECT
            m.funcionario_codigo,
            m.funcionario_nome,
            m.item_codigo,
            m.item_nome,
            m.item_categoria,
            m.item_unidade,
            COALESCE(SUM(-m.quantidade), 0) AS quantidade,
            MAX(m.data_movimentacao) AS ultima_movimentacao,
            MAX(m.validade) AS validade
        FROM movements m
        WHERE TRIM(COALESCE(m.funcionario_codigo, '')) <> ''
        GROUP BY
            m.funcionario_codigo,
            m.funcionario_nome,
            m.item_codigo,
            m.item_nome,
            m.item_categoria,
            m.item_unidade
        HAVING quantidade > 0
        ORDER BY m.funcionario_nome, m.item_nome
    """).fetchall()

    conn.close()

    return [dict(row) for row in rows]

def get_movements():
    conn = get_conn()

    rows = conn.execute("""
        SELECT *
        FROM movements
        ORDER BY data_movimentacao DESC, id DESC
    """).fetchall()

    conn.close()
    return [dict(row) for row in rows]

def get_item_movements(codigo: str):
    conn = get_conn()

    rows = conn.execute("""
        SELECT *
        FROM movements
        WHERE item_codigo = ?
        ORDER BY data_movimentacao DESC, id DESC
    """, (codigo,)).fetchall()

    conn.close()
    return [dict(row) for row in rows]

def get_item_employee_holders(codigo: str):
    rows = []

    for item in get_employee_items():
        if item.get("item_codigo") == codigo:
            rows.append(item)

    return rows

def get_item_services(codigo: str):
    conn = get_conn()

    rows = conn.execute("""
        SELECT
            s.codigo_servico,
            s.nome_servico,
            c.nome_razao,
            s.status,
            s.prazo
        FROM service_required_items sri
        INNER JOIN services s ON s.codigo_servico = sri.codigo_servico
        LEFT JOIN clients c ON c.codigo = s.cliente_codigo
        WHERE sri.item_codigo = ?
        ORDER BY s.created_at DESC
    """, (codigo,)).fetchall()

    conn.close()
    return [dict(row) for row in rows]

def get_item_epi_expenses(codigo: str):
    conn = get_conn()

    rows = conn.execute("""
        SELECT
            m.funcionario_codigo,
            m.funcionario_nome,
            ABS(m.quantidade) AS quantidade,
            COALESCE(i.valor_unitario, 0) AS valor_unitario,
            ABS(m.quantidade) * COALESCE(i.valor_unitario, 0) AS valor_total,
            m.data_movimentacao,
            m.service_codigo
        FROM movements m
        LEFT JOIN items i ON i.codigo = m.item_codigo
        WHERE m.item_codigo = ?
          AND m.quantidade < 0
        ORDER BY m.data_movimentacao DESC, m.id DESC
    """, (codigo,)).fetchall()

    conn.close()

    expenses = []

    for row in rows:
        item = dict(row)
        item["valor_unitario"] = f"{float(item.get('valor_unitario') or 0):.2f}"
        item["valor_total"] = f"{float(item.get('valor_total') or 0):.2f}"
        expenses.append(item)

    return expenses

def get_services(client_code=None):
    conn = get_conn()

    if client_code:
        rows = conn.execute("""
            SELECT
                s.codigo_servico,
                c.nome_razao,
                s.nome_servico,
                s.descricao,
                s.valor,
                s.prazo,
                s.status,
                s.observacoes
            FROM services s
            LEFT JOIN clients c ON c.codigo = s.cliente_codigo
            WHERE s.cliente_codigo = ?
            ORDER BY s.created_at DESC
        """, (client_code,)).fetchall()
    else:
        rows = conn.execute("""
            SELECT
                s.codigo_servico,
                c.nome_razao,
                s.nome_servico,
                s.descricao,
                s.valor,
                s.prazo,
                s.status,
                s.observacoes
            FROM services s
            LEFT JOIN clients c ON c.codigo = s.cliente_codigo
            ORDER BY s.created_at DESC
        """).fetchall()

    conn.close()

    services = []
    for row in rows:
        service = dict(row)
        service["valor"] = f"{float(service.get('valor') or 0):.2f}"
        services.append(service)

    return services

def get_service_by_code(codigo_servico):
    conn = get_conn()

    row = conn.execute("""
        SELECT
            s.*,
            c.nome_razao
        FROM services s
        LEFT JOIN clients c ON c.codigo = s.cliente_codigo
        WHERE s.codigo_servico = ?
    """, (codigo_servico,)).fetchone()

    conn.close()

    if not row:
        return None

    service = dict(row)
    service["valor"] = f"{float(service.get('valor') or 0):.2f}"
    return service

def get_service_employees(codigo_servico):
    conn = get_conn()

    rows = conn.execute("""
        SELECT e.codigo, e.nome
        FROM service_employees se
        INNER JOIN employees e ON e.codigo = se.funcionario_codigo
        WHERE se.codigo_servico = ?
        ORDER BY e.nome
    """, (codigo_servico,)).fetchall()

    conn.close()

    return [dict(row) for row in rows]

def get_service_required_items(codigo_servico):
    conn = get_conn()

    rows = conn.execute("""
        SELECT i.codigo, i.nome
        FROM service_required_items sri
        INNER JOIN items i ON i.codigo = sri.item_codigo
        WHERE sri.codigo_servico = ?
        ORDER BY i.nome
    """, (codigo_servico,)).fetchall()

    conn.close()

    return [dict(row) for row in rows]

def get_service_options():
    conn = get_conn()

    rows = conn.execute("""
        SELECT codigo_servico, nome_servico
        FROM services
        ORDER BY created_at DESC
    """).fetchall()

    conn.close()

    values = [""]
    values.extend([f"{row['codigo_servico']} - {row['nome_servico']}" for row in rows])

    return values

def get_client_options():
    clients = get_clients()
    return [f"{client['codigo']} - {client['nome_razao']}" for client in clients] or [""]

def get_item_options():
    items = get_items()
    return [f"{item['codigo']} - {item['nome']}" for item in items] or [""]

def get_epi_options():
    items = get_items()
    epis = [item for item in items if str(item.get("categoria", "")).upper() == "EPI"]
    return [f"{item['codigo']} - {item['nome']}" for item in epis] or [""]

def link_employee_to_service(codigo_servico, funcionario_codigo):
    if not codigo_servico or not funcionario_codigo:
        raise ValueError("Selecione o serviço e o funcionário.")

    conn = get_conn()

    conn.execute("""
        INSERT OR IGNORE INTO service_employees (
            codigo_servico,
            funcionario_codigo,
            created_at
        )
        VALUES (?, ?, ?)
    """, (
        codigo_servico,
        funcionario_codigo,
        now_iso()
    ))

    conn.commit()
    conn.close()

def link_required_item_to_service(codigo_servico, item_codigo):
    if not codigo_servico or not item_codigo:
        raise ValueError("Selecione o serviço e o EPI obrigatório.")

    conn = get_conn()

    conn.execute("""
        INSERT OR IGNORE INTO service_required_items (
            codigo_servico,
            item_codigo,
            created_at
        )
        VALUES (?, ?, ?)
    """, (
        codigo_servico,
        item_codigo,
        now_iso()
    ))

    conn.commit()
    conn.close()

def get_epi_expense_report():
    conn = get_conn()

    rows = conn.execute("""
        SELECT
            m.funcionario_codigo,
            m.funcionario_nome,
            m.item_codigo,
            m.item_nome,
            ABS(m.quantidade) AS quantidade,
            COALESCE(i.valor_unitario, 0) AS valor_unitario,
            ABS(m.quantidade) * COALESCE(i.valor_unitario, 0) AS valor_total,
            m.data_movimentacao,
            m.service_codigo,
            m.destino,
            m.observacoes
        FROM movements m
        LEFT JOIN items i ON i.codigo = m.item_codigo
        WHERE m.quantidade < 0
          AND TRIM(COALESCE(m.funcionario_codigo, '')) <> ''
          AND UPPER(COALESCE(m.item_categoria, '')) = 'EPI'
        ORDER BY m.funcionario_nome, m.data_movimentacao DESC
    """).fetchall()

    conn.close()

    report = []

    for row in rows:
        item = dict(row)
        item["valor_unitario"] = f"{float(item.get('valor_unitario') or 0):.2f}"
        item["valor_total"] = f"{float(item.get('valor_total') or 0):.2f}"
        report.append(item)

    return report

def get_service_epi_alerts():
    conn = get_conn()

    services = conn.execute("""
        SELECT *
        FROM services
        WHERE TRIM(COALESCE(prazo, '')) <> ''
          AND status NOT IN ('Concluído', 'Cancelado')
        ORDER BY prazo
    """).fetchall()

    conn.close()

    alerts = []

    for service in services:
        delta = days_until(service["prazo"])

        if delta is None:
            continue

        if delta > 7:
            continue

        employees = get_service_employees(service["codigo_servico"])
        required_items = get_service_required_items(service["codigo_servico"])

        if not employees or not required_items:
            continue

        for employee in employees:
            for item in required_items:
                saldo = get_employee_item_balance(employee["codigo"], item["codigo"])

                if saldo <= 0:
                    if delta < 0:
                        status_data = f"Atrasado há {abs(delta)} dia(s)"
                    elif delta == 0:
                        status_data = "Serviço hoje"
                    else:
                        status_data = f"Faltam {delta} dia(s)"

                    alerts.append({
                        "servico": service["codigo_servico"],
                        "nome_servico": service["nome_servico"],
                        "data": service["prazo"],
                        "funcionario_codigo": employee["codigo"],
                        "funcionario": employee["nome"],
                        "item_codigo": item["codigo"],
                        "epi_faltando": item["nome"],
                        "status": status_data
                    })

    return alerts

def get_validity_status(validade: str):
    delta = days_until(validade)

    if delta is None:
        return None, None

    if delta < 0:
        return f"Vencido há {abs(delta)} dia(s)", "VENCIDO"

    if delta == 0:
        return "Vence hoje", "URGENTE"

    if delta == 1:
        return "Vence amanhã", "URGENTE"

    if delta <= 7:
        return f"Vence em {delta} dia(s)", "URGENTE"

    if delta <= 30:
        return f"Vence em {delta} dia(s)", "ATENÇÃO"

    return None, None

def get_validity_alerts():
    alerts = []

    for item in get_employee_items():
        validade = item.get("validade", "")
        status, nivel = get_validity_status(validade)

        if not status:
            continue

        alerts.append({
            "origem": "Funcionário",
            "responsavel": item.get("funcionario_nome", ""),
            "item_codigo": item.get("item_codigo", ""),
            "item_nome": item.get("item_nome", ""),
            "quantidade": item.get("quantidade", ""),
            "validade": validade,
            "status": status,
            "nivel": nivel
        })

    for item in get_items():
        saldo = int(item.get("saldo") or 0)
        validade = item.get("validade", "")

        if saldo <= 0:
            continue

        status, nivel = get_validity_status(validade)

        if not status:
            continue

        alerts.append({
            "origem": "Estoque",
            "responsavel": "Almoxarifado",
            "item_codigo": item.get("codigo", ""),
            "item_nome": item.get("nome", ""),
            "quantidade": saldo,
            "validade": validade,
            "status": status,
            "nivel": nivel
        })

    def order_key(row):
        validade = row.get("validade", "")

        try:
            return datetime.strptime(validade, "%Y-%m-%d").date()
        except ValueError:
            return date.max

    return sorted(alerts, key=order_key)

def get_critical_stock_report():
    rows = []

    for item in get_items():
        try:
            saldo = int(item.get("saldo") or 0)
        except ValueError:
            saldo = 0

        if saldo <= 0:
            row = dict(item)
            row["status_estoque"] = "Critico" if saldo < 0 else "Zerado"
            rows.append(row)

    return rows

def format_money(value) -> str:
    return f"{float(value or 0):.2f}"

def get_next_document_number(tipo_code: str) -> str:
    conn = get_conn()
    row = conn.execute(
        "SELECT prefix FROM document_types WHERE code = ?",
        (tipo_code,)
    ).fetchone()
    prefix = row["prefix"] if row else tipo_code
    number = 1

    while True:
        document_number = f"{prefix}-{number:06d}"
        exists = conn.execute(
            "SELECT id FROM documents WHERE numero = ?",
            (document_number,)
        ).fetchone()

        if not exists:
            conn.close()
            return document_number

        number += 1

def get_document_types(direction=None):
    conn = get_conn()

    if direction:
        rows = conn.execute("""
            SELECT *
            FROM document_types
            WHERE direction = ?
            ORDER BY name
        """, (direction,)).fetchall()
    else:
        rows = conn.execute("""
            SELECT *
            FROM document_types
            ORDER BY direction, name
        """).fetchall()

    conn.close()
    return [dict(row) for row in rows]

def get_document_type_options(direction=None):
    rows = get_document_types(direction)
    return [f"{row['code']} - {row['name']}" for row in rows] or [""]

def create_product(nome, descricao, categoria, tipo_produto, unidade, ncm, codigo_barras, valor_custo, valor_venda, estoque_atual, estoque_minimo, status, image_source, observacoes):
    nome = (nome or "").strip()

    if not nome:
        raise ValueError("Informe o nome do produto.")

    codigo = get_next_code("products", "codigo", "PROD")
    image_path = copy_product_image(image_source, codigo) if image_source else ""

    conn = get_conn()
    conn.execute("""
        INSERT INTO products (
            codigo, nome, descricao, categoria, tipo_produto, unidade, ncm,
            codigo_barras, valor_custo, valor_venda, estoque_atual, estoque_minimo,
            status, image_path, observacoes, created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        codigo,
        nome,
        (descricao or "").strip(),
        (categoria or "").strip(),
        (tipo_produto or "").strip(),
        (unidade or "").strip(),
        (ncm or "").strip(),
        (codigo_barras or "").strip(),
        parse_float_br(valor_custo),
        parse_float_br(valor_venda),
        parse_float_br(estoque_atual),
        parse_float_br(estoque_minimo),
        status or "Ativo",
        image_path,
        (observacoes or "").strip(),
        now_iso(),
        now_iso()
    ))
    conn.commit()
    conn.close()
    return codigo

def update_product(codigo, nome, descricao, categoria, tipo_produto, unidade, ncm, codigo_barras, valor_custo, valor_venda, estoque_atual, estoque_minimo, status, image_source, observacoes):
    codigo = (codigo or "").strip()
    nome = (nome or "").strip()

    if not codigo:
        raise ValueError("Selecione um produto para editar.")

    if not nome:
        raise ValueError("Informe o nome do produto.")

    current = get_product_by_code(codigo)

    if not current:
        raise ValueError("Produto não encontrado.")

    image_path = copy_product_image(image_source, codigo) if image_source else current.get("image_path", "")
    conn = get_conn()
    conn.execute("""
        UPDATE products
        SET nome = ?,
            descricao = ?,
            categoria = ?,
            tipo_produto = ?,
            unidade = ?,
            ncm = ?,
            codigo_barras = ?,
            valor_custo = ?,
            valor_venda = ?,
            estoque_atual = ?,
            estoque_minimo = ?,
            status = ?,
            image_path = ?,
            observacoes = ?,
            updated_at = ?
        WHERE codigo = ?
    """, (
        nome,
        (descricao or "").strip(),
        (categoria or "").strip(),
        (tipo_produto or "").strip(),
        (unidade or "").strip(),
        (ncm or "").strip(),
        (codigo_barras or "").strip(),
        parse_float_br(valor_custo),
        parse_float_br(valor_venda),
        parse_float_br(estoque_atual),
        parse_float_br(estoque_minimo),
        status or "Ativo",
        image_path,
        (observacoes or "").strip(),
        now_iso(),
        codigo
    ))
    conn.commit()
    conn.close()
    return codigo

def get_products():
    conn = get_conn()
    rows = conn.execute("""
        SELECT *
        FROM products
        ORDER BY nome
    """).fetchall()
    conn.close()

    products = []
    for row in rows:
        product = dict(row)
        product["valor_custo"] = format_money(product.get("valor_custo"))
        product["valor_venda"] = format_money(product.get("valor_venda"))
        product["estoque_atual"] = format_money(product.get("estoque_atual"))
        product["estoque_minimo"] = format_money(product.get("estoque_minimo"))
        product["imagem"] = "Sim" if product.get("image_path") else "Não"
        products.append(product)

    return products

def get_product_by_code(codigo):
    conn = get_conn()
    row = conn.execute("SELECT * FROM products WHERE codigo = ?", (codigo,)).fetchone()
    conn.close()

    if not row:
        return None

    return dict(row)

def get_product_options():
    products = get_products()
    return [f"{product['codigo']} - {product['nome']}" for product in products] or [""]

def create_partner(tipo_pessoa, documento, nome_razao, nome_fantasia, inscricao_estadual, inscricao_municipal, telefone, celular, email, endereco, numero, bairro, cidade, estado, cep, pais, responsavel, tipo_parceiro, status, observacoes):
    nome_razao = (nome_razao or "").strip()
    documento = normalize_digits(documento)

    if not nome_razao:
        raise ValueError("Informe o nome ou razão social.")

    if documento and len(documento) not in (11, 14):
        raise ValueError("Informe um CPF ou CNPJ válido.")

    codigo = get_next_code("partners", "codigo", "PAR")
    conn = get_conn()

    if documento and conn.execute("SELECT codigo FROM partners WHERE documento = ?", (documento,)).fetchone():
        conn.close()
        raise ValueError("Já existe parceiro com este CPF/CNPJ.")

    conn.execute("""
        INSERT INTO partners (
            codigo, tipo_pessoa, documento, nome_razao, nome_fantasia,
            inscricao_estadual, inscricao_municipal, telefone, celular, email,
            endereco, numero, bairro, cidade, estado, cep, pais, responsavel,
            tipo_parceiro, status, observacoes, created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        codigo,
        tipo_pessoa or "Jurídica",
        documento,
        nome_razao,
        (nome_fantasia or "").strip(),
        (inscricao_estadual or "").strip(),
        (inscricao_municipal or "").strip(),
        (telefone or "").strip(),
        (celular or "").strip(),
        (email or "").strip(),
        (endereco or "").strip(),
        (numero or "").strip(),
        (bairro or "").strip(),
        (cidade or "").strip(),
        (estado or "").strip().upper(),
        normalize_digits(cep),
        (pais or "Brasil").strip(),
        (responsavel or "").strip(),
        tipo_parceiro or "Cliente",
        status or "Ativo",
        (observacoes or "").strip(),
        now_iso(),
        now_iso()
    ))
    conn.commit()
    conn.close()
    return codigo

def update_partner(codigo, tipo_pessoa, documento, nome_razao, nome_fantasia, inscricao_estadual, inscricao_municipal, telefone, celular, email, endereco, numero, bairro, cidade, estado, cep, pais, responsavel, tipo_parceiro, status, observacoes):
    codigo = (codigo or "").strip()
    nome_razao = (nome_razao or "").strip()
    documento = normalize_digits(documento)

    if not codigo:
        raise ValueError("Selecione um parceiro para editar.")

    if not nome_razao:
        raise ValueError("Informe o nome ou razão social.")

    if documento and len(documento) not in (11, 14):
        raise ValueError("Informe um CPF ou CNPJ válido.")

    conn = get_conn()

    if not conn.execute("SELECT codigo FROM partners WHERE codigo = ?", (codigo,)).fetchone():
        conn.close()
        raise ValueError("Parceiro não encontrado.")

    exists = conn.execute(
        "SELECT codigo FROM partners WHERE documento = ? AND codigo <> ?",
        (documento, codigo)
    ).fetchone() if documento else None

    if exists:
        conn.close()
        raise ValueError("Já existe outro parceiro com este CPF/CNPJ.")

    conn.execute("""
        UPDATE partners
        SET tipo_pessoa = ?,
            documento = ?,
            nome_razao = ?,
            nome_fantasia = ?,
            inscricao_estadual = ?,
            inscricao_municipal = ?,
            telefone = ?,
            celular = ?,
            email = ?,
            endereco = ?,
            numero = ?,
            bairro = ?,
            cidade = ?,
            estado = ?,
            cep = ?,
            pais = ?,
            responsavel = ?,
            tipo_parceiro = ?,
            status = ?,
            observacoes = ?,
            updated_at = ?
        WHERE codigo = ?
    """, (
        tipo_pessoa or "Jurídica",
        documento,
        nome_razao,
        (nome_fantasia or "").strip(),
        (inscricao_estadual or "").strip(),
        (inscricao_municipal or "").strip(),
        (telefone or "").strip(),
        (celular or "").strip(),
        (email or "").strip(),
        (endereco or "").strip(),
        (numero or "").strip(),
        (bairro or "").strip(),
        (cidade or "").strip(),
        (estado or "").strip().upper(),
        normalize_digits(cep),
        (pais or "Brasil").strip(),
        (responsavel or "").strip(),
        tipo_parceiro or "Cliente",
        status or "Ativo",
        (observacoes or "").strip(),
        now_iso(),
        codigo
    ))
    conn.commit()
    conn.close()
    return codigo

def get_partners(kind=None):
    conn = get_conn()
    rows = conn.execute("""
        SELECT *
        FROM partners
        ORDER BY nome_razao
    """).fetchall()
    conn.close()

    partners = []
    for row in rows:
        partner = dict(row)
        partner["documento"] = format_document(partner.get("documento", ""))

        if kind and kind.lower() not in str(partner.get("tipo_parceiro", "")).lower():
            continue

        partners.append(partner)

    return partners

def get_partner_by_code(codigo):
    conn = get_conn()
    row = conn.execute("SELECT * FROM partners WHERE codigo = ?", (codigo,)).fetchone()
    conn.close()

    if not row:
        return None

    partner = dict(row)
    partner["documento"] = format_document(partner.get("documento", ""))
    return partner

def get_partner_options(kind=None):
    partners = get_partners(kind)
    return [f"{partner['nome_razao']} ({partner['codigo']})" for partner in partners] or [""]

def option_code(text: str) -> str:
    text = (text or "").strip()
    match = re.search(r"\(([^()]+)\)\s*$", text)

    if match:
        return match.group(1).strip()

    return text.split(" - ")[0] if text else ""

def create_document(tipo_code, partner_code, data_emissao, data_validade, prazo_entrega, data_prevista_entrega, vendedor_responsavel, condicao_pagamento, forma_pagamento, status, observacoes, product_code, quantidade, valor_unitario, desconto, acrescimo):
    if not tipo_code:
        raise ValueError("Selecione o tipo do documento.")

    if not partner_code:
        raise ValueError("Selecione o parceiro.")

    product = get_product_by_code(product_code)

    if not product:
        raise ValueError("Selecione um produto cadastrado.")

    type_row = next((row for row in get_document_types() if row["code"] == tipo_code), None)

    if not type_row:
        raise ValueError("Tipo de documento inválido.")

    quantidade_valor = parse_float_br(quantidade)
    valor = parse_float_br(valor_unitario)
    desconto_valor = parse_float_br(desconto)
    acrescimo_valor = parse_float_br(acrescimo)

    if quantidade_valor <= 0:
        raise ValueError("Quantidade deve ser maior que zero.")

    subtotal = quantidade_valor * valor
    total = max(subtotal - desconto_valor + acrescimo_valor, 0)
    numero = get_next_document_number(tipo_code)
    emissao = validate_date(data_emissao or today_iso(), "data de emissão", allow_empty=False)
    validade = validate_date(data_validade, "data de validade", allow_empty=True)
    prazo = validate_date(prazo_entrega, "prazo de entrega", allow_empty=True)
    entrega = validate_date(data_prevista_entrega, "data prevista de entrega", allow_empty=True)

    conn = get_conn()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO documents (
            numero, tipo_code, direcao, partner_code, data_emissao, data_validade,
            prazo_entrega, data_prevista_entrega, vendedor_responsavel,
            condicao_pagamento, forma_pagamento, status, subtotal,
            desconto_total, acrescimo_total, total, observacoes,
            fiscal_status, created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        numero,
        tipo_code,
        type_row["direction"],
        partner_code,
        emissao,
        validade,
        prazo,
        entrega,
        (vendedor_responsavel or "").strip(),
        (condicao_pagamento or "").strip(),
        (forma_pagamento or "").strip(),
        status or "Rascunho",
        subtotal,
        desconto_valor,
        acrescimo_valor,
        total,
        (observacoes or "").strip(),
        "Não fiscal",
        now_iso(),
        now_iso()
    ))
    document_id = cur.lastrowid
    cur.execute("""
        INSERT INTO document_items (
            document_id, product_code, produto_codigo, produto_nome, descricao,
            quantidade, unidade, valor_unitario, desconto, acrescimo, total,
            ncm, tipo_produto, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        document_id,
        product["codigo"],
        product["codigo"],
        product["nome"],
        product.get("descricao", ""),
        quantidade_valor,
        product.get("unidade", ""),
        valor,
        desconto_valor,
        acrescimo_valor,
        total,
        product.get("ncm", ""),
        product.get("tipo_produto", ""),
        now_iso()
    ))
    conn.commit()
    conn.close()
    return document_id, numero

def update_document(document_id, tipo_code, partner_code, data_emissao, data_validade, prazo_entrega, data_prevista_entrega, vendedor_responsavel, condicao_pagamento, forma_pagamento, status, observacoes, product_code, quantidade, valor_unitario, desconto, acrescimo):
    if not document_id:
        raise ValueError("Selecione um documento para editar.")

    if not tipo_code:
        raise ValueError("Selecione o tipo do documento.")

    if not partner_code:
        raise ValueError("Selecione o parceiro.")

    product = get_product_by_code(product_code)

    if not product:
        raise ValueError("Selecione um produto cadastrado.")

    type_row = next((row for row in get_document_types() if row["code"] == tipo_code), None)

    if not type_row:
        raise ValueError("Tipo de documento inválido.")

    quantidade_valor = parse_float_br(quantidade)
    valor = parse_float_br(valor_unitario)
    desconto_valor = parse_float_br(desconto)
    acrescimo_valor = parse_float_br(acrescimo)

    if quantidade_valor <= 0:
        raise ValueError("Quantidade deve ser maior que zero.")

    subtotal = quantidade_valor * valor
    total = max(subtotal - desconto_valor + acrescimo_valor, 0)
    emissao = validate_date(data_emissao or today_iso(), "data de emissão", allow_empty=False)
    validade = validate_date(data_validade, "data de validade", allow_empty=True)
    prazo = validate_date(prazo_entrega, "prazo de entrega", allow_empty=True)
    entrega = validate_date(data_prevista_entrega, "data prevista de entrega", allow_empty=True)

    conn = get_conn()

    if not conn.execute("SELECT id FROM documents WHERE id = ?", (document_id,)).fetchone():
        conn.close()
        raise ValueError("Documento não encontrado.")

    conn.execute("""
        UPDATE documents
        SET tipo_code = ?,
            direcao = ?,
            partner_code = ?,
            data_emissao = ?,
            data_validade = ?,
            prazo_entrega = ?,
            data_prevista_entrega = ?,
            vendedor_responsavel = ?,
            condicao_pagamento = ?,
            forma_pagamento = ?,
            status = ?,
            subtotal = ?,
            desconto_total = ?,
            acrescimo_total = ?,
            total = ?,
            observacoes = ?,
            updated_at = ?
        WHERE id = ?
    """, (
        tipo_code,
        type_row["direction"],
        partner_code,
        emissao,
        validade,
        prazo,
        entrega,
        (vendedor_responsavel or "").strip(),
        (condicao_pagamento or "").strip(),
        (forma_pagamento or "").strip(),
        status or "Rascunho",
        subtotal,
        desconto_valor,
        acrescimo_valor,
        total,
        (observacoes or "").strip(),
        now_iso(),
        document_id
    ))
    conn.execute("DELETE FROM document_items WHERE document_id = ?", (document_id,))
    conn.execute("""
        INSERT INTO document_items (
            document_id, product_code, produto_codigo, produto_nome, descricao,
            quantidade, unidade, valor_unitario, desconto, acrescimo, total,
            ncm, tipo_produto, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        document_id,
        product["codigo"],
        product["codigo"],
        product["nome"],
        product.get("descricao", ""),
        quantidade_valor,
        product.get("unidade", ""),
        valor,
        desconto_valor,
        acrescimo_valor,
        total,
        product.get("ncm", ""),
        product.get("tipo_produto", ""),
        now_iso()
    ))
    conn.commit()
    conn.close()
    return document_id

def get_documents(direction=None, tipo_codes=None):
    conn = get_conn()
    params = []
    where = []

    if direction:
        where.append("d.direcao = ?")
        params.append(direction)

    if tipo_codes:
        placeholders = ",".join("?" for _ in tipo_codes)
        where.append(f"d.tipo_code IN ({placeholders})")
        params.extend(tipo_codes)

    where_sql = f"WHERE {' AND '.join(where)}" if where else ""
    rows = conn.execute(f"""
        SELECT
            d.id,
            d.numero,
            dt.name AS tipo_documento,
            d.tipo_code,
            d.direcao,
            d.data_emissao,
            d.data_validade,
            d.prazo_entrega,
            d.data_prevista_entrega,
            d.data_faturamento,
            d.status,
            d.subtotal,
            d.desconto_total,
            d.acrescimo_total,
            d.total,
            d.fiscal_status,
            p.codigo AS partner_code,
            p.nome_razao AS parceiro,
            d.observacoes
        FROM documents d
        LEFT JOIN document_types dt ON dt.code = d.tipo_code
        LEFT JOIN partners p ON p.codigo = d.partner_code
        {where_sql}
        ORDER BY d.created_at DESC, d.id DESC
    """, params).fetchall()
    conn.close()

    documents = []
    for row in rows:
        document = dict(row)
        for date_field in ("data_emissao", "data_validade", "prazo_entrega", "data_prevista_entrega", "data_faturamento"):
            document[f"_{date_field}_iso"] = document.get(date_field, "")
            document[date_field] = format_date_br(document.get(date_field, ""))
        document["subtotal"] = format_money(document.get("subtotal"))
        document["desconto_total"] = format_money(document.get("desconto_total"))
        document["acrescimo_total"] = format_money(document.get("acrescimo_total"))
        document["total"] = format_money(document.get("total"))
        documents.append(document)

    return documents

def get_business_pending_documents_report():
    return [
        document for document in get_documents()
        if document.get("status") not in FINAL_DOCUMENT_STATUSES
    ]

def get_business_low_stock_report():
    rows = []

    for product in get_products():
        current_stock = parse_float_br(product.get("estoque_atual", 0))
        minimum_stock = parse_float_br(product.get("estoque_minimo", 0))

        if minimum_stock > 0 and current_stock <= minimum_stock:
            rows.append(product)

    return rows

def get_document_by_id(document_id):
    conn = get_conn()
    document = conn.execute("""
        SELECT
            d.*,
            dt.name AS tipo_documento,
            p.nome_razao AS parceiro,
            p.documento AS parceiro_documento,
            p.email AS parceiro_email,
            p.telefone AS parceiro_telefone,
            p.endereco AS parceiro_endereco,
            p.cidade AS parceiro_cidade,
            p.estado AS parceiro_estado
        FROM documents d
        LEFT JOIN document_types dt ON dt.code = d.tipo_code
        LEFT JOIN partners p ON p.codigo = d.partner_code
        WHERE d.id = ?
    """, (document_id,)).fetchone()
    items = conn.execute("""
        SELECT *
        FROM document_items
        WHERE document_id = ?
        ORDER BY id
    """, (document_id,)).fetchall()
    conn.close()

    if not document:
        return None, []

    return dict(document), [dict(row) for row in items]

def get_document_items(document_id):
    _document, items = get_document_by_id(document_id)
    return items

def upsert_account_payable_for_document(document_id, data_vencimento="", status="Em aberto"):
    document, _items = get_document_by_id(document_id)

    if not document or document.get("direcao") != "compra":
        return None

    vencimento = validate_date(
        data_vencimento or document.get("data_validade") or document.get("data_prevista_entrega") or document.get("data_emissao"),
        "vencimento da conta",
        allow_empty=False
    )
    status = status or "Em aberto"

    if document.get("status") == "Cancelado":
        status = "Cancelado"

    conn = get_conn()
    existing = conn.execute(
        "SELECT id FROM accounts_payable WHERE document_id = ?",
        (document_id,)
    ).fetchone()

    values = (
        document_id,
        document.get("numero", ""),
        document.get("partner_code", ""),
        f"{document.get('tipo_documento', 'Compra')} {document.get('numero', '')}",
        parse_float_br(document.get("total", 0)),
        document.get("data_emissao", ""),
        vencimento,
        status,
        document.get("observacoes", ""),
        now_iso(),
    )

    if existing:
        conn.execute("""
            UPDATE accounts_payable
            SET document_id = ?,
                numero_documento = ?,
                partner_code = ?,
                descricao = ?,
                valor = ?,
                data_emissao = ?,
                data_vencimento = ?,
                status = ?,
                observacoes = ?,
                updated_at = ?
            WHERE id = ?
        """, values + (existing["id"],))
        payable_id = existing["id"]
    else:
        conn.execute("""
            INSERT INTO accounts_payable (
                document_id, numero_documento, partner_code, descricao, valor,
                data_emissao, data_vencimento, status, observacoes, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, values + (now_iso(),))
        payable_id = conn.execute("SELECT last_insert_rowid() AS id").fetchone()["id"]

    conn.commit()
    conn.close()
    return payable_id

def get_account_payable_by_document_id(document_id):
    conn = get_conn()
    row = conn.execute(
        "SELECT * FROM accounts_payable WHERE document_id = ?",
        (document_id,)
    ).fetchone()
    conn.close()
    return dict(row) if row else None

def get_accounts_payable():
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            ap.*,
            p.nome_razao AS parceiro
        FROM accounts_payable ap
        LEFT JOIN partners p ON p.codigo = ap.partner_code
        ORDER BY ap.data_vencimento ASC, ap.id DESC
    """).fetchall()
    conn.close()

    result = []
    today = date.today().isoformat()

    for row in rows:
        item = dict(row)
        if item.get("status") == "Em aberto" and item.get("data_vencimento") and item["data_vencimento"] < today:
            item["status"] = "Vencido"
        item["_data_emissao_iso"] = item.get("data_emissao", "")
        item["_data_vencimento_iso"] = item.get("data_vencimento", "")
        item["data_emissao"] = format_date_br(item.get("data_emissao", ""))
        item["data_vencimento"] = format_date_br(item.get("data_vencimento", ""))
        item["valor"] = format_money(item.get("valor"))
        result.append(item)

    return result

def log_audit(user_id, action, entity, entity_id="", details=""):
    conn = get_conn()
    conn.execute("""
        INSERT INTO audit_logs (user_id, action, entity, entity_id, details, created_at)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (user_id, action, entity, str(entity_id or ""), details or "", now_iso()))
    conn.commit()
    conn.close()

def get_scoped_code(table_name, prefix, where_sql="", params=()):
    conn = get_conn()
    sql = f"SELECT COUNT(*) AS qtd FROM {table_name}"

    if where_sql:
        sql += f" WHERE {where_sql}"

    count = conn.execute(sql, params).fetchone()["qtd"]
    number = int(count or 0) + 1

    while True:
        code = f"{prefix}-{number:05d}"
        exists = conn.execute(f"SELECT id FROM {table_name} WHERE codigo = ?", (code,)).fetchone()

        if not exists:
            conn.close()
            return code

        number += 1

def create_auxiliary_record(categoria, nome, partner_code, cidade, valor, status, descricao, attachment_source=""):
    nome = (nome or "").strip()

    if not nome:
        raise ValueError("Informe o nome do cadastro.")

    codigo = get_scoped_code("auxiliary_records", "AUX")
    attachment_path = copy_attachment(attachment_source, codigo) if attachment_source else ""
    conn = get_conn()
    conn.execute("""
        INSERT INTO auxiliary_records (
            codigo, categoria, nome, partner_code, cidade, valor, status,
            descricao, attachment_path, created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        codigo,
        categoria or "Geral",
        nome,
        partner_code or None,
        (cidade or "").strip(),
        parse_float_br(valor),
        status or "Ativo",
        (descricao or "").strip(),
        attachment_path,
        now_iso(),
        now_iso()
    ))
    conn.commit()
    conn.close()
    return codigo

def update_auxiliary_record(record_id, categoria, nome, partner_code, cidade, valor, status, descricao, attachment_source="", current_attachment=""):
    if not record_id:
        raise ValueError("Selecione um cadastro para editar.")

    nome = (nome or "").strip()

    if not nome:
        raise ValueError("Informe o nome do cadastro.")

    attachment_path = current_attachment or ""

    if attachment_source:
        attachment_path = copy_attachment(attachment_source, f"AUX-{record_id}")

    conn = get_conn()
    conn.execute("""
        UPDATE auxiliary_records
        SET categoria = ?,
            nome = ?,
            partner_code = ?,
            cidade = ?,
            valor = ?,
            status = ?,
            descricao = ?,
            attachment_path = ?,
            updated_at = ?
        WHERE id = ?
    """, (
        categoria or "Geral",
        nome,
        partner_code or None,
        (cidade or "").strip(),
        parse_float_br(valor),
        status or "Ativo",
        (descricao or "").strip(),
        attachment_path,
        now_iso(),
        record_id
    ))
    conn.commit()
    conn.close()
    return record_id

def get_auxiliary_records():
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            ar.*,
            p.nome_razao AS parceiro
        FROM auxiliary_records ar
        LEFT JOIN partners p ON p.codigo = ar.partner_code
        ORDER BY ar.categoria, ar.nome
    """).fetchall()
    conn.close()

    result = []

    for row in rows:
        item = dict(row)
        item["valor"] = format_money(item.get("valor"))
        item["anexo"] = "Sim" if item.get("attachment_path") else ""
        result.append(item)

    return result

def get_auxiliary_record_by_id(record_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM auxiliary_records WHERE id = ?", (record_id,)).fetchone()
    conn.close()
    return dict(row) if row else None

def get_auxiliary_options(category=None):
    records = get_auxiliary_records()

    if category:
        records = [record for record in records if record.get("categoria") == category and record.get("status") != "Inativo"]

    return [record.get("nome", "") for record in records if record.get("nome")] or [""]

def resolve_partner_text(text, kind=None):
    text = (text or "").strip()

    if not text:
        return None

    partner = get_partner_by_code(option_code(text))

    if partner:
        return partner

    text_lower = text.lower()

    for partner in get_partners(kind):
        fields = [
            partner.get("nome_razao", ""),
            partner.get("nome_fantasia", ""),
            partner.get("documento", ""),
            partner.get("codigo", ""),
        ]

        if any(text_lower == str(field).lower() or text_lower in str(field).lower() for field in fields if field):
            return partner

    return None

def format_partner_option(partner):
    if not partner:
        return ""

    return f"{partner.get('nome_razao') or partner.get('codigo')} ({partner.get('codigo')})"

def create_erp_record(module_code, categoria, titulo, partner_code, product_code, employee_code, project_code, cost_center, data_inicio, data_fim, status, valor_previsto, valor_real, prioridade, responsavel, descricao, observacoes, attachment_source="", created_by=None):
    config = ERP_MODULE_CONFIGS.get(module_code, {})
    titulo = (titulo or "").strip()

    if not titulo:
        raise ValueError("Informe o título do registro.")

    codigo = get_scoped_code("erp_records", config.get("prefix", "ERP"), "module_code = ?", (module_code,))
    attachment_path = copy_attachment(attachment_source, codigo) if attachment_source else ""
    inicio = validate_date(data_inicio, "data inicial", allow_empty=True)
    fim = validate_date(data_fim, "data final", allow_empty=True)
    conn = get_conn()
    conn.execute("""
        INSERT INTO erp_records (
            module_code, codigo, categoria, titulo, partner_code, product_code,
            employee_code, project_code, cost_center, data_inicio, data_fim,
            status, valor_previsto, valor_real, prioridade, responsavel,
            descricao, observacoes, attachment_path, created_by, created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        module_code,
        codigo,
        categoria or "",
        titulo,
        partner_code or None,
        product_code or None,
        employee_code or "",
        project_code or "",
        (cost_center or "").strip(),
        inicio,
        fim,
        status or "",
        parse_float_br(valor_previsto),
        parse_float_br(valor_real),
        prioridade or "",
        (responsavel or "").strip(),
        (descricao or "").strip(),
        (observacoes or "").strip(),
        attachment_path,
        created_by,
        now_iso(),
        now_iso()
    ))
    conn.commit()
    conn.close()
    return codigo

def update_erp_record(record_id, module_code, categoria, titulo, partner_code, product_code, employee_code, project_code, cost_center, data_inicio, data_fim, status, valor_previsto, valor_real, prioridade, responsavel, descricao, observacoes, attachment_source="", current_attachment=""):
    if not record_id:
        raise ValueError("Selecione um registro para editar.")

    titulo = (titulo or "").strip()

    if not titulo:
        raise ValueError("Informe o título do registro.")

    attachment_path = current_attachment or ""

    if attachment_source:
        attachment_path = copy_attachment(attachment_source, f"ERP-{record_id}")

    conn = get_conn()
    conn.execute("""
        UPDATE erp_records
        SET categoria = ?,
            titulo = ?,
            partner_code = ?,
            product_code = ?,
            employee_code = ?,
            project_code = ?,
            cost_center = ?,
            data_inicio = ?,
            data_fim = ?,
            status = ?,
            valor_previsto = ?,
            valor_real = ?,
            prioridade = ?,
            responsavel = ?,
            descricao = ?,
            observacoes = ?,
            attachment_path = ?,
            updated_at = ?
        WHERE id = ?
          AND module_code = ?
    """, (
        categoria or "",
        titulo,
        partner_code or None,
        product_code or None,
        employee_code or "",
        project_code or "",
        (cost_center or "").strip(),
        validate_date(data_inicio, "data inicial", allow_empty=True),
        validate_date(data_fim, "data final", allow_empty=True),
        status or "",
        parse_float_br(valor_previsto),
        parse_float_br(valor_real),
        prioridade or "",
        (responsavel or "").strip(),
        (descricao or "").strip(),
        (observacoes or "").strip(),
        attachment_path,
        now_iso(),
        record_id,
        module_code
    ))
    conn.commit()
    conn.close()
    return record_id

def get_erp_records(module_code):
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            er.*,
            p.nome_razao AS parceiro,
            pr.nome AS produto
        FROM erp_records er
        LEFT JOIN partners p ON p.codigo = er.partner_code
        LEFT JOIN products pr ON pr.codigo = er.product_code
        WHERE er.module_code = ?
        ORDER BY er.updated_at DESC, er.id DESC
    """, (module_code,)).fetchall()
    conn.close()

    result = []

    for row in rows:
        item = dict(row)
        item["data_inicio"] = format_date_br(item.get("data_inicio", ""))
        item["data_fim"] = format_date_br(item.get("data_fim", ""))
        item["valor_previsto"] = format_money(item.get("valor_previsto"))
        item["valor_real"] = format_money(item.get("valor_real"))
        item["anexo"] = "Sim" if item.get("attachment_path") else ""
        result.append(item)

    return result

def get_erp_record_by_id(record_id, module_code):
    conn = get_conn()
    row = conn.execute(
        "SELECT * FROM erp_records WHERE id = ? AND module_code = ?",
        (record_id, module_code)
    ).fetchone()
    conn.close()
    return dict(row) if row else None

def calculate_proposal_values(horas_normais, horas_50, horas_100, valor_hora, custos_viagem, custos_materiais, custos_terceiros, impostos_percentual, margem_percentual):
    base_horas = (
        parse_float_br(horas_normais) * parse_float_br(valor_hora)
        + parse_float_br(horas_50) * parse_float_br(valor_hora) * 1.5
        + parse_float_br(horas_100) * parse_float_br(valor_hora) * 2
    )
    custo_total = base_horas + parse_float_br(custos_viagem) + parse_float_br(custos_materiais) + parse_float_br(custos_terceiros)
    impostos = parse_float_br(impostos_percentual) / 100
    margem = parse_float_br(margem_percentual) / 100
    preco_venda = custo_total

    if impostos + margem < 0.95:
        preco_venda = custo_total / max(1 - impostos - margem, 0.05)

    lucro_previsto = preco_venda - custo_total
    return custo_total, preco_venda, lucro_previsto

def create_proposal(partner_code, projeto, revisao, data_emissao, validade, prazo, pagamento, condicoes, horas_normais, horas_50, horas_100, valor_hora, custos_viagem, custos_materiais, custos_terceiros, impostos_percentual, margem_percentual, incluso, nao_incluso, status, motivo_status, attachment_source=""):
    if not partner_code:
        raise ValueError("Selecione o cliente.")

    projeto = (projeto or "").strip()

    if not projeto:
        raise ValueError("Informe o projeto/escopo da proposta.")

    codigo = get_scoped_code("commercial_proposals", "PROP")
    attachment_path = copy_attachment(attachment_source, codigo) if attachment_source else ""
    custo_total, preco_venda, lucro_previsto = calculate_proposal_values(
        horas_normais, horas_50, horas_100, valor_hora,
        custos_viagem, custos_materiais, custos_terceiros,
        impostos_percentual, margem_percentual
    )
    conn = get_conn()
    conn.execute("""
        INSERT INTO commercial_proposals (
            codigo, partner_code, projeto, revisao, data_emissao, validade,
            prazo, pagamento, condicoes, horas_normais, horas_50, horas_100,
            valor_hora, custos_viagem, custos_materiais, custos_terceiros,
            impostos_percentual, margem_percentual, custo_total, preco_venda,
            lucro_previsto, incluso, nao_incluso, status, motivo_status,
            attachment_path, created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        codigo,
        partner_code,
        projeto,
        int(parse_float_br(revisao)),
        validate_date(data_emissao or today_br(), "data de emissão", allow_empty=False),
        validate_date(validade, "validade", allow_empty=True),
        (prazo or "").strip(),
        (pagamento or "").strip(),
        (condicoes or "").strip(),
        parse_float_br(horas_normais),
        parse_float_br(horas_50),
        parse_float_br(horas_100),
        parse_float_br(valor_hora),
        parse_float_br(custos_viagem),
        parse_float_br(custos_materiais),
        parse_float_br(custos_terceiros),
        parse_float_br(impostos_percentual),
        parse_float_br(margem_percentual),
        custo_total,
        preco_venda,
        lucro_previsto,
        (incluso or "").strip(),
        (nao_incluso or "").strip(),
        status or "Em elaboração",
        (motivo_status or "").strip(),
        attachment_path,
        now_iso(),
        now_iso()
    ))
    conn.commit()
    conn.close()
    return codigo

def update_proposal(record_id, partner_code, projeto, revisao, data_emissao, validade, prazo, pagamento, condicoes, horas_normais, horas_50, horas_100, valor_hora, custos_viagem, custos_materiais, custos_terceiros, impostos_percentual, margem_percentual, incluso, nao_incluso, status, motivo_status, attachment_source="", current_attachment=""):
    if not record_id:
        raise ValueError("Selecione uma proposta para editar.")

    if not partner_code:
        raise ValueError("Selecione o cliente.")

    projeto = (projeto or "").strip()

    if not projeto:
        raise ValueError("Informe o projeto/escopo da proposta.")

    attachment_path = current_attachment or ""

    if attachment_source:
        attachment_path = copy_attachment(attachment_source, f"PROP-{record_id}")

    custo_total, preco_venda, lucro_previsto = calculate_proposal_values(
        horas_normais, horas_50, horas_100, valor_hora,
        custos_viagem, custos_materiais, custos_terceiros,
        impostos_percentual, margem_percentual
    )
    conn = get_conn()
    conn.execute("""
        UPDATE commercial_proposals
        SET partner_code = ?,
            projeto = ?,
            revisao = ?,
            data_emissao = ?,
            validade = ?,
            prazo = ?,
            pagamento = ?,
            condicoes = ?,
            horas_normais = ?,
            horas_50 = ?,
            horas_100 = ?,
            valor_hora = ?,
            custos_viagem = ?,
            custos_materiais = ?,
            custos_terceiros = ?,
            impostos_percentual = ?,
            margem_percentual = ?,
            custo_total = ?,
            preco_venda = ?,
            lucro_previsto = ?,
            incluso = ?,
            nao_incluso = ?,
            status = ?,
            motivo_status = ?,
            attachment_path = ?,
            updated_at = ?
        WHERE id = ?
    """, (
        partner_code,
        projeto,
        int(parse_float_br(revisao)),
        validate_date(data_emissao or today_br(), "data de emissão", allow_empty=False),
        validate_date(validade, "validade", allow_empty=True),
        (prazo or "").strip(),
        (pagamento or "").strip(),
        (condicoes or "").strip(),
        parse_float_br(horas_normais),
        parse_float_br(horas_50),
        parse_float_br(horas_100),
        parse_float_br(valor_hora),
        parse_float_br(custos_viagem),
        parse_float_br(custos_materiais),
        parse_float_br(custos_terceiros),
        parse_float_br(impostos_percentual),
        parse_float_br(margem_percentual),
        custo_total,
        preco_venda,
        lucro_previsto,
        (incluso or "").strip(),
        (nao_incluso or "").strip(),
        status or "Em elaboração",
        (motivo_status or "").strip(),
        attachment_path,
        now_iso(),
        record_id
    ))
    conn.commit()
    conn.close()
    return record_id

def get_proposals():
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            cp.*,
            p.nome_razao AS cliente
        FROM commercial_proposals cp
        LEFT JOIN partners p ON p.codigo = cp.partner_code
        ORDER BY cp.updated_at DESC, cp.id DESC
    """).fetchall()
    conn.close()
    result = []

    for row in rows:
        item = dict(row)
        item["revisao_label"] = f"Rev. {int(item.get('revisao') or 0):02d}"
        item["data_emissao"] = format_date_br(item.get("data_emissao", ""))
        item["validade"] = format_date_br(item.get("validade", ""))
        item["custo_total"] = format_money(item.get("custo_total"))
        item["preco_venda"] = format_money(item.get("preco_venda"))
        item["lucro_previsto"] = format_money(item.get("lucro_previsto"))
        item["margem_percentual"] = format_money(item.get("margem_percentual"))
        item["anexo"] = "Sim" if item.get("attachment_path") else ""
        result.append(item)

    return result

def get_proposal_by_id(record_id):
    conn = get_conn()
    row = conn.execute("""
        SELECT
            cp.*,
            p.nome_razao AS cliente
        FROM commercial_proposals cp
        LEFT JOIN partners p ON p.codigo = cp.partner_code
        WHERE cp.id = ?
    """, (record_id,)).fetchone()
    conn.close()
    return dict(row) if row else None

def create_financial_entry(tipo, descricao, partner_code, document_id, project_code, cost_center, valor, data_emissao, data_vencimento, status, forma_pagamento, conta_bancaria, observacoes, data_baixa="", attachment_source=""):
    descricao = (descricao or "").strip()

    if not descricao:
        raise ValueError("Informe a descrição financeira.")

    prefix = "REC" if tipo == "Entrada" else "PAG"
    codigo = get_scoped_code("financial_entries", prefix, "tipo = ?", (tipo,))
    attachment_path = copy_attachment(attachment_source, codigo) if attachment_source else ""
    conn = get_conn()
    conn.execute("""
        INSERT INTO financial_entries (
            codigo, tipo, descricao, partner_code, document_id, project_code,
            cost_center, valor, data_emissao, data_vencimento, data_baixa, status,
            forma_pagamento, conta_bancaria, observacoes, attachment_path,
            created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        codigo,
        tipo,
        descricao,
        partner_code or None,
        document_id or None,
        project_code or "",
        (cost_center or "").strip(),
        parse_float_br(valor),
        validate_date(data_emissao, "data de emissão", allow_empty=True),
        validate_date(data_vencimento, "data de vencimento", allow_empty=True),
        validate_date(data_baixa, "data de baixa", allow_empty=True),
        status or "Aberto",
        (forma_pagamento or "").strip(),
        (conta_bancaria or "").strip(),
        (observacoes or "").strip(),
        attachment_path,
        now_iso(),
        now_iso()
    ))
    conn.commit()
    conn.close()
    return codigo

def update_financial_entry(record_id, tipo, descricao, partner_code, document_id, project_code, cost_center, valor, data_emissao, data_vencimento, data_baixa, status, forma_pagamento, conta_bancaria, observacoes, attachment_source="", current_attachment=""):
    if not record_id:
        raise ValueError("Selecione um lançamento para editar.")

    descricao = (descricao or "").strip()

    if not descricao:
        raise ValueError("Informe a descrição financeira.")

    attachment_path = current_attachment or ""

    if attachment_source:
        attachment_path = copy_attachment(attachment_source, f"FIN-{record_id}")

    conn = get_conn()
    conn.execute("""
        UPDATE financial_entries
        SET tipo = ?,
            descricao = ?,
            partner_code = ?,
            document_id = ?,
            project_code = ?,
            cost_center = ?,
            valor = ?,
            data_emissao = ?,
            data_vencimento = ?,
            data_baixa = ?,
            status = ?,
            forma_pagamento = ?,
            conta_bancaria = ?,
            observacoes = ?,
            attachment_path = ?,
            updated_at = ?
        WHERE id = ?
    """, (
        tipo,
        descricao,
        partner_code or None,
        document_id or None,
        project_code or "",
        (cost_center or "").strip(),
        parse_float_br(valor),
        validate_date(data_emissao, "data de emissão", allow_empty=True),
        validate_date(data_vencimento, "data de vencimento", allow_empty=True),
        validate_date(data_baixa, "data de baixa", allow_empty=True),
        status or "Aberto",
        (forma_pagamento or "").strip(),
        (conta_bancaria or "").strip(),
        (observacoes or "").strip(),
        attachment_path,
        now_iso(),
        record_id
    ))
    conn.commit()
    conn.close()
    return record_id

def get_financial_entries():
    conn = get_conn()
    rows = conn.execute("""
        SELECT
            fe.*,
            p.nome_razao AS parceiro
        FROM financial_entries fe
        LEFT JOIN partners p ON p.codigo = fe.partner_code
        ORDER BY fe.data_vencimento ASC, fe.id DESC
    """).fetchall()
    conn.close()
    result = []

    for row in rows:
        item = dict(row)
        item["data_emissao"] = format_date_br(item.get("data_emissao", ""))
        item["data_vencimento"] = format_date_br(item.get("data_vencimento", ""))
        item["data_baixa"] = format_date_br(item.get("data_baixa", ""))
        item["valor"] = format_money(item.get("valor"))
        item["anexo"] = "Sim" if item.get("attachment_path") else ""
        result.append(item)

    return result

def get_financial_entry_by_id(record_id):
    conn = get_conn()
    row = conn.execute("SELECT * FROM financial_entries WHERE id = ?", (record_id,)).fetchone()
    conn.close()
    return dict(row) if row else None

def export_xlsx(default_name: str, columns: list, rows: list):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        messagebox.showerror("Exportação XLSX", "Instale a dependência openpyxl.")
        return

    ensure_app_dirs()
    path = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel", "*.xlsx")],
        initialdir=str(EXPORTS_DIR),
        initialfile=f"{default_name}.xlsx"
    )

    if not path:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append(columns)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for row in rows:
        ws.append([row.get(col, "") for col in columns])

    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 45)

    wb.save(path)
    messagebox.showinfo("Exportação XLSX", "Arquivo XLSX exportado com sucesso.")

def generate_document_pdf(document_id):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError:
        messagebox.showerror("PDF", "Instale a dependência reportlab.")
        return

    document, items = get_document_by_id(document_id)

    if not document:
        messagebox.showerror("PDF", "Documento não encontrado.")
        return

    ensure_app_dirs()
    path = filedialog.asksaveasfilename(
        defaultextension=".pdf",
        filetypes=[("PDF", "*.pdf")],
        initialdir=str(PDFS_DIR),
        initialfile=f"{document['numero']}.pdf"
    )

    if not path:
        return

    font_name = "Helvetica"
    arial = Path("C:/Windows/Fonts/arial.ttf")

    if arial.exists():
        try:
            pdfmetrics.registerFont(TTFont("Arial", str(arial)))
            font_name = "Arial"
        except Exception:
            font_name = "Helvetica"

    styles = getSampleStyleSheet()
    styles["Normal"].fontName = font_name
    styles["Title"].fontName = font_name
    story = []
    doc = SimpleDocTemplate(path, pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)

    story.append(Paragraph(APP_NAME, styles["Title"]))
    story.append(Paragraph(f"{document.get('tipo_documento', '')} - {document.get('numero', '')}", styles["Title"]))
    story.append(Spacer(1, 0.35 * cm))
    story.append(Paragraph(f"Emissão: {format_date_br(document.get('data_emissao', ''))}", styles["Normal"]))
    story.append(Paragraph(f"Parceiro: {document.get('parceiro', '')}", styles["Normal"]))
    story.append(Paragraph(f"CPF/CNPJ: {format_document(document.get('parceiro_documento', ''))}", styles["Normal"]))
    story.append(Paragraph(f"Status: {document.get('status', '')}", styles["Normal"]))
    story.append(Spacer(1, 0.35 * cm))

    data = [["Produto", "Qtd.", "Un.", "Vlr. unit.", "Desc.", "Total"]]

    for item in items:
        data.append([
            item.get("produto_nome", ""),
            format_money(item.get("quantidade")),
            item.get("unidade", ""),
            format_money(item.get("valor_unitario")),
            format_money(item.get("desconto")),
            format_money(item.get("total")),
        ])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    story.append(Spacer(1, 0.35 * cm))
    story.append(Paragraph(f"Subtotal: R$ {format_money(document.get('subtotal'))}", styles["Normal"]))
    story.append(Paragraph(f"Desconto: R$ {format_money(document.get('desconto_total'))}", styles["Normal"]))
    story.append(Paragraph(f"Acréscimo: R$ {format_money(document.get('acrescimo_total'))}", styles["Normal"]))
    story.append(Paragraph(f"Total geral: R$ {format_money(document.get('total'))}", styles["Title"]))
    story.append(Spacer(1, 0.45 * cm))
    story.append(Paragraph(f"Observações: {document.get('observacoes', '')}", styles["Normal"]))
    story.append(Spacer(1, 1.0 * cm))
    story.append(Paragraph("Assinatura: ______________________________________________", styles["Normal"]))

    doc.build(story)
    messagebox.showinfo("PDF", "PDF gerado com sucesso.")
    return path

def generate_table_pdf(default_name: str, title: str, columns: list, rows: list):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError:
        messagebox.showerror("PDF", "Instale a dependência reportlab.")
        return

    ensure_app_dirs()
    path = filedialog.asksaveasfilename(
        defaultextension=".pdf",
        filetypes=[("PDF", "*.pdf")],
        initialdir=str(PDFS_DIR),
        initialfile=f"{default_name}.pdf"
    )

    if not path:
        return

    font_name = "Helvetica"
    arial = Path("C:/Windows/Fonts/arial.ttf")

    if arial.exists():
        try:
            pdfmetrics.registerFont(TTFont("Arial", str(arial)))
            font_name = "Arial"
        except Exception:
            font_name = "Helvetica"

    styles = getSampleStyleSheet()
    styles["Normal"].fontName = font_name
    styles["Title"].fontName = font_name
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), rightMargin=1.0 * cm, leftMargin=1.0 * cm, topMargin=1.0 * cm, bottomMargin=1.0 * cm)
    story = [Paragraph(APP_NAME, styles["Title"]), Paragraph(title, styles["Title"]), Spacer(1, 0.35 * cm)]
    data = [columns]

    for row in rows:
        data.append([str(row.get(col, ""))[:80] for col in columns])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    doc.build(story)
    messagebox.showinfo("PDF", "PDF gerado com sucesso.")
    return path

def export_csv(default_name: str, columns: list, rows: list):
    ensure_app_dirs()
    path = filedialog.asksaveasfilename(
        defaultextension=".csv",
        filetypes=[("CSV", "*.csv")],
        initialdir=str(EXPORTS_DIR),
        initialfile=f"{default_name}.csv"
    )

    if not path:
        return

    with open(path, "w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file, delimiter=";")
        writer.writerow(columns)

        for row in rows:
            writer.writerow([row.get(col, "") for col in columns])

    messagebox.showinfo("Exportação", "Arquivo exportado com sucesso.")

class TreeFrame(ttk.Frame):
    def __init__(self, master, columns):
        super().__init__(master)

        self.tree = ttk.Treeview(
            self,
            columns=columns,
            show="headings",
            height=15
        )

        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=130, anchor="w")

        y_scroll = ttk.Scrollbar(self, orient="vertical", command=self.tree.yview)
        x_scroll = ttk.Scrollbar(self, orient="horizontal", command=self.tree.xview)

        self.tree.configure(
            yscrollcommand=y_scroll.set,
            xscrollcommand=x_scroll.set
        )

        self.tree.tag_configure("VENCIDO", background="#3B1114", foreground="#FCA5A5")
        self.tree.tag_configure("URGENTE", background="#3A2F08", foreground="#FDE68A")
        self.tree.tag_configure("ATENCAO", background="#3A2F08", foreground="#FDE68A")
        self.tree.tag_configure("OK", background="#0F2F24", foreground="#86EFAC")

        self.tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")

        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, weight=1)

class ScrollableFrame(ttk.Frame):
    def __init__(self, master):
        super().__init__(master)

        self.canvas = tk.Canvas(
            self,
            borderwidth=0,
            highlightthickness=0,
            background=COLOR_BG
        )
        self.scrollbar = ttk.Scrollbar(
            self,
            orient="vertical",
            command=self.canvas.yview
        )

        self.scrollable_frame = ttk.Frame(self.canvas, style="App.TFrame")

        self.scrollable_frame.bind(
            "<Configure>",
            lambda event: self.canvas.configure(
                scrollregion=self.canvas.bbox("all")
            )
        )

        self.canvas_window = self.canvas.create_window(
            (0, 0),
            window=self.scrollable_frame,
            anchor="nw"
        )

        self.canvas.configure(yscrollcommand=self.scrollbar.set)

        self.canvas.pack(side="left", fill="both", expand=True)
        self.scrollbar.pack(side="right", fill="y")

        self.canvas.bind("<Configure>", self.resize_frame)
        self.canvas.bind("<Enter>", self.bind_mousewheel)
        self.canvas.bind("<Leave>", self.unbind_mousewheel)
        self.scrollable_frame.bind("<Enter>", self.bind_mousewheel)
        self.scrollable_frame.bind("<Leave>", self.unbind_mousewheel)

    def resize_frame(self, event):
        self.canvas.itemconfig(self.canvas_window, width=event.width)

    def bind_mousewheel(self, event=None):
        self.canvas.bind_all("<MouseWheel>", self.on_mousewheel)

    def unbind_mousewheel(self, event=None):
        self.canvas.unbind_all("<MouseWheel>")

    def on_mousewheel(self, event):
        self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

class ModernButton(tk.Canvas):
    def __init__(self, master, text, command=None, variant="primary", height=36, width=132):
        super().__init__(master, height=height, width=width, bg=COLOR_BG, highlightthickness=0, bd=0, cursor="hand2")
        self.text = text
        self.command = command
        self.variant = variant
        self.hover = False
        self.bind("<Configure>", lambda _event: self.redraw())
        self.bind("<Enter>", lambda _event: self.set_hover(True))
        self.bind("<Leave>", lambda _event: self.set_hover(False))
        self.bind("<Button-1>", lambda _event: self.invoke())
        self.redraw()

    def colors(self):
        if self.variant == "primary":
            return (COLOR_PRIMARY_HOVER if self.hover else COLOR_PRIMARY, COLOR_PRIMARY_HOVER, "#FFFFFF")
        if self.variant == "danger":
            return ("#B91C1C" if self.hover else COLOR_DANGER, "#B91C1C", "#FFFFFF")
        return (COLOR_BUTTON_HOVER if self.hover else COLOR_BUTTON, COLOR_BORDER, COLOR_TEXT)

    def set_hover(self, value):
        self.hover = value
        self.redraw()

    def invoke(self):
        if self.command:
            self.command()

    def redraw(self):
        self.delete("all")
        fill, outline, text_color = self.colors()
        width = max(self.winfo_width(), 80)
        height = max(self.winfo_height(), 30)
        draw_round_rect(self, 1, 1, width - 1, height - 1, 10, fill, outline, 1)
        self.create_text(width / 2, height / 2, text=self.text, fill=text_color, font=("Segoe UI", 10, "bold"))

class BaseFrame(ttk.Frame):
    def make_label(self, parent, text, row, col, sticky="w"):
        label = ttk.Label(parent, text=text)
        label.grid(row=row, column=col, padx=6, pady=(6, 2), sticky=sticky)
        return label

    def make_entry(self, parent, row, col, width=30, show=None):
        entry = ttk.Entry(parent, width=width, show=show)
        entry.grid(row=row, column=col, padx=6, pady=(0, 6), sticky="ew")
        return entry

    def make_date_entry(self, parent, row, col, width=16):
        entry = self.make_entry(parent, row, col, width=width)
        self.bind_date_entry(entry)
        return entry

    def make_combo(self, parent, row, col, values, width=28):
        combo = ttk.Combobox(parent, values=values, width=width, state="readonly")
        combo.grid(row=row, column=col, padx=6, pady=(0, 6), sticky="ew")

        if values:
            combo.set(values[0])

        return combo

    def clear_entries(self, *entries):
        for entry in entries:
            entry.delete(0, "end")

    def set_entry(self, entry, value):
        entry.delete(0, "end")
        entry.insert(0, "" if value is None else str(value))

    def set_combo(self, combo, value):
        combo.set("" if value is None else str(value))

    def bind_date_entry(self, entry):
        entry.bind("<FocusOut>", lambda _event: self.normalize_date_entry(entry))
        entry.bind("<Return>", lambda _event: self.normalize_date_entry(entry))
        return entry

    def normalize_date_entry(self, entry):
        value = entry.get().strip()

        if not value:
            return

        try:
            formatted = format_date_br(validate_date(value, "data", allow_empty=True))
        except ValueError:
            return

        entry.delete(0, "end")
        entry.insert(0, formatted)

    def create_tree(self, parent, columns):
        return TreeFrame(parent, columns)

    def fill_tree(self, tree_frame, rows, columns):
        tree = tree_frame.tree
        tree.delete(*tree.get_children())

        for row in rows:
            values = [row.get(col, "") for col in columns]

            nivel = str(row.get("nivel", "")).upper().strip()
            status = str(row.get("status", "")).lower().strip()
            tag = ""

            if nivel == "VENCIDO" or "vencido" in status or "atrasado" in status:
                tag = "VENCIDO"
            elif nivel == "URGENTE" or "hoje" in status or "amanhã" in status or "amanha" in status:
                tag = "URGENTE"
            elif nivel in ("ATENÇÃO", "ATENCAO") or "vence em" in status or "faltam" in status:
                tag = "ATENCAO"
            elif nivel == "OK":
                tag = "OK"

            if tag:
                tree.insert("", "end", values=values, tags=(tag,))
            else:
                tree.insert("", "end", values=values)

    def get_selected_row(self, tree_frame, columns):
        selected = tree_frame.tree.selection()

        if not selected:
            return None

        values = tree_frame.tree.item(selected[0], "values")
        return dict(zip(columns, values))

    def bind_double_click(self, tree_frame, columns, callback):
        def handle_double_click(event=None):
            row = self.get_selected_row(tree_frame, columns)

            if row:
                callback(row)

        tree_frame.tree.bind("<Double-1>", handle_double_click)

    def open_item_detail(self, item_codigo):
        if item_codigo:
            ItemDetailWindow(self.winfo_toplevel(), item_codigo)

    def open_employee_detail(self, funcionario_codigo):
        if funcionario_codigo:
            EmployeeDetailWindow(self.winfo_toplevel(), funcionario_codigo)

    def open_client_detail(self, cliente_codigo):
        if cliente_codigo:
            ClientDetailWindow(self.winfo_toplevel(), cliente_codigo)

    def open_service_detail(self, codigo_servico):
        if codigo_servico:
            ServiceDetailWindow(self.winfo_toplevel(), codigo_servico)

    def open_product_detail(self, codigo):
        if codigo:
            ProductDetailWindow(self.winfo_toplevel(), codigo)

    def open_partner_detail(self, codigo):
        if codigo:
            PartnerDetailWindow(self.winfo_toplevel(), codigo)

    def open_document_detail(self, document_id):
        if document_id:
            DocumentDetailWindow(self.winfo_toplevel(), document_id)

class DetailWindow(tk.Toplevel):
    def __init__(self, master, title, geometry="980x700"):
        super().__init__(master)
        self.title(title)
        self.geometry(geometry)
        self.minsize(760, 520)
        self.configure(background=COLOR_BG)
        self.preview_image = None

    def add_section_title(self, parent, text):
        ttk.Label(
            parent,
            text=text,
            font=("Segoe UI", 13, "bold"),
            style="Title.TLabel"
        ).pack(anchor="w", pady=(12, 6))

    def fill_tree(self, tree_frame, rows, columns):
        tree = tree_frame.tree
        tree.delete(*tree.get_children())

        for row in rows:
            tree.insert("", "end", values=[row.get(col, "") for col in columns])

    def add_info_grid(self, parent, rows):
        grid = ttk.Frame(parent, style="Card.TFrame", padding=10)
        grid.pack(fill="x", pady=(0, 8))
        grid.columnconfigure(1, weight=1)
        grid.columnconfigure(3, weight=1)

        for index, (label, value) in enumerate(rows):
            row = index // 2
            col = (index % 2) * 2
            ttk.Label(grid, text=f"{label}:", style="Muted.TLabel").grid(row=row, column=col, sticky="w", padx=(0, 6), pady=3)
            ttk.Label(grid, text=str(value or ""), wraplength=320).grid(row=row, column=col + 1, sticky="w", pady=3)

class ItemDetailWindow(DetailWindow):
    def __init__(self, master, item_codigo):
        self.item = get_item_by_code(item_codigo)

        if not self.item:
            messagebox.showerror("Item", "Item nao encontrado.")
            return

        super().__init__(master, f"Item {item_codigo}")
        self.build()

    def build(self):
        scroll = ScrollableFrame(self)
        scroll.pack(fill="both", expand=True)
        container = scroll.scrollable_frame

        header = ttk.Frame(container, style="Card.TFrame", padding=12)
        header.pack(fill="x", padx=12, pady=12)
        header.columnconfigure(1, weight=1)

        self.preview_image = load_tk_image(self.item.get("image_path"), (220, 220))

        image_box = ttk.Frame(header, style="Card.TFrame")
        image_box.grid(row=0, column=0, rowspan=2, sticky="nw", padx=(0, 14))

        if self.preview_image:
            ttk.Label(image_box, image=self.preview_image, style="Card.TLabel").pack()
        else:
            ttk.Label(
                image_box,
                text="Sem imagem",
                style="Muted.TLabel",
                width=24,
                anchor="center"
            ).pack(ipadx=18, ipady=44)

        ttk.Label(
            header,
            text=self.item.get("nome", ""),
            font=("Segoe UI", 20, "bold"),
            style="Title.TLabel"
        ).grid(row=0, column=1, sticky="w")

        ttk.Label(
            header,
            text=f"Codigo {self.item.get('codigo', '')} | Categoria {self.item.get('categoria', '')}",
            style="Muted.TLabel"
        ).grid(row=1, column=1, sticky="w", pady=(4, 0))

        info_rows = [
            ("Unidade", self.item.get("unidade")),
            ("CA", self.item.get("ca")),
            ("Valor unitario", f"R$ {float(self.item.get('valor_unitario') or 0):.2f}"),
            ("Saldo atual", self.item.get("saldo")),
            ("Validade", self.item.get("validade")),
            ("Imagem", self.item.get("image_path")),
            ("Observacoes", self.item.get("observacoes")),
        ]

        self.add_info_grid(container, info_rows)

        movement_columns = [
            "tipo",
            "quantidade",
            "data_movimentacao",
            "validade",
            "funcionario_nome",
            "service_codigo",
            "destino",
            "observacoes"
        ]
        self.add_section_title(container, "Movimentacoes relacionadas")
        movement_tree = TreeFrame(container, movement_columns)
        movement_tree.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        self.fill_tree(movement_tree, get_item_movements(self.item["codigo"]), movement_columns)

        holder_columns = ["funcionario_codigo", "funcionario_nome", "quantidade", "ultima_movimentacao", "validade"]
        self.add_section_title(container, "Funcionarios com este item")
        holder_tree = TreeFrame(container, holder_columns)
        holder_tree.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        self.fill_tree(holder_tree, get_item_employee_holders(self.item["codigo"]), holder_columns)

        service_columns = ["codigo_servico", "nome_servico", "nome_razao", "status", "prazo"]
        self.add_section_title(container, "Servicos relacionados")
        service_tree = TreeFrame(container, service_columns)
        service_tree.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        self.fill_tree(service_tree, get_item_services(self.item["codigo"]), service_columns)

        expense_columns = ["funcionario_codigo", "funcionario_nome", "quantidade", "valor_unitario", "valor_total", "data_movimentacao", "service_codigo"]
        self.add_section_title(container, "Gastos relacionados")
        expense_tree = TreeFrame(container, expense_columns)
        expense_tree.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        self.fill_tree(expense_tree, get_item_epi_expenses(self.item["codigo"]), expense_columns)

class EmployeeDetailWindow(DetailWindow):
    def __init__(self, master, employee_codigo):
        self.employee = get_employee_by_code(employee_codigo)

        if not self.employee:
            messagebox.showerror("Funcionario", "Funcionario nao encontrado.")
            return

        super().__init__(master, f"Funcionario {employee_codigo}", "820x620")
        self.build()

    def build(self):
        container = ttk.Frame(self, padding=14)
        container.pack(fill="both", expand=True)

        ttk.Label(container, text=self.employee.get("nome", ""), font=("Segoe UI", 18, "bold"), style="Title.TLabel").pack(anchor="w")

        self.add_info_grid(container, [
            ("Codigo", self.employee.get("codigo")),
            ("Documento", format_document(self.employee.get("documento", ""))),
            ("Cargo", self.employee.get("cargo")),
            ("Telefone", self.employee.get("telefone")),
            ("E-mail", self.employee.get("email")),
            ("Status", self.employee.get("status")),
            ("Observacoes", self.employee.get("observacoes")),
        ])

        columns = ["item_codigo", "item_nome", "quantidade", "ultima_movimentacao", "validade"]
        rows = [row for row in get_employee_items() if row.get("funcionario_codigo") == self.employee.get("codigo")]

        self.add_section_title(container, "Itens com este funcionario")
        tree_frame = TreeFrame(container, columns)
        tree_frame.pack(fill="both", expand=True)
        self.fill_tree(tree_frame, rows, columns)

class ClientDetailWindow(DetailWindow):
    def __init__(self, master, client_codigo):
        self.client = get_client_by_code(client_codigo)

        if not self.client:
            messagebox.showerror("Cliente", "Cliente/empresa nao encontrado.")
            return

        super().__init__(master, f"Cliente {client_codigo}", "860x620")
        self.build()

    def build(self):
        container = ttk.Frame(self, padding=14)
        container.pack(fill="both", expand=True)

        ttk.Label(container, text=self.client.get("nome_razao", ""), font=("Segoe UI", 18, "bold"), style="Title.TLabel").pack(anchor="w")

        self.add_info_grid(container, [
            ("Codigo", self.client.get("codigo")),
            ("Tipo", self.client.get("tipo_pessoa")),
            ("Documento", self.client.get("documento")),
            ("Telefone", self.client.get("telefone")),
            ("E-mail", self.client.get("email")),
            ("Responsavel", self.client.get("responsavel")),
            ("Contato", self.client.get("contato")),
            ("Endereco", self.client.get("endereco")),
            ("Observacoes", self.client.get("observacoes")),
        ])

        columns = ["codigo_servico", "nome_razao", "nome_servico", "descricao", "valor", "prazo", "status"]
        rows = get_services(self.client.get("codigo"))

        self.add_section_title(container, "Historico de servicos")
        tree_frame = TreeFrame(container, columns)
        tree_frame.pack(fill="both", expand=True)
        self.fill_tree(tree_frame, rows, columns)

class ServiceDetailWindow(DetailWindow):
    def __init__(self, master, service_codigo):
        self.service = get_service_by_code(service_codigo)

        if not self.service:
            messagebox.showerror("Servico", "Servico nao encontrado.")
            return

        super().__init__(master, f"Servico {service_codigo}", "860x650")
        self.build()

    def build(self):
        container = ttk.Frame(self, padding=14)
        container.pack(fill="both", expand=True)

        ttk.Label(container, text=self.service.get("nome_servico", ""), font=("Segoe UI", 18, "bold"), style="Title.TLabel").pack(anchor="w")

        self.add_info_grid(container, [
            ("Codigo", self.service.get("codigo_servico")),
            ("Cliente", self.service.get("nome_razao")),
            ("Valor", self.service.get("valor")),
            ("Prazo", self.service.get("prazo")),
            ("Status", self.service.get("status")),
            ("Descricao", self.service.get("descricao")),
            ("Observacoes", self.service.get("observacoes")),
        ])

        employee_columns = ["codigo", "nome"]
        required_columns = ["codigo", "nome"]

        self.add_section_title(container, "Funcionarios vinculados")
        employee_tree = TreeFrame(container, employee_columns)
        employee_tree.pack(fill="both", expand=True, pady=(0, 8))
        self.fill_tree(employee_tree, get_service_employees(self.service["codigo_servico"]), employee_columns)

        self.add_section_title(container, "EPIs obrigatorios")
        required_tree = TreeFrame(container, required_columns)
        required_tree.pack(fill="both", expand=True)
        self.fill_tree(required_tree, get_service_required_items(self.service["codigo_servico"]), required_columns)

class ProductDetailWindow(DetailWindow):
    def __init__(self, master, codigo):
        self.product = get_product_by_code(codigo)

        if not self.product:
            messagebox.showerror("Produto", "Produto não encontrado.")
            return

        super().__init__(master, f"Produto {codigo}", "940x680")
        self.build()

    def build(self):
        scroll = ScrollableFrame(self)
        scroll.pack(fill="both", expand=True)
        container = scroll.scrollable_frame
        header = ttk.Frame(container, style="Card.TFrame", padding=12)
        header.pack(fill="x", padx=12, pady=12)
        header.columnconfigure(1, weight=1)
        self.preview_image = load_tk_image(self.product.get("image_path"), (210, 210))

        if self.preview_image:
            ttk.Label(header, image=self.preview_image, style="Card.TLabel").grid(row=0, column=0, rowspan=2, padx=(0, 14), sticky="nw")
        else:
            ttk.Label(header, text="Sem imagem", style="Muted.TLabel", width=24, anchor="center").grid(row=0, column=0, rowspan=2, padx=(0, 14), sticky="nw")

        ttk.Label(header, text=self.product.get("nome", ""), font=("Segoe UI", 20, "bold"), style="Title.TLabel").grid(row=0, column=1, sticky="w")
        ttk.Label(header, text=f"Código {self.product.get('codigo', '')} | NCM {self.product.get('ncm', '')}", style="Muted.TLabel").grid(row=1, column=1, sticky="w")

        self.add_info_grid(container, [
            ("Descrição", self.product.get("descricao")),
            ("Categoria", self.product.get("categoria")),
            ("Tipo", self.product.get("tipo_produto")),
            ("Unidade", self.product.get("unidade")),
            ("Código de barras", self.product.get("codigo_barras")),
            ("Custo", f"R$ {format_money(self.product.get('valor_custo'))}"),
            ("Venda", f"R$ {format_money(self.product.get('valor_venda'))}"),
            ("Estoque atual", format_money(self.product.get("estoque_atual"))),
            ("Estoque mínimo", format_money(self.product.get("estoque_minimo"))),
            ("Status", self.product.get("status")),
            ("Observações", self.product.get("observacoes")),
        ])

        related = []
        for document in get_documents():
            items = get_document_items(document["id"])
            if any(item.get("product_code") == self.product.get("codigo") for item in items):
                related.append(document)

        columns = ["id", "numero", "tipo_documento", "parceiro", "data_emissao", "status", "total"]
        self.add_section_title(container, "Documentos relacionados")
        tree = TreeFrame(container, columns)
        tree.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        self.fill_tree(tree, related, columns)

class PartnerDetailWindow(DetailWindow):
    def __init__(self, master, codigo):
        self.partner = get_partner_by_code(codigo)

        if not self.partner:
            messagebox.showerror("Parceiro", "Parceiro não encontrado.")
            return

        super().__init__(master, f"Parceiro {codigo}", "920x650")
        self.build()

    def build(self):
        container = ttk.Frame(self, padding=14)
        container.pack(fill="both", expand=True)
        ttk.Label(container, text=self.partner.get("nome_razao", ""), font=("Segoe UI", 18, "bold"), style="Title.TLabel").pack(anchor="w")

        self.add_info_grid(container, [
            ("Código", self.partner.get("codigo")),
            ("Tipo de pessoa", self.partner.get("tipo_pessoa")),
            ("CPF/CNPJ", self.partner.get("documento")),
            ("Nome fantasia", self.partner.get("nome_fantasia")),
            ("Inscrição estadual", self.partner.get("inscricao_estadual")),
            ("Inscrição municipal", self.partner.get("inscricao_municipal")),
            ("Telefone", self.partner.get("telefone")),
            ("Celular", self.partner.get("celular")),
            ("E-mail", self.partner.get("email")),
            ("Cidade/UF", f"{self.partner.get('cidade', '')}/{self.partner.get('estado', '')}"),
            ("Tipo de parceiro", self.partner.get("tipo_parceiro")),
            ("Status", self.partner.get("status")),
            ("Observações", self.partner.get("observacoes")),
        ])

        rows = [row for row in get_documents() if row.get("partner_code") == self.partner.get("codigo")]
        columns = ["id", "numero", "tipo_documento", "direcao", "data_emissao", "status", "total"]
        self.add_section_title(container, "Compras, vendas e faturamentos")
        tree = TreeFrame(container, columns)
        tree.pack(fill="both", expand=True)
        self.fill_tree(tree, rows, columns)

class DocumentDetailWindow(DetailWindow):
    def __init__(self, master, document_id):
        self.document, self.items = get_document_by_id(document_id)

        if not self.document:
            messagebox.showerror("Documento", "Documento não encontrado.")
            return

        super().__init__(master, f"Documento {self.document.get('numero', '')}", "940x680")
        self.build()

    def build(self):
        container = ttk.Frame(self, padding=14)
        container.pack(fill="both", expand=True)
        top = ttk.Frame(container)
        top.pack(fill="x", pady=(0, 8))
        ttk.Label(top, text=f"{self.document.get('tipo_documento', '')} {self.document.get('numero', '')}", font=("Segoe UI", 18, "bold"), style="Title.TLabel").pack(side="left")
        ttk.Button(top, text="Gerar PDF", style="Primary.TButton", command=lambda: generate_document_pdf(self.document["id"])).pack(side="right", padx=4)

        self.add_info_grid(container, [
            ("Parceiro", self.document.get("parceiro")),
            ("Emissão", format_date_br(self.document.get("data_emissao"))),
            ("Validade", format_date_br(self.document.get("data_validade"))),
            ("Prazo entrega", format_date_br(self.document.get("prazo_entrega"))),
            ("Condição pagamento", self.document.get("condicao_pagamento")),
            ("Forma pagamento", self.document.get("forma_pagamento")),
            ("Status", self.document.get("status")),
            ("Status fiscal", self.document.get("fiscal_status")),
            ("Subtotal", f"R$ {format_money(self.document.get('subtotal'))}"),
            ("Desconto", f"R$ {format_money(self.document.get('desconto_total'))}"),
            ("Acréscimo", f"R$ {format_money(self.document.get('acrescimo_total'))}"),
            ("Total", f"R$ {format_money(self.document.get('total'))}"),
            ("Observações", self.document.get("observacoes")),
        ])

        columns = ["produto_codigo", "produto_nome", "quantidade", "unidade", "valor_unitario", "desconto", "acrescimo", "total", "ncm"]
        self.add_section_title(container, "Itens do documento")
        tree = TreeFrame(container, columns)
        tree.pack(fill="both", expand=True)
        self.fill_tree(tree, self.items, columns)

class LoginFrame(BaseFrame):
    def __init__(self, master, app):
        super().__init__(master, padding=20)
        self.app = app
        self.build()

    def build(self):
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)

        box = ttk.LabelFrame(self, text=" ", padding=22)
        box.grid(row=0, column=0)

        logo_box = ttk.Frame(box)
        logo_box.grid(row=0, column=0, pady=(0, 18), sticky="ew")

        logo_image = self.app.get_logo_image()

        if logo_image:
            ttk.Label(logo_box, image=logo_image).pack(pady=(0, 6))

        self.make_label(box, "Login ou e-mail:", 1, 0)
        self.login_entry = self.make_entry(box, 2, 0, width=40)

        self.make_label(box, "Senha:", 3, 0)
        self.password_entry = self.make_entry(box, 4, 0, width=40, show="*")

        login_button = ttk.Button(
            box,
            text="Entrar",
            style="Primary.TButton",
            command=self.do_login
        )
        login_button.grid(row=5, column=0, sticky="ew", pady=(8, 4))

        register_button = ttk.Button(
            box,
            text="Cadastrar-se",
            command=self.app.show_register
        )
        register_button.grid(row=6, column=0, sticky="ew", pady=4)

        info = ttk.Label(
            box,
            text="Faça o Login para entrar no App",
            foreground="gray"
        )
        info.grid(row=7, column=0, pady=(14, 0))

        self.login_entry.bind("<Return>", lambda event: self.do_login())
        self.password_entry.bind("<Return>", lambda event: self.do_login())

    def do_login(self):
        login = self.login_entry.get().strip()
        password = self.password_entry.get().strip()

        if not login or not password:
            messagebox.showwarning("Login", "Informe login e senha.")
            return

        user = authenticate_user(login, password)

        if user == "pending":
            messagebox.showwarning(
                "Cadastro pendente",
                "Seu cadastro ainda está pendente de aprovação."
            )
            return

        if user == "blocked":
            messagebox.showerror(
                "Cadastro bloqueado",
                "Seu usuário está bloqueado."
            )
            return

        if not user:
            messagebox.showerror("Login", "Login ou senha inválidos.")
            return

        self.app.show_main(user)

class RegisterFrame(BaseFrame):
    def __init__(self, master, app):
        super().__init__(master, padding=20)
        self.app = app
        self.build()

    def build(self):
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)

        scroll = ScrollableFrame(self)
        scroll.grid(row=0, column=0, sticky="nsew")
        container = scroll.scrollable_frame

        box = ttk.LabelFrame(container, text="Cadastro de novo usuário", padding=20)
        box.pack(padx=20, pady=20)

        title = ttk.Label(
            box,
            text="Solicitar acesso",
            font=("Segoe UI", 18, "bold")
        )
        title.grid(row=0, column=0, columnspan=2, pady=(0, 12))

        self.make_label(box, "Tipo:", 1, 0)
        self.tipo = self.make_combo(box, 2, 0, ["Jurídica", "Física"], width=40)

        self.make_label(box, "Nome ou razão social:", 3, 0)
        self.nome = self.make_entry(box, 4, 0, width=42)

        self.make_label(box, "CPF ou CNPJ:", 5, 0)
        self.doc = self.make_entry(box, 6, 0, width=42)

        self.make_label(box, "Telefone:", 7, 0)
        self.tel = self.make_entry(box, 8, 0, width=42)

        self.make_label(box, "E-mail:", 9, 0)
        self.email = self.make_entry(box, 10, 0, width=42)

        self.make_label(box, "Responsável:", 11, 0)
        self.responsavel = self.make_entry(box, 12, 0, width=42)

        self.make_label(box, "Endereço:", 13, 0)
        self.endereco = self.make_entry(box, 14, 0, width=42)

        self.make_label(box, "Login desejado:", 15, 0)
        self.username = self.make_entry(box, 16, 0, width=42)

        self.make_label(box, "Senha:", 17, 0)
        self.password = self.make_entry(box, 18, 0, width=42, show="*")

        self.make_label(box, "Confirmar senha:", 19, 0)
        self.confirm = self.make_entry(box, 20, 0, width=42, show="*")

        ttk.Button(
            box,
            text="Solicitar cadastro",
            style="Primary.TButton",
            command=self.create_account
        ).grid(row=21, column=0, sticky="ew", pady=(10, 4))

        ttk.Button(
            box,
            text="Voltar",
            command=self.app.show_login
        ).grid(row=22, column=0, sticky="ew", pady=4)

    def create_account(self):
        if self.password.get() != self.confirm.get():
            messagebox.showwarning("Cadastro", "As senhas não conferem.")
            return

        try:
            codigo = create_client_user(
                nome_razao=self.nome.get(),
                documento=self.doc.get(),
                telefone=self.tel.get(),
                email=self.email.get(),
                username=self.username.get(),
                password=self.password.get(),
                tipo_pessoa=self.tipo.get(),
                responsavel=self.responsavel.get(),
                endereco=self.endereco.get(),
                status="pending"
            )

            messagebox.showinfo(
                "Cadastro solicitado",
                f"Cadastro criado com sucesso.\nCódigo: {codigo}\n\nAguarde aprovação do administrador."
            )

            self.app.show_login()

        except Exception as error:
            messagebox.showerror("Cadastro", str(error))

class DashboardTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=12)
        self.alert_columns = ["servico", "nome_servico", "data", "funcionario_codigo", "funcionario", "item_codigo", "epi_faltando", "status"]
        self.validity_columns = ["origem", "responsavel", "item_codigo", "item_nome", "quantidade", "validade", "status", "nivel"]
        self.employee_item_columns = ["funcionario_codigo", "funcionario_nome", "item_codigo", "item_nome", "quantidade", "ultima_movimentacao", "validade"]
        self.build()

    def build(self):
        title = ttk.Label(
            self,
            text="Painel inicial",
            font=("Segoe UI", 16, "bold")
        )
        title.pack(anchor="w", pady=(0, 12))

        stats = self.get_stats()

        cards = ttk.Frame(self)
        cards.pack(fill="x", pady=(0, 12))

        for index, (name, value) in enumerate(stats.items()):
            card = ttk.LabelFrame(cards, text=name, padding=12)
            card.grid(row=index // 4, column=index % 4, padx=5, pady=5, sticky="ew")

            ttk.Label(
                card,
                text=str(value),
                font=("Segoe UI", 18, "bold")
            ).pack()

        for column in range(4):
            cards.columnconfigure(column, weight=1)

        alerts_box = ttk.LabelFrame(self, text="Alertas automáticos de EPI por serviço", padding=8)
        alerts_box.pack(fill="both", expand=True, pady=(0, 8))

        self.alert_tree = self.create_tree(alerts_box, self.alert_columns)
        self.alert_tree.pack(fill="both", expand=True)
        self.fill_tree(self.alert_tree, get_service_epi_alerts(), self.alert_columns)
        self.bind_double_click(
            self.alert_tree,
            self.alert_columns,
            lambda row: self.open_item_detail(row.get("item_codigo"))
        )

        validity_box = ttk.LabelFrame(self, text="Alertas de validade de itens e EPI", padding=8)
        validity_box.pack(fill="both", expand=True, pady=(0, 8))

        self.validity_tree = self.create_tree(validity_box, self.validity_columns)
        self.validity_tree.pack(fill="both", expand=True)
        self.fill_tree(self.validity_tree, get_validity_alerts(), self.validity_columns)
        self.bind_double_click(
            self.validity_tree,
            self.validity_columns,
            lambda row: self.open_item_detail(row.get("item_codigo"))
        )

        employee_items_box = ttk.LabelFrame(self, text="Itens atualmente com funcionários", padding=8)
        employee_items_box.pack(fill="both", expand=True)

        self.employee_items_tree = self.create_tree(employee_items_box, self.employee_item_columns)
        self.employee_items_tree.pack(fill="both", expand=True)
        self.fill_tree(self.employee_items_tree, get_employee_items(), self.employee_item_columns)
        self.bind_double_click(
            self.employee_items_tree,
            self.employee_item_columns,
            lambda row: self.open_item_detail(row.get("item_codigo"))
        )

    def get_stats(self):
        conn = get_conn()
        cur = conn.cursor()
        validity_alerts = get_validity_alerts()
        epi_alerts = get_service_epi_alerts()
        epi_expenses = get_epi_expense_report()
        total_epi_expense = sum(parse_float_br(row.get("valor_total", 0)) for row in epi_expenses)

        stats = {
            "Clientes": cur.execute("SELECT COUNT(*) qtd FROM clients").fetchone()["qtd"],
            "Funcionários": cur.execute("SELECT COUNT(*) qtd FROM employees").fetchone()["qtd"],
            "Itens": cur.execute("SELECT COUNT(*) qtd FROM items").fetchone()["qtd"],
            "Saldo estoque": cur.execute("SELECT COALESCE(SUM(quantidade), 0) qtd FROM movements").fetchone()["qtd"],
            "Movimentações": cur.execute("SELECT COUNT(*) qtd FROM movements").fetchone()["qtd"],
            "Serviços": cur.execute("SELECT COUNT(*) qtd FROM services").fetchone()["qtd"],
            "Itens vencidos": len([row for row in validity_alerts if row.get("nivel") == "VENCIDO"]),
            "Itens vencendo": len([row for row in validity_alerts if row.get("nivel") in ("URGENTE", "ATENCAO", "ATENÇÃO")]),
            "EPIs faltando": len(epi_alerts),
            "Pendentes": cur.execute("SELECT COUNT(*) qtd FROM users WHERE status = 'pending'").fetchone()["qtd"],
            "Gastos EPI": f"R$ {total_epi_expense:.2f}",
        }

        conn.close()
        return stats

class ApprovalTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.selected_user_id = None
        self.columns = [
            "id",
            "codigo",
            "nome_razao",
            "documento",
            "email",
            "telefone",
            "username",
            "status",
            "created_at"
        ]
        self.build()
        self.refresh()

    def build(self):
        top = ttk.Frame(self)
        top.pack(fill="x", pady=(0, 8))

        ttk.Label(
            top,
            text="Aprovação de usuários",
            font=("Segoe UI", 14, "bold")
        ).pack(side="left")

        ttk.Button(
            top,
            text="Atualizar",
            command=self.refresh
        ).pack(side="right")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)

        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.select_user)

        buttons = ttk.Frame(self)
        buttons.pack(fill="x", pady=8)

        ttk.Button(
            buttons,
            text="Aprovar selecionado",
            command=lambda: self.change_status("approved")
        ).pack(side="left", padx=4)

        ttk.Button(
            buttons,
            text="Bloquear selecionado",
            command=lambda: self.change_status("blocked")
        ).pack(side="left", padx=4)

    def refresh(self):
        rows = get_pending_users()

        for row in rows:
            row["documento"] = format_document(row.get("documento", ""))

        self.fill_tree(self.tree_frame, rows, self.columns)
        self.selected_user_id = None

    def select_user(self, event=None):
        selected = self.tree_frame.tree.selection()

        if not selected:
            self.selected_user_id = None
            return

        values = self.tree_frame.tree.item(selected[0], "values")
        self.selected_user_id = values[0]

    def change_status(self, status):
        if not self.selected_user_id:
            messagebox.showwarning("Usuário", "Selecione um usuário na tabela.")
            return

        update_user_status(self.selected_user_id, status)

        if status == "approved":
            messagebox.showinfo("Usuário", "Usuário aprovado com sucesso.")
        else:
            messagebox.showinfo("Usuário", "Usuário bloqueado com sucesso.")

        self.refresh()

class AccessControlTab(BaseFrame):
    def __init__(self, master, current_user):
        super().__init__(master, padding=10)
        self.current_user = current_user
        self.selected_user_id = None
        self.columns = [
            "id",
            "username",
            "email",
            "perfil",
            "status_acesso",
            "client_code",
            "nome_razao",
            "documento",
            "created_at"
        ]
        self.build()
        self.refresh()

    def build(self):
        top = ttk.Frame(self)
        top.pack(fill="x", pady=(0, 8))

        ttk.Label(
            top,
            text="Controle de acessos",
            font=("Segoe UI", 14, "bold")
        ).pack(side="left")

        ttk.Button(
            top,
            text="Atualizar",
            command=self.refresh
        ).pack(side="right")

        info = ttk.Label(
            self,
            text="Use esta tela para liberar acesso total ao sistema, manter acesso de cliente ou bloquear usuários. O perfil Administrador libera as mesmas abas do admin.",
            foreground="gray"
        )
        info.pack(anchor="w", pady=(0, 8))

        form = ttk.LabelFrame(self, text="Alterar perfil e status do usuário", padding=10)
        form.pack(fill="x", pady=(0, 8))

        self.make_label(form, "Perfil:", 0, 0)
        self.role_combo = self.make_combo(form, 1, 0, list(ROLE_LABELS.values()), width=28)

        self.make_label(form, "Status:", 0, 1)
        self.status_combo = self.make_combo(form, 1, 1, ["Aprovado", "Pendente", "Bloqueado"], width=28)

        ttk.Button(
            form,
            text="Salvar acesso selecionado",
            command=self.save_access
        ).grid(row=1, column=2, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Liberar acesso total",
            command=self.give_admin_access
        ).grid(row=1, column=3, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Manter como cliente",
            command=self.give_client_access
        ).grid(row=1, column=4, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Bloquear",
            command=self.block_access
        ).grid(row=1, column=5, padx=6, pady=4, sticky="ew")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)

        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.select_user)

    def refresh(self):
        self.fill_tree(self.tree_frame, get_users_access(), self.columns)
        self.selected_user_id = None

    def select_user(self, event=None):
        selected = self.tree_frame.tree.selection()

        if not selected:
            self.selected_user_id = None
            return

        values = self.tree_frame.tree.item(selected[0], "values")
        self.selected_user_id = values[0]

        perfil = values[3]
        status = values[4]

        if perfil:
            self.role_combo.set(perfil)

        if status:
            self.status_combo.set(status)

    def validate_selected_user(self, new_role, new_status):
        if not self.selected_user_id:
            raise ValueError("Selecione um usuário na tabela.")

        current_id = str(self.current_user.get("id", ""))
        selected_id = str(self.selected_user_id)

        if selected_id == current_id and (new_role != "admin" or new_status != "approved"):
            raise ValueError("Por segurança, não é permitido remover ou bloquear o próprio acesso administrativo.")

    def save_access(self):
        try:
            new_role = role_value(self.role_combo.get())
            new_status = status_value(self.status_combo.get())

            self.validate_selected_user(new_role, new_status)
            update_user_access(self.selected_user_id, new_role, new_status)

            messagebox.showinfo(
                "Acessos",
                "Acesso atualizado com sucesso. O usuário deve sair e entrar novamente para carregar as novas permissões."
            )

            self.refresh()

        except Exception as error:
            messagebox.showerror("Acessos", str(error))

    def give_admin_access(self):
        try:
            self.validate_selected_user("admin", "approved")
            update_user_access(self.selected_user_id, "admin", "approved")

            messagebox.showinfo(
                "Acessos",
                "Acesso total liberado. O usuário terá as mesmas abas do administrador ao entrar novamente."
            )

            self.refresh()

        except Exception as error:
            messagebox.showerror("Acessos", str(error))

    def give_client_access(self):
        try:
            self.validate_selected_user("client", "approved")
            update_user_access(self.selected_user_id, "client", "approved")

            messagebox.showinfo(
                "Acessos",
                "Usuário liberado como cliente."
            )

            self.refresh()

        except Exception as error:
            messagebox.showerror("Acessos", str(error))

    def block_access(self):
        try:
            self.validate_selected_user(role_value(self.role_combo.get()), "blocked")
            update_user_access(self.selected_user_id, role_value(self.role_combo.get()), "blocked")

            messagebox.showinfo(
                "Acessos",
                "Usuário bloqueado com sucesso."
            )

            self.refresh()

        except Exception as error:
            messagebox.showerror("Acessos", str(error))

class ClientsTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = [
            "codigo",
            "tipo_pessoa",
            "nome_razao",
            "documento",
            "telefone",
            "email",
            "responsavel",
            "username",
            "status"
        ]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text="Cadastrar cliente", padding=10)
        form.pack(fill="x", pady=(0, 8))

        form.columnconfigure(0, weight=1)
        form.columnconfigure(1, weight=1)
        form.columnconfigure(2, weight=1)
        form.columnconfigure(3, weight=1)

        self.make_label(form, "Tipo:", 0, 0)
        self.tipo = self.make_combo(form, 1, 0, ["Jurídica", "Física"])

        self.make_label(form, "Nome/Razão:", 0, 1)
        self.nome = self.make_entry(form, 1, 1)

        self.make_label(form, "CPF/CNPJ:", 0, 2)
        self.doc = self.make_entry(form, 1, 2)

        self.make_label(form, "Telefone:", 0, 3)
        self.tel = self.make_entry(form, 1, 3)

        self.make_label(form, "E-mail:", 2, 0)
        self.email = self.make_entry(form, 3, 0)

        self.make_label(form, "Responsável:", 2, 1)
        self.resp = self.make_entry(form, 3, 1)

        self.make_label(form, "Endereço:", 2, 2)
        self.end = self.make_entry(form, 3, 2)

        self.make_label(form, "Login:", 2, 3)
        self.user = self.make_entry(form, 3, 3)

        self.make_label(form, "Senha:", 4, 0)
        self.password = self.make_entry(form, 5, 0, show="*")

        ttk.Button(
            form,
            text="Salvar cliente",
            command=self.save
        ).grid(row=5, column=1, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Atualizar lista",
            command=self.refresh
        ).grid(row=5, column=2, padx=6, pady=4, sticky="ew")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_client_detail(row.get("codigo"))
        )

    def save(self):
        try:
            create_client_user(
                nome_razao=self.nome.get(),
                documento=self.doc.get(),
                telefone=self.tel.get(),
                email=self.email.get(),
                username=self.user.get(),
                password=self.password.get(),
                tipo_pessoa=self.tipo.get(),
                responsavel=self.resp.get(),
                endereco=self.end.get(),
                status="approved"
            )

            messagebox.showinfo("Clientes", "Cliente cadastrado com sucesso.")

            self.clear_entries(
                self.nome,
                self.doc,
                self.tel,
                self.email,
                self.resp,
                self.end,
                self.user,
                self.password
            )

            self.refresh()

        except Exception as error:
            messagebox.showerror("Clientes", str(error))

    def refresh(self):
        self.fill_tree(self.tree_frame, get_clients(), self.columns)

class EmployeesTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = [
            "codigo",
            "nome",
            "documento",
            "cargo",
            "telefone",
            "email",
            "status",
            "observacoes"
        ]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text="Cadastro de funcionário", padding=10)
        form.pack(fill="x", pady=(0, 8))

        self.make_label(form, "Nome:", 0, 0)
        self.nome = self.make_entry(form, 1, 0)

        self.make_label(form, "CPF/Documento:", 0, 1)
        self.documento = self.make_entry(form, 1, 1)

        self.make_label(form, "Cargo/Função:", 0, 2)
        self.cargo = self.make_entry(form, 1, 2)

        self.make_label(form, "Telefone:", 0, 3)
        self.telefone = self.make_entry(form, 1, 3)

        self.make_label(form, "E-mail:", 2, 0)
        self.email = self.make_entry(form, 3, 0)

        self.make_label(form, "Status:", 2, 1)
        self.status = self.make_combo(form, 3, 1, ["Ativo", "Inativo"])

        self.make_label(form, "Observações:", 2, 2)
        self.obs = self.make_entry(form, 3, 2, width=50)

        ttk.Button(
            form,
            text="Salvar funcionário",
            command=self.save
        ).grid(row=3, column=3, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Atualizar lista",
            command=self.refresh
        ).grid(row=3, column=4, padx=6, pady=4, sticky="ew")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_employee_detail(row.get("codigo"))
        )

    def save(self):
        try:
            codigo = create_employee(
                self.nome.get(),
                self.documento.get(),
                self.cargo.get(),
                self.telefone.get(),
                self.email.get(),
                self.status.get(),
                self.obs.get()
            )

            messagebox.showinfo("Funcionários", f"Funcionário cadastrado com sucesso.\nCódigo: {codigo}")

            self.clear_entries(
                self.nome,
                self.documento,
                self.cargo,
                self.telefone,
                self.email,
                self.obs
            )

            self.refresh()

        except Exception as error:
            messagebox.showerror("Funcionários", str(error))

    def refresh(self):
        self.fill_tree(self.tree_frame, get_employees(), self.columns)

class ItemsTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.selected_image_source = ""
        self.item_image_preview = None
        self.columns = [
            "codigo",
            "nome",
            "categoria",
            "unidade",
            "saldo",
            "ca",
            "valor_unitario",
            "validade",
            "imagem",
            "observacoes"
        ]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text="Cadastro de item", padding=10)
        form.pack(fill="x", pady=(0, 8))

        self.make_label(form, "Nome:", 0, 0)
        self.nome = self.make_entry(form, 1, 0)

        self.make_label(form, "Categoria:", 0, 1)
        self.categoria = self.make_combo(
            form,
            1,
            1,
            ["EPI", "Ferramenta", "Material", "Uniforme", "Peça", "Outro"]
        )

        self.make_label(form, "Unidade:", 0, 2)
        self.unidade = self.make_combo(
            form,
            1,
            2,
            ["un", "par", "caixa", "kit", "m", "kg", "litro"]
        )

        self.make_label(form, "CA:", 0, 3)
        self.ca = self.make_entry(form, 1, 3)

        self.make_label(form, "Valor unitário:", 2, 0)
        self.valor_unitario = self.make_entry(form, 3, 0)

        self.make_label(form, "Observações:", 2, 1)
        self.obs = self.make_entry(form, 3, 1, width=60)

        self.make_label(form, "Imagem:", 2, 2)
        image_actions = ttk.Frame(form)
        image_actions.grid(row=3, column=2, padx=6, pady=(0, 6), sticky="ew")

        ttk.Button(
            image_actions,
            text="Selecionar",
            command=self.choose_image
        ).pack(side="left", padx=(0, 4))

        ttk.Button(
            image_actions,
            text="Limpar",
            command=self.clear_selected_image
        ).pack(side="left")

        self.image_preview_label = ttk.Label(
            form,
            text="Sem imagem",
            style="Muted.TLabel"
        )
        self.image_preview_label.grid(row=3, column=3, padx=6, pady=(0, 6), sticky="w")

        ttk.Button(
            form,
            text="Salvar item",
            command=self.save
        ).grid(row=5, column=0, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Atualizar lista",
            command=self.refresh
        ).grid(row=5, column=1, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Abrir detalhe",
            command=self.open_selected_item_detail
        ).grid(row=5, column=2, padx=6, pady=4, sticky="ew")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_item_detail(row.get("codigo"))
        )

    def choose_image(self):
        path = filedialog.askopenfilename(
            title="Selecionar imagem do item",
            filetypes=[
                ("Imagens", "*.png;*.jpg;*.jpeg"),
                ("PNG", "*.png"),
                ("JPG", "*.jpg;*.jpeg"),
                ("Todos os arquivos", "*.*"),
            ]
        )

        if not path:
            return

        self.selected_image_source = path
        self.update_image_preview(path)

    def update_image_preview(self, path):
        self.item_image_preview = load_tk_image(path, (72, 72))

        if self.item_image_preview:
            self.image_preview_label.configure(image=self.item_image_preview, text="")
        else:
            self.image_preview_label.configure(image="", text=Path(path).name)

    def clear_selected_image(self):
        self.selected_image_source = ""
        self.item_image_preview = None
        self.image_preview_label.configure(image="", text="Sem imagem")

    def save(self):
        nome = self.nome.get().strip()

        if not nome:
            messagebox.showwarning("Itens", "Informe o nome do item.")
            return

        try:
            valor_unitario = parse_float_br(self.valor_unitario.get())
            codigo = get_next_code("items", "codigo", "ITM")
            image_path = copy_item_image(self.selected_image_source, codigo) if self.selected_image_source else ""

            conn = get_conn()
            conn.execute("""
                INSERT INTO items (
                    codigo,
                    nome,
                    categoria,
                    unidade,
                    ca,
                    valor_unitario,
                    image_path,
                    observacoes,
                    created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                codigo,
                nome,
                self.categoria.get(),
                self.unidade.get(),
                self.ca.get().strip(),
                valor_unitario,
                image_path,
                self.obs.get().strip(),
                now_iso()
            ))

            conn.commit()
            conn.close()

            messagebox.showinfo("Itens", f"Item cadastrado com sucesso.\nCódigo: {codigo}")

            self.clear_entries(self.nome, self.ca, self.valor_unitario, self.obs)
            self.clear_selected_image()
            self.refresh()

        except Exception as error:
            messagebox.showerror("Itens", str(error))

    def refresh(self):
        self.fill_tree(self.tree_frame, get_items(), self.columns)

    def open_selected_item_detail(self):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            messagebox.showwarning("Itens", "Selecione um item na tabela.")
            return

        self.open_item_detail(row.get("codigo"))

class MovementsTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = [
            "item_codigo",
            "item_nome",
            "item_categoria",
            "item_unidade",
            "tipo",
            "quantidade",
            "data_movimentacao",
            "validade",
            "funcionario_nome",
            "service_codigo",
            "destino",
            "observacoes"
        ]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text="Registrar movimentação", padding=10)
        form.pack(fill="x", pady=(0, 8))

        self.make_label(form, "Item:", 0, 0)
        self.item = self.make_combo(form, 1, 0, get_item_options(), width=45)

        self.make_label(form, "Tipo:", 0, 1)
        self.tipo = self.make_combo(form, 1, 1, ["Entrada", "Saída"])

        self.make_label(form, "Quantidade:", 0, 2)
        self.quantidade = self.make_entry(form, 1, 2)

        self.make_label(form, "Data:", 0, 3)
        self.data = self.make_date_entry(form, 1, 3)
        self.data.insert(0, today_br())

        self.make_label(form, "Validade:", 2, 0)
        self.validade = self.make_date_entry(form, 3, 0)

        self.make_label(form, "Funcionário:", 2, 1)
        self.funcionario = self.make_combo(form, 3, 1, get_employee_options(), width=45)

        self.make_label(form, "Serviço:", 2, 2)
        self.servico = self.make_combo(form, 3, 2, get_service_options(), width=45)

        self.make_label(form, "Destino:", 4, 0)
        self.destino = self.make_entry(form, 5, 0)

        self.make_label(form, "Observações:", 4, 1)
        self.obs = self.make_entry(form, 5, 1, width=50)

        ttk.Button(
            form,
            text="Registrar",
            command=self.save
        ).grid(row=5, column=2, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Atualizar",
            command=self.refresh
        ).grid(row=5, column=3, padx=6, pady=4, sticky="ew")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_item_detail(row.get("item_codigo"))
        )

    def refresh_options(self):
        item_values = get_item_options()
        self.item["values"] = item_values
        self.item.set(item_values[0] if item_values else "")

        employee_values = get_employee_options()
        self.funcionario["values"] = employee_values
        self.funcionario.set(employee_values[0] if employee_values else "")

        service_values = get_service_options()
        self.servico["values"] = service_values
        self.servico.set(service_values[0] if service_values else "")

    def save(self):
        item_text = self.item.get().strip()

        if not item_text:
            messagebox.showwarning("Movimentação", "Selecione um item.")
            return

        item_codigo = item_text.split(" - ")[0]

        conn = get_conn()
        item = conn.execute(
            "SELECT * FROM items WHERE codigo = ?",
            (item_codigo,)
        ).fetchone()

        if not item:
            conn.close()
            messagebox.showerror("Movimentação", "Item não encontrado.")
            return

        try:
            quantidade = int(self.quantidade.get())

            if quantidade <= 0:
                raise ValueError("Quantidade deve ser maior que zero.")

            data_mov = validate_date(self.data.get(), "data", allow_empty=False)
            validade = validate_date(self.validade.get(), "validade", allow_empty=True)

            funcionario_codigo = ""
            funcionario_nome = ""

            funcionario_text = self.funcionario.get().strip()

            if funcionario_text:
                funcionario_codigo = funcionario_text.split(" - ")[0]
                employee = get_employee_by_code(funcionario_codigo)

                if employee:
                    funcionario_nome = employee["nome"]

            service_codigo = ""
            service_text = self.servico.get().strip()

            if service_text:
                service_codigo = service_text.split(" - ")[0]

            if self.tipo.get() == "Saída":
                saldo = get_item_balance(item_codigo)

                if saldo < quantidade:
                    raise ValueError("Saldo insuficiente para saída.")

                if not funcionario_codigo:
                    raise ValueError("Para saída de EPI/material, selecione o funcionário que receberá o item.")

                quantidade = quantidade * -1

            conn.execute("""
                INSERT INTO movements (
                    item_codigo,
                    item_nome,
                    item_categoria,
                    item_unidade,
                    tipo,
                    quantidade,
                    data_movimentacao,
                    validade,
                    funcionario_codigo,
                    funcionario_nome,
                    service_codigo,
                    destino,
                    observacoes,
                    created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                item["codigo"],
                item["nome"],
                item["categoria"],
                item["unidade"],
                self.tipo.get(),
                quantidade,
                data_mov,
                validade,
                funcionario_codigo,
                funcionario_nome,
                service_codigo,
                self.destino.get().strip(),
                self.obs.get().strip(),
                now_iso()
            ))

            conn.commit()
            conn.close()

            messagebox.showinfo("Movimentação", "Movimentação registrada com sucesso.")

            self.clear_entries(
                self.quantidade,
                self.validade,
                self.destino,
                self.obs
            )

            self.data.delete(0, "end")
            self.data.insert(0, today_br())

            self.refresh()

        except Exception as error:
            conn.close()
            messagebox.showerror("Movimentação", str(error))

    def refresh(self):
        self.refresh_options()
        self.fill_tree(self.tree_frame, get_movements(), self.columns)

class ServicesTab(BaseFrame):
    def __init__(self, master, user):
        super().__init__(master, padding=10)
        self.user = user
        self.columns = [
            "codigo_servico",
            "nome_razao",
            "nome_servico",
            "descricao",
            "valor",
            "prazo",
            "status",
            "observacoes"
        ]
        self.employee_columns = ["codigo", "nome"]
        self.required_columns = ["codigo", "nome"]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text="Serviços", padding=10)
        form.pack(fill="x", pady=(0, 8))

        self.make_label(form, "Cliente:", 0, 0)
        self.cliente = self.make_combo(form, 1, 0, get_client_options(), width=45)

        self.make_label(form, "Nome do serviço:", 0, 1)
        self.nome = self.make_entry(form, 1, 1)

        self.make_label(form, "Valor:", 0, 2)
        self.valor = self.make_entry(form, 1, 2)

        self.make_label(form, "Prazo/Data:", 0, 3)
        self.prazo = self.make_date_entry(form, 1, 3)

        self.make_label(form, "Status:", 2, 0)
        self.status = self.make_combo(
            form,
            3,
            0,
            ["Pendente", "Em andamento", "Concluído", "Cancelado"]
        )

        self.make_label(form, "Descrição:", 2, 1)
        self.descricao = self.make_entry(form, 3, 1, width=45)

        self.make_label(form, "Observações:", 2, 2)
        self.obs = self.make_entry(form, 3, 2, width=45)

        if self.user.get("role") == "admin":
            ttk.Button(
                form,
                text="Salvar serviço",
                command=self.save
            ).grid(row=3, column=3, padx=6, pady=4, sticky="ew")

        ttk.Button(
            form,
            text="Atualizar",
            command=self.refresh
        ).grid(row=3, column=4, padx=6, pady=4, sticky="ew")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True, pady=(0, 8))
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_service_detail(row.get("codigo_servico"))
        )

        if self.user.get("role") == "admin":
            link_box = ttk.LabelFrame(self, text="Vincular funcionários e EPIs ao serviço", padding=10)
            link_box.pack(fill="x", pady=(0, 8))

            self.make_label(link_box, "Serviço:", 0, 0)
            self.service_link = self.make_combo(link_box, 1, 0, get_service_options(), width=45)

            self.make_label(link_box, "Funcionário:", 0, 1)
            self.employee_link = self.make_combo(link_box, 1, 1, get_employee_options(), width=45)

            ttk.Button(
                link_box,
                text="Adicionar funcionário ao serviço",
                command=self.add_employee_to_service
            ).grid(row=1, column=2, padx=6, sticky="ew")

            self.make_label(link_box, "EPI obrigatório:", 2, 0)
            self.required_item_link = self.make_combo(link_box, 3, 0, get_epi_options(), width=45)

            ttk.Button(
                link_box,
                text="Adicionar EPI obrigatório",
                command=self.add_required_item_to_service
            ).grid(row=3, column=1, padx=6, sticky="ew")

    def save(self):
        try:
            cliente_text = self.cliente.get().strip()

            if not cliente_text:
                raise ValueError("Selecione um cliente.")

            cliente_codigo = cliente_text.split(" - ")[0]

            if not self.nome.get().strip():
                raise ValueError("Informe o nome do serviço.")

            codigo_servico = get_next_code("services", "codigo_servico", "SRV")
            prazo = validate_date(self.prazo.get(), "prazo", allow_empty=True)
            valor = parse_float_br(self.valor.get())

            conn = get_conn()
            conn.execute("""
                INSERT INTO services (
                    codigo_servico,
                    cliente_codigo,
                    nome_servico,
                    descricao,
                    valor,
                    prazo,
                    status,
                    observacoes,
                    created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                codigo_servico,
                cliente_codigo,
                self.nome.get().strip(),
                self.descricao.get().strip(),
                valor,
                prazo,
                self.status.get(),
                self.obs.get().strip(),
                now_iso()
            ))

            conn.commit()
            conn.close()

            messagebox.showinfo("Serviços", f"Serviço cadastrado com sucesso.\nCódigo: {codigo_servico}")

            self.clear_entries(
                self.nome,
                self.valor,
                self.prazo,
                self.descricao,
                self.obs
            )

            self.refresh()

        except Exception as error:
            messagebox.showerror("Serviços", str(error))

    def add_employee_to_service(self):
        try:
            service_text = self.service_link.get().strip()
            employee_text = self.employee_link.get().strip()

            codigo_servico = service_text.split(" - ")[0] if service_text else ""
            funcionario_codigo = employee_text.split(" - ")[0] if employee_text else ""

            link_employee_to_service(codigo_servico, funcionario_codigo)

            messagebox.showinfo("Serviços", "Funcionário vinculado ao serviço.")

        except Exception as error:
            messagebox.showerror("Serviços", str(error))

    def add_required_item_to_service(self):
        try:
            service_text = self.service_link.get().strip()
            item_text = self.required_item_link.get().strip()

            codigo_servico = service_text.split(" - ")[0] if service_text else ""
            item_codigo = item_text.split(" - ")[0] if item_text else ""

            link_required_item_to_service(codigo_servico, item_codigo)

            messagebox.showinfo("Serviços", "EPI obrigatório vinculado ao serviço.")

        except Exception as error:
            messagebox.showerror("Serviços", str(error))

    def refresh(self):
        client_code = self.user.get("client_code") if self.user.get("role") == "client" else None
        self.fill_tree(self.tree_frame, get_services(client_code), self.columns)

        if self.user.get("role") == "admin" and hasattr(self, "service_link"):
            service_values = get_service_options()
            self.service_link["values"] = service_values
            self.service_link.set(service_values[0] if service_values else "")

            employee_values = get_employee_options()
            self.employee_link["values"] = employee_values
            self.employee_link.set(employee_values[0] if employee_values else "")

            required_values = get_epi_options()
            self.required_item_link["values"] = required_values
            self.required_item_link.set(required_values[0] if required_values else "")

class HistoryTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = [
            "item_codigo",
            "item_nome",
            "item_categoria",
            "item_unidade",
            "tipo",
            "quantidade",
            "data_movimentacao",
            "validade",
            "funcionario_nome",
            "service_codigo",
            "destino",
            "observacoes"
        ]
        self.build()
        self.refresh()

    def build(self):
        top = ttk.Frame(self)
        top.pack(fill="x", pady=(0, 8))

        ttk.Label(
            top,
            text="Histórico de movimentações",
            font=("Segoe UI", 14, "bold")
        ).pack(side="left")

        ttk.Button(
            top,
            text="Exportar CSV",
            command=self.export
        ).pack(side="right", padx=4)

        ttk.Button(
            top,
            text="Atualizar",
            command=self.refresh
        ).pack(side="right", padx=4)

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_item_detail(row.get("item_codigo"))
        )

    def refresh(self):
        self.fill_tree(self.tree_frame, get_movements(), self.columns)

    def export(self):
        export_csv("historico_movimentacoes", self.columns, get_movements())

class EpiExpenseReportTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = [
            "funcionario_codigo",
            "funcionario_nome",
            "item_codigo",
            "item_nome",
            "quantidade",
            "valor_unitario",
            "valor_total",
            "data_movimentacao",
            "service_codigo",
            "destino",
            "observacoes"
        ]
        self.build()
        self.refresh()

    def build(self):
        top = ttk.Frame(self)
        top.pack(fill="x", pady=(0, 8))

        ttk.Label(
            top,
            text="Relatório de gastos de EPI por funcionário",
            font=("Segoe UI", 14, "bold")
        ).pack(side="left")

        ttk.Button(
            top,
            text="Exportar CSV",
            command=self.export
        ).pack(side="right", padx=4)

        ttk.Button(
            top,
            text="Atualizar",
            command=self.refresh
        ).pack(side="right", padx=4)

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_item_detail(row.get("item_codigo"))
        )

    def refresh(self):
        self.fill_tree(self.tree_frame, get_epi_expense_report(), self.columns)

    def export(self):
        export_csv("gastos_epi_por_funcionario", self.columns, get_epi_expense_report())

class ValidityAlertsTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = [
            "origem",
            "responsavel",
            "item_codigo",
            "item_nome",
            "quantidade",
            "validade",
            "status",
            "nivel"
        ]
        self.build()
        self.refresh()

    def build(self):
        top = ttk.Frame(self)
        top.pack(fill="x", pady=(0, 8))

        ttk.Label(
            top,
            text="Alertas de validade de itens e EPI",
            font=("Segoe UI", 14, "bold")
        ).pack(side="left")

        ttk.Button(
            top,
            text="Exportar CSV",
            command=self.export
        ).pack(side="right", padx=4)

        ttk.Button(
            top,
            text="Atualizar",
            command=self.refresh
        ).pack(side="right", padx=4)

        info = ttk.Label(
            self,
            text="O sistema mostra itens vencidos, itens que vencem hoje e itens que vencem em até 30 dias.",
            foreground="gray"
        )
        info.pack(anchor="w", pady=(0, 8))

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_item_detail(row.get("item_codigo"))
        )

    def refresh(self):
        self.fill_tree(self.tree_frame, get_validity_alerts(), self.columns)

    def export(self):
        export_csv("alertas_validade_itens_epi", self.columns, get_validity_alerts())

class CompaniesTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = [
            "codigo",
            "nome_razao",
            "documento",
            "telefone",
            "email",
            "responsavel",
            "contato",
            "endereco",
            "observacoes",
            "status"
        ]
        self.build()
        self.refresh()

    def build(self):
        top = ttk.Frame(self)
        top.pack(fill="x", pady=(0, 8))

        ttk.Label(
            top,
            text="Empresas",
            font=("Segoe UI", 14, "bold")
        ).pack(side="left")

        ttk.Button(
            top,
            text="Atualizar",
            command=self.refresh
        ).pack(side="right", padx=4)

        ttk.Label(
            self,
            text="Estrutura preparada a partir dos clientes pessoa juridica, com espaco para contatos, documentos, endereco e historico de servicos.",
            style="Muted.TLabel",
            wraplength=920
        ).pack(anchor="w", pady=(0, 8))

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.bind_double_click(
            self.tree_frame,
            self.columns,
            lambda row: self.open_client_detail(row.get("codigo"))
        )

    def refresh(self):
        self.fill_tree(self.tree_frame, get_companies(), self.columns)

class ReportsTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = ["relatorio", "status", "exportacao", "proxima_fase"]
        self.build()

    def build(self):
        top = ttk.Frame(self)
        top.pack(fill="x", pady=(0, 8))

        ttk.Label(
            top,
            text="Relatorios gerenciais",
            font=("Segoe UI", 14, "bold")
        ).pack(side="left")

        buttons = ttk.Frame(self)
        buttons.pack(fill="x", pady=(0, 8))

        ttk.Button(buttons, text="Exportar estoque critico", command=self.export_critical_stock).pack(side="left", padx=4)
        ttk.Button(buttons, text="Exportar historico", command=self.export_history).pack(side="left", padx=4)
        ttk.Button(buttons, text="Exportar gastos EPI", command=self.export_epi).pack(side="left", padx=4)
        ttk.Button(buttons, text="Exportar validades", command=self.export_validity).pack(side="left", padx=4)

        rows = [
            {"relatorio": "Estoque critico", "status": "Disponivel", "exportacao": "CSV", "proxima_fase": "Parametros por saldo minimo"},
            {"relatorio": "Itens vencidos", "status": "Disponivel", "exportacao": "CSV", "proxima_fase": "Filtros por periodo"},
            {"relatorio": "Itens vencendo", "status": "Disponivel", "exportacao": "CSV", "proxima_fase": "Filtros por janela de dias"},
            {"relatorio": "EPIs por funcionario", "status": "Disponivel", "exportacao": "CSV", "proxima_fase": "Tela analitica"},
            {"relatorio": "Gastos por funcionario", "status": "Disponivel", "exportacao": "CSV", "proxima_fase": "Exportacao Excel"},
            {"relatorio": "Gastos por periodo", "status": "Preparado", "exportacao": "CSV futuro", "proxima_fase": "Filtro de datas"},
            {"relatorio": "Gastos por servico", "status": "Preparado", "exportacao": "CSV futuro", "proxima_fase": "Agrupamento por servico"},
            {"relatorio": "Consumo por cliente", "status": "Preparado", "exportacao": "CSV futuro", "proxima_fase": "Agrupamento por cliente"},
            {"relatorio": "Servicos pendentes", "status": "Preparado", "exportacao": "CSV futuro", "proxima_fase": "Filtro de status"},
            {"relatorio": "Produtividade operacional", "status": "Planejado", "exportacao": "Excel futuro", "proxima_fase": "Indicadores operacionais"},
        ]

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.fill_tree(self.tree_frame, rows, self.columns)

    def export_critical_stock(self):
        columns = ["codigo", "nome", "categoria", "unidade", "saldo", "ca", "valor_unitario", "validade", "status_estoque", "observacoes"]
        export_csv("estoque_critico", columns, get_critical_stock_report())

    def export_history(self):
        columns = ["item_codigo", "item_nome", "item_categoria", "item_unidade", "tipo", "quantidade", "data_movimentacao", "validade", "funcionario_nome", "service_codigo", "destino", "observacoes"]
        export_csv("historico_movimentacoes", columns, get_movements())

    def export_epi(self):
        columns = ["funcionario_codigo", "funcionario_nome", "item_codigo", "item_nome", "quantidade", "valor_unitario", "valor_total", "data_movimentacao", "service_codigo", "destino", "observacoes"]
        export_csv("gastos_epi_por_funcionario", columns, get_epi_expense_report())

    def export_validity(self):
        columns = ["origem", "responsavel", "item_codigo", "item_nome", "quantidade", "validade", "status", "nivel"]
        export_csv("alertas_validade_itens_epi", columns, get_validity_alerts())

class PartnerLookupMixin:
    def setup_partner_lookup(self, combo, info_label=None, kind=None):
        self.partner = combo
        self.partner_info_label = info_label
        self.partner_kind = kind
        self.partner_options = get_partner_options(kind)
        self.partner.configure(state="normal")
        self.partner["values"] = self.partner_options
        self.partner.bind("<KeyRelease>", self.filter_partner_options)
        self.partner.bind("<<ComboboxSelected>>", lambda _event: self.load_partner_details())
        self.partner.bind("<Return>", lambda _event: self.load_partner_details())
        self.partner.bind("<FocusOut>", lambda _event: self.load_partner_details())

    def filter_partner_options(self, event=None):
        if event and event.keysym in ("Return", "Escape", "Up", "Down", "Left", "Right", "Tab"):
            return

        typed = self.partner.get().strip().lower()
        options = self.partner_options or get_partner_options(getattr(self, "partner_kind", None))
        self.partner["values"] = [option for option in options if typed in option.lower()] if typed else options

    def refresh_partner_options(self):
        current = self.partner.get().strip()
        self.partner_options = get_partner_options(getattr(self, "partner_kind", None))
        self.partner["values"] = self.partner_options

        if not current and self.partner_options:
            self.partner.set(self.partner_options[0])

    def resolve_selected_partner(self):
        partner = resolve_partner_text(self.partner.get(), getattr(self, "partner_kind", None))

        if partner:
            option = format_partner_option(partner)
            if option:
                self.partner.set(option)

        return partner

    def load_partner_details(self):
        if not getattr(self, "partner_info_label", None):
            return

        partner = self.resolve_selected_partner()

        if not partner:
            self.partner_info_label.configure(text="")
            return

        details = [
            partner.get("nome_razao", ""),
            f"CPF/CNPJ: {partner.get('documento', '')}",
            f"Tipo: {partner.get('tipo_parceiro', '')}",
            f"Cidade: {partner.get('cidade', '')}/{partner.get('estado', '')}",
            f"Telefone: {partner.get('telefone', '')}",
            f"E-mail: {partner.get('email', '')}",
        ]
        self.partner_info_label.configure(text=" | ".join([detail for detail in details if detail and not detail.endswith(": ")]))

class AttachmentActionsMixin:
    def init_attachment_state(self):
        self.selected_attachment_source = ""
        self.current_attachment = ""

    def choose_attachment(self):
        path = filedialog.askopenfilename(
            title="Selecionar anexo",
            filetypes=[("Arquivos", "*.pdf;*.png;*.jpg;*.jpeg;*.xlsx;*.xls;*.docx;*.xml"), ("Todos os arquivos", "*.*")]
        )

        if path:
            self.selected_attachment_source = path
            self.set_attachment_label()

    def clear_attachment(self):
        self.selected_attachment_source = ""
        self.current_attachment = ""
        self.set_attachment_label()

    def set_attachment_label(self):
        label = getattr(self, "attachment_label", None)

        if not label:
            return

        path = self.selected_attachment_source or self.current_attachment
        label.configure(text=Path(path).name if path else "Sem anexo")

class AuxiliaryRegistersTab(AttachmentActionsMixin, PartnerLookupMixin, BaseFrame):
    def __init__(self, master, current_user=None):
        super().__init__(master, padding=10)
        self.current_user = current_user or {}
        self.selected_record_id = None
        self.init_attachment_state()
        self.columns = ["id", "codigo", "categoria", "nome", "parceiro", "cidade", "valor", "status", "descricao", "anexo"]
        self.build()
        self.refresh()

    def build(self):
        ttk.Label(self, text="Cadastros auxiliares", font=("Segoe UI", 14, "bold"), style="Title.TLabel").pack(anchor="w", pady=(0, 8))
        form = ttk.LabelFrame(self, text="Cadastro geral", padding=10)
        form.pack(fill="x", pady=(0, 8))

        for column in range(4):
            form.columnconfigure(column, weight=1, uniform="aux_form")

        self.make_label(form, "Categoria:", 0, 0)
        self.categoria = self.make_combo(form, 1, 0, AUXILIARY_CATEGORIES, width=30)
        self.make_label(form, "Nome:", 0, 1)
        self.nome = self.make_entry(form, 1, 1, width=34)
        self.make_label(form, "Parceiro vinculado:", 0, 2)
        partner = self.make_combo(form, 1, 2, get_partner_options(), width=36)
        self.make_label(form, "Cidade/local:", 0, 3)
        self.cidade = self.make_entry(form, 1, 3, width=22)
        self.make_label(form, "Valor:", 2, 0)
        self.valor = self.make_entry(form, 3, 0, width=16)
        self.make_label(form, "Status:", 2, 1)
        self.status = self.make_combo(form, 3, 1, ["Ativo", "Inativo", "Pendente"], width=18)
        self.make_label(form, "Descrição:", 2, 2)
        self.descricao = self.make_entry(form, 3, 2, width=40)
        self.make_label(form, "Anexo:", 2, 3)
        attachment = ttk.Frame(form)
        attachment.grid(row=3, column=3, padx=6, pady=(0, 6), sticky="ew")
        ttk.Button(attachment, text="Anexar", command=self.choose_attachment).pack(side="left", padx=(0, 4))
        ttk.Button(attachment, text="Limpar", command=self.clear_attachment).pack(side="left")
        self.attachment_label = ttk.Label(form, text="Sem anexo", style="Muted.TLabel")
        self.attachment_label.grid(row=4, column=3, padx=6, pady=(0, 6), sticky="w")
        self.partner_info_label = ttk.Label(self, text="", style="Muted.TLabel", wraplength=980)
        self.partner_info_label.pack(anchor="w", pady=(0, 6))
        self.setup_partner_lookup(partner, self.partner_info_label)

        self.save_button = ttk.Button(form, text="Salvar cadastro", style="Primary.TButton", command=self.save)
        self.save_button.grid(row=5, column=0, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Novo", command=self.clear_form).grid(row=5, column=1, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Exportar XLSX", command=self.export_xlsx).grid(row=5, column=2, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Gerar PDF", command=self.export_pdf).grid(row=5, column=3, padx=6, pady=4, sticky="ew")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.load_selected_record)

    def clear_form(self):
        self.selected_record_id = None
        self.clear_entries(self.nome, self.cidade, self.valor, self.descricao)
        self.status.set("Ativo")
        if AUXILIARY_CATEGORIES:
            self.categoria.set(AUXILIARY_CATEGORIES[0])
        self.refresh_partner_options()
        self.partner_info_label.configure(text="")
        self.init_attachment_state()
        self.set_attachment_label()
        self.save_button.configure(text="Salvar cadastro")

    def load_selected_record(self, event=None):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            return

        record = get_auxiliary_record_by_id(row.get("id"))

        if not record:
            return

        partner = get_partner_by_code(record.get("partner_code"))
        self.selected_record_id = record.get("id")
        self.set_combo(self.categoria, record.get("categoria"))
        self.set_entry(self.nome, record.get("nome"))
        self.set_combo(self.partner, format_partner_option(partner))
        self.set_entry(self.cidade, record.get("cidade"))
        self.set_entry(self.valor, format_money(record.get("valor")))
        self.set_combo(self.status, record.get("status"))
        self.set_entry(self.descricao, record.get("descricao"))
        self.selected_attachment_source = ""
        self.current_attachment = record.get("attachment_path") or ""
        self.set_attachment_label()
        self.load_partner_details()
        self.save_button.configure(text="Atualizar cadastro")

    def save(self):
        try:
            partner = self.resolve_selected_partner()
            partner_code = partner.get("codigo") if partner else ""

            if self.selected_record_id:
                record_id = update_auxiliary_record(
                    self.selected_record_id, self.categoria.get(), self.nome.get(), partner_code, self.cidade.get(),
                    self.valor.get(), self.status.get(), self.descricao.get(), self.selected_attachment_source, self.current_attachment
                )
                message = "Cadastro atualizado com sucesso."
            else:
                record_id = create_auxiliary_record(
                    self.categoria.get(), self.nome.get(), partner_code, self.cidade.get(), self.valor.get(),
                    self.status.get(), self.descricao.get(), self.selected_attachment_source
                )
                message = f"Cadastro criado com sucesso.\nCódigo: {record_id}"

            log_audit(self.current_user.get("id"), "salvar", "auxiliary_records", str(record_id), self.nome.get())
            messagebox.showinfo("Cadastros auxiliares", message)
            self.clear_form()
            self.refresh()
        except Exception as error:
            messagebox.showerror("Cadastros auxiliares", str(error))

    def refresh(self):
        self.refresh_partner_options()
        self.fill_tree(self.tree_frame, get_auxiliary_records(), self.columns)

    def export_xlsx(self):
        export_xlsx("cadastros_auxiliares", self.columns, get_auxiliary_records())

    def export_pdf(self):
        generate_table_pdf("cadastros_auxiliares", "Cadastros auxiliares", self.columns, get_auxiliary_records())

class GenericProcessTab(AttachmentActionsMixin, PartnerLookupMixin, BaseFrame):
    def __init__(self, master, module_code, current_user=None, partner_kind=None):
        super().__init__(master, padding=10)
        self.module_code = module_code
        self.current_user = current_user or {}
        self.partner_kind = partner_kind
        self.config = ERP_MODULE_CONFIGS[module_code]
        self.selected_record_id = None
        self.init_attachment_state()
        self.columns = ["id", "codigo", "categoria", "titulo", "parceiro", "produto", "data_inicio", "data_fim", "status", "valor_previsto", "valor_real", "cost_center", "responsavel", "prioridade", "anexo"]
        self.build()
        self.refresh()

    def build(self):
        ttk.Label(self, text=self.config["title"], font=("Segoe UI", 14, "bold"), style="Title.TLabel").pack(anchor="w", pady=(0, 8))
        form = ttk.LabelFrame(self, text="Registro", padding=10)
        form.pack(fill="x", pady=(0, 8))

        for column in range(4):
            form.columnconfigure(column, weight=1, uniform="erp_form")

        self.make_label(form, "Categoria:", 0, 0)
        self.categoria = self.make_combo(form, 1, 0, self.config["categories"], width=28)
        self.make_label(form, "Título:", 0, 1)
        self.titulo = self.make_entry(form, 1, 1, width=36)
        self.make_label(form, "Parceiro:", 0, 2)
        partner = self.make_combo(form, 1, 2, get_partner_options(self.partner_kind), width=36)
        self.make_label(form, "Produto/serviço:", 0, 3)
        self.product = self.make_combo(form, 1, 3, get_product_options(), width=32)
        self.make_label(form, "Funcionário:", 2, 0)
        self.employee = self.make_combo(form, 3, 0, get_employee_options(), width=30)
        self.make_label(form, "Projeto/OS:", 2, 1)
        self.project_code = self.make_entry(form, 3, 1, width=20)
        self.make_label(form, "Centro de custo:", 2, 2)
        self.cost_center = self.make_combo(form, 3, 2, get_auxiliary_options("Centro de custo"), width=28)
        self.cost_center.configure(state="normal")
        self.make_label(form, "Responsável:", 2, 3)
        self.responsavel = self.make_entry(form, 3, 3, width=28)
        self.make_label(form, "Data inicial:", 4, 0)
        self.data_inicio = self.make_date_entry(form, 5, 0, width=16)
        self.make_label(form, "Data final:", 4, 1)
        self.data_fim = self.make_date_entry(form, 5, 1, width=16)
        self.make_label(form, "Status:", 4, 2)
        self.status = self.make_combo(form, 5, 2, self.config["statuses"], width=24)
        self.make_label(form, "Prioridade:", 4, 3)
        self.prioridade = self.make_combo(form, 5, 3, ["Normal", "Baixa", "Alta", "Crítica"], width=18)
        self.make_label(form, "Valor previsto:", 6, 0)
        self.valor_previsto = self.make_entry(form, 7, 0, width=18)
        self.make_label(form, "Valor real:", 6, 1)
        self.valor_real = self.make_entry(form, 7, 1, width=18)
        self.make_label(form, "Descrição:", 6, 2)
        self.descricao = self.make_entry(form, 7, 2, width=36)
        self.make_label(form, "Observações:", 6, 3)
        self.observacoes = self.make_entry(form, 7, 3, width=36)
        self.make_label(form, "Anexo:", 8, 0)
        attachment = ttk.Frame(form)
        attachment.grid(row=9, column=0, padx=6, pady=(0, 6), sticky="ew")
        ttk.Button(attachment, text="Anexar", command=self.choose_attachment).pack(side="left", padx=(0, 4))
        ttk.Button(attachment, text="Limpar", command=self.clear_attachment).pack(side="left")
        self.attachment_label = ttk.Label(form, text="Sem anexo", style="Muted.TLabel")
        self.attachment_label.grid(row=9, column=1, padx=6, pady=(0, 6), sticky="w")
        self.setup_partner_lookup(partner, None, self.partner_kind)

        self.save_button = ttk.Button(form, text="Salvar registro", style="Primary.TButton", command=self.save)
        self.save_button.grid(row=10, column=0, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Novo", command=self.clear_form).grid(row=10, column=1, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Exportar XLSX", command=self.export_xlsx).grid(row=10, column=2, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Gerar PDF", command=self.export_pdf).grid(row=10, column=3, padx=6, pady=4, sticky="ew")

        self.partner_info_label = ttk.Label(self, text="", style="Muted.TLabel", wraplength=980)
        self.partner_info_label.pack(anchor="w", pady=(0, 6))
        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.load_selected_record)

    def option_by_code(self, combo, code):
        code = str(code or "")

        for option in combo["values"]:
            if option.startswith(f"{code} - ") or option.endswith(f"({code})"):
                return option

        return ""

    def clear_form(self):
        self.selected_record_id = None
        self.clear_entries(self.titulo, self.project_code, self.responsavel, self.data_inicio, self.data_fim, self.valor_previsto, self.valor_real, self.descricao, self.observacoes)
        if self.config["categories"]:
            self.categoria.set(self.config["categories"][0])
        if self.config["statuses"]:
            self.status.set(self.config["statuses"][0])
        self.prioridade.set("Normal")
        self.product["values"] = get_product_options()
        self.employee["values"] = get_employee_options()
        self.cost_center["values"] = get_auxiliary_options("Centro de custo")
        self.refresh_partner_options()
        self.partner_info_label.configure(text="")
        self.init_attachment_state()
        self.set_attachment_label()
        self.save_button.configure(text="Salvar registro")

    def load_partner_details(self):
        partner = self.resolve_selected_partner()

        if not partner:
            self.partner_info_label.configure(text="")
            return

        details = [
            partner.get("nome_razao", ""),
            f"CPF/CNPJ: {partner.get('documento', '')}",
            f"Tipo: {partner.get('tipo_parceiro', '')}",
            f"Cidade: {partner.get('cidade', '')}/{partner.get('estado', '')}",
            f"Telefone: {partner.get('telefone', '')}",
            f"E-mail: {partner.get('email', '')}",
        ]
        self.partner_info_label.configure(text=" | ".join([detail for detail in details if detail and not detail.endswith(": ")]))

    def load_selected_record(self, event=None):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            return

        record = get_erp_record_by_id(row.get("id"), self.module_code)

        if not record:
            return

        partner = get_partner_by_code(record.get("partner_code"))
        self.selected_record_id = record.get("id")
        self.set_combo(self.categoria, record.get("categoria"))
        self.set_entry(self.titulo, record.get("titulo"))
        self.set_combo(self.partner, format_partner_option(partner))
        self.set_combo(self.product, self.option_by_code(self.product, record.get("product_code")))
        self.set_combo(self.employee, self.option_by_code(self.employee, record.get("employee_code")))
        self.set_entry(self.project_code, record.get("project_code"))
        self.set_combo(self.cost_center, record.get("cost_center"))
        self.set_entry(self.responsavel, record.get("responsavel"))
        self.set_entry(self.data_inicio, format_date_br(record.get("data_inicio")))
        self.set_entry(self.data_fim, format_date_br(record.get("data_fim")))
        self.set_combo(self.status, record.get("status"))
        self.set_combo(self.prioridade, record.get("prioridade") or "Normal")
        self.set_entry(self.valor_previsto, format_money(record.get("valor_previsto")))
        self.set_entry(self.valor_real, format_money(record.get("valor_real")))
        self.set_entry(self.descricao, record.get("descricao"))
        self.set_entry(self.observacoes, record.get("observacoes"))
        self.selected_attachment_source = ""
        self.current_attachment = record.get("attachment_path") or ""
        self.set_attachment_label()
        self.load_partner_details()
        self.save_button.configure(text="Atualizar registro")

    def save(self):
        try:
            partner = self.resolve_selected_partner()
            partner_code = partner.get("codigo") if partner else ""
            product_code = option_code(self.product.get())
            employee_code = option_code(self.employee.get())

            if self.selected_record_id:
                record_id = update_erp_record(
                    self.selected_record_id, self.module_code, self.categoria.get(), self.titulo.get(),
                    partner_code, product_code, employee_code, self.project_code.get(), self.cost_center.get(),
                    self.data_inicio.get(), self.data_fim.get(), self.status.get(), self.valor_previsto.get(),
                    self.valor_real.get(), self.prioridade.get(), self.responsavel.get(), self.descricao.get(),
                    self.observacoes.get(), self.selected_attachment_source, self.current_attachment
                )
                message = "Registro atualizado com sucesso."
            else:
                record_id = create_erp_record(
                    self.module_code, self.categoria.get(), self.titulo.get(), partner_code, product_code, employee_code,
                    self.project_code.get(), self.cost_center.get(), self.data_inicio.get(), self.data_fim.get(),
                    self.status.get(), self.valor_previsto.get(), self.valor_real.get(), self.prioridade.get(),
                    self.responsavel.get(), self.descricao.get(), self.observacoes.get(), self.selected_attachment_source,
                    self.current_user.get("id")
                )
                message = f"Registro criado com sucesso.\nCódigo: {record_id}"

            log_audit(self.current_user.get("id"), "salvar", self.module_code, str(record_id), self.titulo.get())
            messagebox.showinfo(self.config["title"], message)
            self.clear_form()
            self.refresh()
        except Exception as error:
            messagebox.showerror(self.config["title"], str(error))

    def refresh(self):
        self.refresh_partner_options()
        self.product["values"] = get_product_options()
        self.employee["values"] = get_employee_options()
        self.cost_center["values"] = get_auxiliary_options("Centro de custo")
        self.fill_tree(self.tree_frame, get_erp_records(self.module_code), self.columns)

    def export_xlsx(self):
        export_xlsx(self.module_code, self.columns, get_erp_records(self.module_code))

    def export_pdf(self):
        generate_table_pdf(self.module_code, self.config["title"], self.columns, get_erp_records(self.module_code))

class CommercialPipelineTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "comercial", current_user, "Cliente")

class ContractsTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "contratos", current_user, "Cliente")

class ProjectsTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "projetos", current_user, "Cliente")

class ScheduleTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "escalas", current_user)

class TimeTrackingTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "horas", current_user)

class HumanResourcesTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "rh", current_user, "Colaborador")

class TravelTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "viagens", current_user)

class ReceivingTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "recebimento", current_user, "Fornecedor")

class EquipmentTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "equipamentos", current_user)

class FiscalDocumentsTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "fiscal", current_user)

class OperationalReportsTab(GenericProcessTab):
    def __init__(self, master, current_user=None):
        super().__init__(master, "operacional", current_user, "Cliente")

class ProposalsTab(AttachmentActionsMixin, PartnerLookupMixin, BaseFrame):
    def __init__(self, master, current_user=None):
        super().__init__(master, padding=10)
        self.current_user = current_user or {}
        self.selected_record_id = None
        self.partner_kind = "Cliente"
        self.init_attachment_state()
        self.columns = ["id", "codigo", "cliente", "projeto", "revisao_label", "data_emissao", "validade", "status", "custo_total", "preco_venda", "lucro_previsto", "margem_percentual", "anexo"]
        self.build()
        self.refresh()

    def build(self):
        ttk.Label(self, text="Propostas comerciais", font=("Segoe UI", 14, "bold"), style="Title.TLabel").pack(anchor="w", pady=(0, 8))
        form = ttk.LabelFrame(self, text="Proposta", padding=10)
        form.pack(fill="x", pady=(0, 8))

        for column in range(4):
            form.columnconfigure(column, weight=1, uniform="proposal_form")

        self.make_label(form, "Cliente:", 0, 0)
        partner = self.make_combo(form, 1, 0, get_partner_options("Cliente"), width=36)
        self.make_label(form, "Projeto/escopo:", 0, 1)
        self.projeto = self.make_entry(form, 1, 1, width=42)
        self.make_label(form, "Revisão:", 0, 2)
        self.revisao = self.make_entry(form, 1, 2, width=10)
        self.revisao.insert(0, "0")
        self.make_label(form, "Status:", 0, 3)
        self.status = self.make_combo(form, 1, 3, ["Em elaboração", "Enviada", "Aguardando cliente", "Aprovada", "Recusada", "Cancelada"], width=24)
        self.make_label(form, "Emissão:", 2, 0)
        self.data_emissao = self.make_date_entry(form, 3, 0, width=16)
        self.data_emissao.insert(0, today_br())
        self.make_label(form, "Validade:", 2, 1)
        self.validade = self.make_date_entry(form, 3, 1, width=16)
        self.make_label(form, "Prazo:", 2, 2)
        self.prazo = self.make_entry(form, 3, 2, width=24)
        self.make_label(form, "Pagamento:", 2, 3)
        self.pagamento = self.make_entry(form, 3, 3, width=28)
        self.make_label(form, "Horas normais:", 4, 0)
        self.horas_normais = self.make_entry(form, 5, 0, width=14)
        self.make_label(form, "Horas 50%:", 4, 1)
        self.horas_50 = self.make_entry(form, 5, 1, width=14)
        self.make_label(form, "Horas 100%:", 4, 2)
        self.horas_100 = self.make_entry(form, 5, 2, width=14)
        self.make_label(form, "Valor hora:", 4, 3)
        self.valor_hora = self.make_entry(form, 5, 3, width=16)
        self.make_label(form, "Viagem:", 6, 0)
        self.custos_viagem = self.make_entry(form, 7, 0, width=16)
        self.make_label(form, "Materiais:", 6, 1)
        self.custos_materiais = self.make_entry(form, 7, 1, width=16)
        self.make_label(form, "Terceiros:", 6, 2)
        self.custos_terceiros = self.make_entry(form, 7, 2, width=16)
        self.make_label(form, "Impostos %:", 6, 3)
        self.impostos_percentual = self.make_entry(form, 7, 3, width=12)
        self.make_label(form, "Margem %:", 8, 0)
        self.margem_percentual = self.make_entry(form, 9, 0, width=12)
        self.make_label(form, "Condições:", 8, 1)
        self.condicoes = self.make_entry(form, 9, 1, width=36)
        self.make_label(form, "Incluso:", 8, 2)
        self.incluso = self.make_entry(form, 9, 2, width=36)
        self.make_label(form, "Não incluso:", 8, 3)
        self.nao_incluso = self.make_entry(form, 9, 3, width=36)
        self.make_label(form, "Motivo/status:", 10, 0)
        self.motivo_status = self.make_entry(form, 11, 0, width=36)
        self.make_label(form, "Anexo:", 10, 1)
        attachment = ttk.Frame(form)
        attachment.grid(row=11, column=1, padx=6, pady=(0, 6), sticky="ew")
        ttk.Button(attachment, text="Anexar", command=self.choose_attachment).pack(side="left", padx=(0, 4))
        ttk.Button(attachment, text="Limpar", command=self.clear_attachment).pack(side="left")
        self.attachment_label = ttk.Label(form, text="Sem anexo", style="Muted.TLabel")
        self.attachment_label.grid(row=11, column=2, padx=6, pady=(0, 6), sticky="w")
        self.setup_partner_lookup(partner, None, "Cliente")

        self.save_button = ttk.Button(form, text="Salvar proposta", style="Primary.TButton", command=self.save)
        self.save_button.grid(row=12, column=0, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Calcular", command=self.calculate_and_show).grid(row=12, column=1, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="PDF do selecionado", command=self.generate_selected_pdf).grid(row=12, column=2, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Novo", command=self.clear_form).grid(row=12, column=3, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Exportar XLSX", command=self.export_xlsx).grid(row=13, column=0, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Gerar PDF geral", command=self.export_pdf).grid(row=13, column=1, padx=6, pady=4, sticky="ew")

        self.result_label = ttk.Label(self, text="", style="Muted.TLabel")
        self.result_label.pack(anchor="w", pady=(0, 6))
        self.partner_info_label = ttk.Label(self, text="", style="Muted.TLabel", wraplength=980)
        self.partner_info_label.pack(anchor="w", pady=(0, 6))
        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.load_selected_record)

    def calculate_values(self):
        return calculate_proposal_values(
            self.horas_normais.get(), self.horas_50.get(), self.horas_100.get(), self.valor_hora.get(),
            self.custos_viagem.get(), self.custos_materiais.get(), self.custos_terceiros.get(),
            self.impostos_percentual.get(), self.margem_percentual.get()
        )

    def calculate_and_show(self):
        try:
            custo_total, preco_venda, lucro_previsto = self.calculate_values()
            self.result_label.configure(text=f"Custo R$ {format_money(custo_total)} | Venda R$ {format_money(preco_venda)} | Lucro previsto R$ {format_money(lucro_previsto)}")
        except Exception as error:
            messagebox.showerror("Propostas", str(error))

    def clear_form(self):
        self.selected_record_id = None
        self.clear_entries(
            self.projeto, self.revisao, self.validade, self.prazo, self.pagamento,
            self.horas_normais, self.horas_50, self.horas_100, self.valor_hora,
            self.custos_viagem, self.custos_materiais, self.custos_terceiros,
            self.impostos_percentual, self.margem_percentual, self.condicoes,
            self.incluso, self.nao_incluso, self.motivo_status
        )
        self.revisao.insert(0, "0")
        self.data_emissao.delete(0, "end")
        self.data_emissao.insert(0, today_br())
        self.status.set("Em elaboração")
        self.refresh_partner_options()
        self.partner_info_label.configure(text="")
        self.result_label.configure(text="")
        self.init_attachment_state()
        self.set_attachment_label()
        self.save_button.configure(text="Salvar proposta")

    def load_partner_details(self):
        partner = self.resolve_selected_partner()

        if not partner:
            self.partner_info_label.configure(text="")
            return

        details = [
            partner.get("nome_razao", ""),
            f"CPF/CNPJ: {partner.get('documento', '')}",
            f"Cidade: {partner.get('cidade', '')}/{partner.get('estado', '')}",
            f"Responsável: {partner.get('responsavel', '')}",
            f"E-mail: {partner.get('email', '')}",
        ]
        self.partner_info_label.configure(text=" | ".join([detail for detail in details if detail and not detail.endswith(": ")]))

    def load_selected_record(self, event=None):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            return

        record = get_proposal_by_id(row.get("id"))

        if not record:
            return

        partner = get_partner_by_code(record.get("partner_code"))
        self.selected_record_id = record.get("id")
        self.set_combo(self.partner, format_partner_option(partner))
        self.set_entry(self.projeto, record.get("projeto"))
        self.set_entry(self.revisao, record.get("revisao"))
        self.set_entry(self.data_emissao, format_date_br(record.get("data_emissao")))
        self.set_entry(self.validade, format_date_br(record.get("validade")))
        self.set_entry(self.prazo, record.get("prazo"))
        self.set_entry(self.pagamento, record.get("pagamento"))
        self.set_entry(self.condicoes, record.get("condicoes"))
        self.set_entry(self.horas_normais, format_money(record.get("horas_normais")))
        self.set_entry(self.horas_50, format_money(record.get("horas_50")))
        self.set_entry(self.horas_100, format_money(record.get("horas_100")))
        self.set_entry(self.valor_hora, format_money(record.get("valor_hora")))
        self.set_entry(self.custos_viagem, format_money(record.get("custos_viagem")))
        self.set_entry(self.custos_materiais, format_money(record.get("custos_materiais")))
        self.set_entry(self.custos_terceiros, format_money(record.get("custos_terceiros")))
        self.set_entry(self.impostos_percentual, format_money(record.get("impostos_percentual")))
        self.set_entry(self.margem_percentual, format_money(record.get("margem_percentual")))
        self.set_entry(self.incluso, record.get("incluso"))
        self.set_entry(self.nao_incluso, record.get("nao_incluso"))
        self.set_combo(self.status, record.get("status"))
        self.set_entry(self.motivo_status, record.get("motivo_status"))
        self.selected_attachment_source = ""
        self.current_attachment = record.get("attachment_path") or ""
        self.set_attachment_label()
        self.load_partner_details()
        self.calculate_and_show()
        self.save_button.configure(text="Atualizar proposta")

    def save(self):
        try:
            partner = self.resolve_selected_partner()
            partner_code = partner.get("codigo") if partner else ""

            if self.selected_record_id:
                record_id = update_proposal(
                    self.selected_record_id, partner_code, self.projeto.get(), self.revisao.get(), self.data_emissao.get(),
                    self.validade.get(), self.prazo.get(), self.pagamento.get(), self.condicoes.get(),
                    self.horas_normais.get(), self.horas_50.get(), self.horas_100.get(), self.valor_hora.get(),
                    self.custos_viagem.get(), self.custos_materiais.get(), self.custos_terceiros.get(),
                    self.impostos_percentual.get(), self.margem_percentual.get(), self.incluso.get(),
                    self.nao_incluso.get(), self.status.get(), self.motivo_status.get(),
                    self.selected_attachment_source, self.current_attachment
                )
                message = "Proposta atualizada com sucesso."
            else:
                record_id = create_proposal(
                    partner_code, self.projeto.get(), self.revisao.get(), self.data_emissao.get(),
                    self.validade.get(), self.prazo.get(), self.pagamento.get(), self.condicoes.get(),
                    self.horas_normais.get(), self.horas_50.get(), self.horas_100.get(), self.valor_hora.get(),
                    self.custos_viagem.get(), self.custos_materiais.get(), self.custos_terceiros.get(),
                    self.impostos_percentual.get(), self.margem_percentual.get(), self.incluso.get(),
                    self.nao_incluso.get(), self.status.get(), self.motivo_status.get(), self.selected_attachment_source
                )
                message = f"Proposta criada com sucesso.\nCódigo: {record_id}"

            log_audit(self.current_user.get("id"), "salvar", "commercial_proposals", str(record_id), self.projeto.get())
            messagebox.showinfo("Propostas", message)
            self.clear_form()
            self.refresh()
        except Exception as error:
            messagebox.showerror("Propostas", str(error))

    def refresh(self):
        self.refresh_partner_options()
        self.fill_tree(self.tree_frame, get_proposals(), self.columns)

    def selected_row(self):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            messagebox.showwarning("Propostas", "Selecione uma proposta na tabela.")
            return None

        return row

    def generate_selected_pdf(self):
        row = self.selected_row()

        if row:
            generate_table_pdf(f"proposta_{row.get('codigo')}", f"Proposta {row.get('codigo')}", self.columns, [row])

    def export_xlsx(self):
        export_xlsx("propostas", self.columns, get_proposals())

    def export_pdf(self):
        generate_table_pdf("propostas", "Propostas comerciais", self.columns, get_proposals())

class FinancialControlTab(AttachmentActionsMixin, PartnerLookupMixin, BaseFrame):
    def __init__(self, master, current_user=None):
        super().__init__(master, padding=10)
        self.current_user = current_user or {}
        self.selected_record_id = None
        self.init_attachment_state()
        self.columns = ["id", "codigo", "tipo", "descricao", "parceiro", "project_code", "cost_center", "valor", "data_emissao", "data_vencimento", "data_baixa", "status", "forma_pagamento", "conta_bancaria", "anexo"]
        self.build()
        self.refresh()

    def build(self):
        ttk.Label(self, text="Financeiro e fluxo de caixa", font=("Segoe UI", 14, "bold"), style="Title.TLabel").pack(anchor="w", pady=(0, 8))
        form = ttk.LabelFrame(self, text="Lançamento financeiro", padding=10)
        form.pack(fill="x", pady=(0, 8))

        for column in range(4):
            form.columnconfigure(column, weight=1, uniform="financial_form")

        self.make_label(form, "Tipo:", 0, 0)
        self.tipo = self.make_combo(form, 1, 0, ["Entrada", "Saída"], width=16)
        self.make_label(form, "Descrição:", 0, 1)
        self.descricao = self.make_entry(form, 1, 1, width=36)
        self.make_label(form, "Parceiro:", 0, 2)
        partner = self.make_combo(form, 1, 2, get_partner_options(), width=36)
        self.make_label(form, "Documento ID:", 0, 3)
        self.document_id = self.make_entry(form, 1, 3, width=12)
        self.make_label(form, "Projeto/OS:", 2, 0)
        self.project_code = self.make_entry(form, 3, 0, width=20)
        self.make_label(form, "Centro de custo:", 2, 1)
        self.cost_center = self.make_combo(form, 3, 1, get_auxiliary_options("Centro de custo"), width=28)
        self.cost_center.configure(state="normal")
        self.make_label(form, "Valor:", 2, 2)
        self.valor = self.make_entry(form, 3, 2, width=16)
        self.make_label(form, "Status:", 2, 3)
        self.status = self.make_combo(form, 3, 3, ["Aberto", "Programado", "Aprovado", "Pago", "Recebido", "Vencido", "Cancelado"], width=20)
        self.make_label(form, "Emissão:", 4, 0)
        self.data_emissao = self.make_date_entry(form, 5, 0, width=16)
        self.data_emissao.insert(0, today_br())
        self.make_label(form, "Vencimento:", 4, 1)
        self.data_vencimento = self.make_date_entry(form, 5, 1, width=16)
        self.make_label(form, "Baixa:", 4, 2)
        self.data_baixa = self.make_date_entry(form, 5, 2, width=16)
        self.make_label(form, "Forma pagamento:", 4, 3)
        self.forma_pagamento = self.make_combo(form, 5, 3, get_auxiliary_options("Forma de pagamento"), width=28)
        self.forma_pagamento.configure(state="normal")
        self.make_label(form, "Conta bancária:", 6, 0)
        self.conta_bancaria = self.make_combo(form, 7, 0, get_auxiliary_options("Banco/conta"), width=28)
        self.conta_bancaria.configure(state="normal")
        self.make_label(form, "Observações:", 6, 1)
        self.observacoes = self.make_entry(form, 7, 1, width=40)
        self.make_label(form, "Anexo:", 6, 2)
        attachment = ttk.Frame(form)
        attachment.grid(row=7, column=2, padx=6, pady=(0, 6), sticky="ew")
        ttk.Button(attachment, text="Anexar", command=self.choose_attachment).pack(side="left", padx=(0, 4))
        ttk.Button(attachment, text="Limpar", command=self.clear_attachment).pack(side="left")
        self.attachment_label = ttk.Label(form, text="Sem anexo", style="Muted.TLabel")
        self.attachment_label.grid(row=7, column=3, padx=6, pady=(0, 6), sticky="w")
        self.setup_partner_lookup(partner, None)

        self.save_button = ttk.Button(form, text="Salvar lançamento", style="Primary.TButton", command=self.save)
        self.save_button.grid(row=8, column=0, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Novo", command=self.clear_form).grid(row=8, column=1, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Exportar XLSX", command=self.export_xlsx).grid(row=8, column=2, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Gerar PDF", command=self.export_pdf).grid(row=8, column=3, padx=6, pady=4, sticky="ew")

        self.partner_info_label = ttk.Label(self, text="", style="Muted.TLabel", wraplength=980)
        self.partner_info_label.pack(anchor="w", pady=(0, 6))
        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.load_selected_record)

    def clear_form(self):
        self.selected_record_id = None
        self.clear_entries(self.descricao, self.document_id, self.project_code, self.valor, self.data_vencimento, self.data_baixa, self.observacoes)
        self.tipo.set("Entrada")
        self.status.set("Aberto")
        self.data_emissao.delete(0, "end")
        self.data_emissao.insert(0, today_br())
        self.cost_center["values"] = get_auxiliary_options("Centro de custo")
        self.forma_pagamento["values"] = get_auxiliary_options("Forma de pagamento")
        self.conta_bancaria["values"] = get_auxiliary_options("Banco/conta")
        self.refresh_partner_options()
        self.partner_info_label.configure(text="")
        self.init_attachment_state()
        self.set_attachment_label()
        self.save_button.configure(text="Salvar lançamento")

    def load_partner_details(self):
        partner = self.resolve_selected_partner()

        if not partner:
            self.partner_info_label.configure(text="")
            return

        details = [
            partner.get("nome_razao", ""),
            f"CPF/CNPJ: {partner.get('documento', '')}",
            f"Tipo: {partner.get('tipo_parceiro', '')}",
            f"Cidade: {partner.get('cidade', '')}/{partner.get('estado', '')}",
            f"E-mail: {partner.get('email', '')}",
        ]
        self.partner_info_label.configure(text=" | ".join([detail for detail in details if detail and not detail.endswith(": ")]))

    def load_selected_record(self, event=None):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            return

        record = get_financial_entry_by_id(row.get("id"))

        if not record:
            return

        partner = get_partner_by_code(record.get("partner_code"))
        self.selected_record_id = record.get("id")
        self.set_combo(self.tipo, record.get("tipo"))
        self.set_entry(self.descricao, record.get("descricao"))
        self.set_combo(self.partner, format_partner_option(partner))
        self.set_entry(self.document_id, record.get("document_id") or "")
        self.set_entry(self.project_code, record.get("project_code"))
        self.set_combo(self.cost_center, record.get("cost_center"))
        self.set_entry(self.valor, format_money(record.get("valor")))
        self.set_entry(self.data_emissao, format_date_br(record.get("data_emissao")))
        self.set_entry(self.data_vencimento, format_date_br(record.get("data_vencimento")))
        self.set_entry(self.data_baixa, format_date_br(record.get("data_baixa")))
        self.set_combo(self.status, record.get("status"))
        self.set_combo(self.forma_pagamento, record.get("forma_pagamento"))
        self.set_combo(self.conta_bancaria, record.get("conta_bancaria"))
        self.set_entry(self.observacoes, record.get("observacoes"))
        self.selected_attachment_source = ""
        self.current_attachment = record.get("attachment_path") or ""
        self.set_attachment_label()
        self.load_partner_details()
        self.save_button.configure(text="Atualizar lançamento")

    def save(self):
        try:
            partner = self.resolve_selected_partner()
            partner_code = partner.get("codigo") if partner else ""
            document_id = self.document_id.get().strip() or None

            if self.selected_record_id:
                record_id = update_financial_entry(
                    self.selected_record_id, self.tipo.get(), self.descricao.get(), partner_code, document_id,
                    self.project_code.get(), self.cost_center.get(), self.valor.get(), self.data_emissao.get(),
                    self.data_vencimento.get(), self.data_baixa.get(), self.status.get(), self.forma_pagamento.get(),
                    self.conta_bancaria.get(), self.observacoes.get(), self.selected_attachment_source, self.current_attachment
                )
                message = "Lançamento atualizado com sucesso."
            else:
                record_id = create_financial_entry(
                    self.tipo.get(), self.descricao.get(), partner_code, document_id, self.project_code.get(),
                    self.cost_center.get(), self.valor.get(), self.data_emissao.get(), self.data_vencimento.get(),
                    self.status.get(), self.forma_pagamento.get(), self.conta_bancaria.get(), self.observacoes.get(),
                    self.data_baixa.get(), self.selected_attachment_source
                )
                message = f"Lançamento criado com sucesso.\nCódigo: {record_id}"

            log_audit(self.current_user.get("id"), "salvar", "financial_entries", str(record_id), self.descricao.get())
            messagebox.showinfo("Financeiro", message)
            self.clear_form()
            self.refresh()
        except Exception as error:
            messagebox.showerror("Financeiro", str(error))

    def refresh(self):
        self.refresh_partner_options()
        self.cost_center["values"] = get_auxiliary_options("Centro de custo")
        self.forma_pagamento["values"] = get_auxiliary_options("Forma de pagamento")
        self.conta_bancaria["values"] = get_auxiliary_options("Banco/conta")
        self.fill_tree(self.tree_frame, get_financial_entries(), self.columns)

    def export_xlsx(self):
        export_xlsx("financeiro", self.columns, get_financial_entries())

    def export_pdf(self):
        generate_table_pdf("financeiro", "Financeiro e fluxo de caixa", self.columns, get_financial_entries())

class BusinessDashboardTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=0)
        self.pending_columns = PENDING_DOCUMENT_COLUMNS
        self.low_stock_columns = LOW_STOCK_COLUMNS
        self.build()

    def build(self):
        stats = self.get_stats()
        self.stats = {
            "Estoque": stats.get("Produtos", 0),
            "Clientes": stats.get("Clientes", stats.get("Parceiros", 0)),
            "Serviços": stats.get("Serviços", 0),
        }
        self.criteria = self.get_dashboard_criteria(stats)
        self.canvas = tk.Canvas(self, bg=COLOR_BG, height=421, highlightthickness=0, bd=0)
        self.canvas.pack(fill="both", expand=True)
        self.canvas.bind("<Configure>", self.redraw_dashboard)

    def redraw_dashboard(self, event=None):
        canvas = self.canvas
        canvas.delete("all")
        width = max(canvas.winfo_width(), 742)
        left = 21
        right = 9
        gap = 14
        card_top = 22
        card_height = 103
        card_width = (width - left - right - (gap * 2)) / 3
        card_data = [
            ("Estoque", self.stats["Estoque"], COLOR_ACCENT),
            ("Clientes", self.stats["Clientes"], COLOR_CYAN),
            ("Serviços", self.stats["Serviços"], COLOR_ACCENT),
        ]

        for index, (label, value, color) in enumerate(card_data):
            x1 = left + index * (card_width + gap)
            x2 = x1 + card_width
            draw_round_rect(canvas, x1, card_top, x2, card_top + card_height, 10, COLOR_PANEL, COLOR_BORDER, 1)
            canvas.create_text(x1 + 15, card_top + 37, text=str(value), anchor="w", fill=color, font=("Segoe UI", 26, "bold"))
            canvas.create_text(x1 + 15, card_top + 76, text=label, anchor="w", fill=COLOR_MUTED, font=("Segoe UI", 12))

        panel_top = 146
        panel_height = 200
        panel_left = left
        panel_right = width - right
        draw_round_rect(canvas, panel_left, panel_top, panel_right, panel_top + panel_height, 10, COLOR_PANEL_2, COLOR_BORDER, 1)
        canvas.create_text(panel_left + 22, panel_top + 27, text="Critérios do sistema", anchor="w", fill=COLOR_ACCENT, font=("Segoe UI", 13, "bold"))

        for y in (panel_top + 51, panel_top + 101, panel_top + 152):
            canvas.create_line(panel_left + 1, y, panel_right - 1, y, fill=COLOR_BORDER, width=1)

        if not self.criteria:
            canvas.create_text(panel_left + 22, panel_top + 87, text="Sem critérios ativos para sinalizar agora.", anchor="w", fill=COLOR_MUTED, font=("Segoe UI", 12))
            return

        icon_x = panel_right - 36
        row_centers = (panel_top + 76, panel_top + 127, panel_top + 177)

        for row, y in zip(self.criteria[:3], row_centers):
            canvas.create_text(panel_left + 22, y - 7, text=row["titulo"], anchor="w", fill=COLOR_TEXT, font=("Segoe UI", 11, "bold"))
            canvas.create_text(panel_left + 22, y + 12, text=row["detalhe"], anchor="w", fill=COLOR_MUTED, font=("Segoe UI", 9))
            self.draw_criteria_icon(canvas, icon_x, y, row["estado"])

    def draw_criteria_icon(self, canvas, x, y, state):
        if state == "ok":
            canvas.create_oval(x - 14, y - 14, x + 14, y + 14, fill="#173C39", outline="")
            canvas.create_text(x, y - 1, text="✓", fill="#22C55E", font=("Segoe UI", 9, "bold"))
        elif state == "bad":
            canvas.create_oval(x - 14, y - 14, x + 14, y + 14, fill="#3B1114", outline="")
            canvas.create_text(x, y - 1, text="X", fill="#FCA5A5", font=("Segoe UI", 9, "bold"))
        elif state == "warn":
            canvas.create_oval(x - 14, y - 14, x + 14, y + 14, fill="#3A2F08", outline="")
            canvas.create_text(x, y - 1, text="!", fill="#FDE68A", font=("Segoe UI", 9, "bold"))

    def get_dashboard_criteria(self, stats):
        criteria = []
        low_stock = int(stats.get("Estoque baixo", 0) or 0)
        products = int(stats.get("Produtos", 0) or 0)

        if products:
            criteria.append({
                "titulo": "Estoque crítico",
                "detalhe": f"{low_stock} item(ns) abaixo do mínimo." if low_stock else "Todos os produtos estão acima do mínimo.",
                "estado": "bad" if low_stock else "ok",
            })

        revenue = parse_float_br(stats.get("Faturamento mês valor", 0))
        expenses = parse_float_br(stats.get("Compras mês valor", 0))

        if revenue > 0 or expenses > 0:
            criteria.append({
                "titulo": "Faturamento do mês",
                "detalhe": f"Vendas R$ {format_money(revenue)} contra compras R$ {format_money(expenses)}.",
                "estado": "ok" if revenue >= expenses and revenue > 0 else "warn",
            })

        overdue = int(stats.get("Contas vencidas", 0) or 0)
        open_payable_total = parse_float_br(stats.get("Contas abertas valor", 0))

        if overdue or open_payable_total:
            high_expense = revenue > 0 and open_payable_total > revenue * 0.7
            criteria.append({
                "titulo": "Despesas e contas a pagar",
                "detalhe": f"{overdue} vencida(s), R$ {format_money(open_payable_total)} em aberto.",
                "estado": "bad" if overdue or high_expense else "ok",
            })

        return criteria

    def get_recent_activity(self):
        items = []

        for document in get_documents()[:4]:
            items.append({
                "data": document.get("data_emissao", ""),
                "titulo": f"{document.get('tipo_documento', 'Documento')} {document.get('numero', '')}",
                "detalhe": f"{document.get('parceiro') or 'Sem parceiro'} | {document.get('status', '')} | R$ {document.get('total', '0,00')}",
            })

        conn = get_conn()
        rows = conn.execute("""
            SELECT tipo, quantidade, data_movimentacao, parceiro_nome, numero_documento
            FROM business_movements
            ORDER BY created_at DESC, id DESC
            LIMIT 4
        """).fetchall()
        conn.close()

        for row in rows:
            items.append({
                "data": row["data_movimentacao"],
                "titulo": row["tipo"],
                "detalhe": f"{format_money(row['quantidade'])} un | {row['parceiro_nome'] or 'Sem parceiro'} | {row['numero_documento'] or 'Sem documento'}",
            })

        items.sort(key=lambda item: str(item.get("data", "")), reverse=True)
        return items[:5]

    def get_stats(self):
        conn = get_conn()
        cur = conn.cursor()
        month = date.today().strftime("%Y-%m")
        total_vendido = cur.execute("SELECT COALESCE(SUM(total), 0) total FROM documents WHERE direcao = 'venda'").fetchone()["total"]
        total_comprado = cur.execute("SELECT COALESCE(SUM(total), 0) total FROM documents WHERE direcao = 'compra'").fetchone()["total"]
        vendido_mes = cur.execute("SELECT COALESCE(SUM(total), 0) total FROM documents WHERE direcao = 'venda' AND substr(data_emissao, 1, 7) = ?", (month,)).fetchone()["total"]
        comprado_mes = cur.execute("SELECT COALESCE(SUM(total), 0) total FROM documents WHERE direcao = 'compra' AND substr(data_emissao, 1, 7) = ?", (month,)).fetchone()["total"]
        contas_vencidas = cur.execute("SELECT COUNT(*) qtd FROM accounts_payable WHERE status = 'Em aberto' AND data_vencimento < ?", (today_iso(),)).fetchone()["qtd"]
        contas_abertas_valor = cur.execute("SELECT COALESCE(SUM(valor), 0) total FROM accounts_payable WHERE status IN ('Em aberto', 'Vencido')").fetchone()["total"]
        pending_placeholders = ",".join("?" for _ in FINAL_DOCUMENT_STATUSES)
        stats = {
            "Produtos": cur.execute("SELECT COUNT(*) qtd FROM products").fetchone()["qtd"],
            "Parceiros": cur.execute("SELECT COUNT(*) qtd FROM partners").fetchone()["qtd"],
            "Clientes": cur.execute("SELECT COUNT(*) qtd FROM partners WHERE lower(tipo_parceiro) LIKE '%cliente%'").fetchone()["qtd"],
            "Serviços": cur.execute("SELECT COUNT(*) qtd FROM services").fetchone()["qtd"],
            "Compras abertas": cur.execute("SELECT COUNT(*) qtd FROM documents WHERE direcao = 'compra' AND status NOT IN ('Recebido', 'Recebida', 'Cancelado')").fetchone()["qtd"],
            "Vendas abertas": cur.execute("SELECT COUNT(*) qtd FROM documents WHERE tipo_code IN ('OV', 'PV') AND status NOT IN ('Faturado', 'Faturado internamente', 'Cancelado')").fetchone()["qtd"],
            "Faturamentos mês": cur.execute("SELECT COUNT(*) qtd FROM documents WHERE tipo_code = 'FT' AND substr(data_emissao, 1, 7) = ?", (month,)).fetchone()["qtd"],
            "Valor vendido": f"R$ {format_money(total_vendido)}",
            "Valor comprado": f"R$ {format_money(total_comprado)}",
            "Faturamento mês valor": vendido_mes,
            "Compras mês valor": comprado_mes,
            "Contas vencidas": contas_vencidas,
            "Contas abertas valor": contas_abertas_valor,
            "Estoque baixo": cur.execute("SELECT COUNT(*) qtd FROM products WHERE estoque_minimo > 0 AND estoque_atual <= estoque_minimo").fetchone()["qtd"],
            "Documentos pendentes": cur.execute(f"SELECT COUNT(*) qtd FROM documents WHERE status NOT IN ({pending_placeholders})", FINAL_DOCUMENT_STATUSES).fetchone()["qtd"],
            "Movimentações": cur.execute("SELECT COUNT(*) qtd FROM business_movements").fetchone()["qtd"],
        }
        conn.close()
        return stats

    def get_low_stock(self):
        return get_business_low_stock_report()

    def get_pending_documents(self):
        return get_business_pending_documents_report()

class ProductsTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.selected_product_code = None
        self.selected_image_source = ""
        self.product_image_preview = None
        self.columns = ["codigo", "nome", "descricao", "categoria", "tipo_produto", "unidade", "ncm", "codigo_barras", "valor_custo", "valor_venda", "estoque_atual", "estoque_minimo", "status", "imagem", "observacoes"]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text="Cadastro de produto", padding=10)
        form.pack(fill="x", pady=(0, 8))

        self.make_label(form, "Nome:", 0, 0)
        self.nome = self.make_entry(form, 1, 0)
        self.make_label(form, "Descrição:", 0, 1)
        self.descricao = self.make_entry(form, 1, 1, width=45)
        self.make_label(form, "Categoria:", 0, 2)
        self.categoria = self.make_entry(form, 1, 2)
        self.make_label(form, "Tipo:", 0, 3)
        self.tipo_produto = self.make_combo(form, 1, 3, ["Matéria-prima", "Revenda", "Serviço", "Produto acabado", "Consumo interno", "Outro"])
        self.make_label(form, "Unidade:", 2, 0)
        self.unidade = self.make_combo(form, 3, 0, ["un", "kg", "m", "litro", "caixa", "serviço", "outro"])
        self.make_label(form, "NCM:", 2, 1)
        self.ncm = self.make_entry(form, 3, 1)
        self.make_label(form, "Código de barras:", 2, 2)
        self.codigo_barras = self.make_entry(form, 3, 2)
        self.make_label(form, "Custo:", 2, 3)
        self.valor_custo = self.make_entry(form, 3, 3)
        self.make_label(form, "Venda:", 4, 0)
        self.valor_venda = self.make_entry(form, 5, 0)
        self.make_label(form, "Estoque atual:", 4, 1)
        self.estoque_atual = self.make_entry(form, 5, 1)
        self.make_label(form, "Estoque mínimo:", 4, 2)
        self.estoque_minimo = self.make_entry(form, 5, 2)
        self.make_label(form, "Status:", 4, 3)
        self.status = self.make_combo(form, 5, 3, ["Ativo", "Inativo"])
        self.make_label(form, "Observações:", 6, 0)
        self.obs = self.make_entry(form, 7, 0, width=50)
        self.make_label(form, "Imagem:", 6, 1)
        image_actions = ttk.Frame(form)
        image_actions.grid(row=7, column=1, padx=6, pady=(0, 6), sticky="ew")
        ttk.Button(image_actions, text="Selecionar", command=self.choose_image).pack(side="left", padx=(0, 4))
        ttk.Button(image_actions, text="Limpar", command=self.clear_selected_image).pack(side="left")
        self.image_preview_label = ttk.Label(form, text="Sem imagem", style="Muted.TLabel")
        self.image_preview_label.grid(row=7, column=2, padx=6, pady=(0, 6), sticky="w")
        self.save_button = ttk.Button(form, text="Salvar produto", style="Primary.TButton", command=self.save)
        self.save_button.grid(row=8, column=0, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Novo produto", command=self.clear_form).grid(row=8, column=1, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Exportar XLSX", command=self.export_xlsx).grid(row=8, column=2, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Gerar PDF", command=self.export_pdf).grid(row=8, column=3, padx=6, pady=4, sticky="ew")

        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.load_selected_product)
        self.bind_double_click(self.tree_frame, self.columns, lambda row: self.open_product_detail(row.get("codigo")))

    def choose_image(self):
        path = filedialog.askopenfilename(title="Selecionar imagem do produto", filetypes=[("Imagens", "*.png;*.jpg;*.jpeg"), ("Todos os arquivos", "*.*")])
        if path:
            self.selected_image_source = path
            self.product_image_preview = load_tk_image(path, (72, 72))
            if self.product_image_preview:
                self.image_preview_label.configure(image=self.product_image_preview, text="")
            else:
                self.image_preview_label.configure(image="", text=Path(path).name)

    def clear_selected_image(self):
        self.selected_image_source = ""
        self.product_image_preview = None
        self.image_preview_label.configure(image="", text="Sem imagem")

    def clear_form(self):
        self.selected_product_code = None
        self.clear_entries(self.nome, self.descricao, self.categoria, self.ncm, self.codigo_barras, self.valor_custo, self.valor_venda, self.estoque_atual, self.estoque_minimo, self.obs)
        self.tipo_produto.set("Revenda")
        self.unidade.set("un")
        self.status.set("Ativo")
        self.clear_selected_image()
        self.save_button.configure(text="Salvar produto")

    def load_selected_product(self, event=None):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            return

        product = get_product_by_code(row.get("codigo"))

        if not product:
            return

        self.selected_product_code = product.get("codigo")
        self.set_entry(self.nome, product.get("nome"))
        self.set_entry(self.descricao, product.get("descricao"))
        self.set_entry(self.categoria, product.get("categoria"))
        self.set_combo(self.tipo_produto, product.get("tipo_produto"))
        self.set_combo(self.unidade, product.get("unidade"))
        self.set_entry(self.ncm, product.get("ncm"))
        self.set_entry(self.codigo_barras, product.get("codigo_barras"))
        self.set_entry(self.valor_custo, format_money(product.get("valor_custo")))
        self.set_entry(self.valor_venda, format_money(product.get("valor_venda")))
        self.set_entry(self.estoque_atual, format_money(product.get("estoque_atual")))
        self.set_entry(self.estoque_minimo, format_money(product.get("estoque_minimo")))
        self.set_combo(self.status, product.get("status"))
        self.set_entry(self.obs, product.get("observacoes"))
        self.selected_image_source = ""
        self.product_image_preview = load_tk_image(product.get("image_path"), (72, 72))

        if self.product_image_preview:
            self.image_preview_label.configure(image=self.product_image_preview, text="")
        else:
            self.image_preview_label.configure(image="", text="Sem imagem")

        self.save_button.configure(text="Atualizar produto")

    def save(self):
        try:
            if self.selected_product_code:
                codigo = update_product(
                    self.selected_product_code, self.nome.get(), self.descricao.get(), self.categoria.get(), self.tipo_produto.get(),
                    self.unidade.get(), self.ncm.get(), self.codigo_barras.get(), self.valor_custo.get(),
                    self.valor_venda.get(), self.estoque_atual.get(), self.estoque_minimo.get(),
                    self.status.get(), self.selected_image_source, self.obs.get()
                )
                messagebox.showinfo("Produtos", f"Produto atualizado com sucesso.\nCódigo: {codigo}")
            else:
                codigo = create_product(
                    self.nome.get(), self.descricao.get(), self.categoria.get(), self.tipo_produto.get(),
                    self.unidade.get(), self.ncm.get(), self.codigo_barras.get(), self.valor_custo.get(),
                    self.valor_venda.get(), self.estoque_atual.get(), self.estoque_minimo.get(),
                    self.status.get(), self.selected_image_source, self.obs.get()
                )
                messagebox.showinfo("Produtos", f"Produto cadastrado com sucesso.\nCódigo: {codigo}")

            self.clear_form()
            self.refresh()
        except Exception as error:
            messagebox.showerror("Produtos", str(error))

    def refresh(self):
        self.fill_tree(self.tree_frame, get_products(), self.columns)

    def export_xlsx(self):
        export_xlsx("produtos", self.columns, get_products())

    def export_pdf(self):
        generate_table_pdf("produtos", "Relatório de produtos", self.columns, get_products())

class PartnersTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.selected_partner_code = None
        self.kind_vars = {}
        self.columns = ["codigo", "tipo_pessoa", "documento", "nome_razao", "nome_fantasia", "inscricao_estadual", "telefone", "celular", "email", "cidade", "estado", "tipo_parceiro", "status"]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text="Cadastro de parceiro", padding=10)
        form.pack(fill="x", pady=(0, 8))
        self.make_label(form, "Tipo pessoa:", 0, 0)
        self.tipo_pessoa = self.make_combo(form, 1, 0, ["Jurídica", "Física"])
        self.make_label(form, "CPF/CNPJ:", 0, 1)
        self.documento = self.make_entry(form, 1, 1)
        self.make_label(form, "Nome/Razão social:", 0, 2)
        self.nome_razao = self.make_entry(form, 1, 2)
        self.make_label(form, "Nome fantasia:", 0, 3)
        self.nome_fantasia = self.make_entry(form, 1, 3)
        self.make_label(form, "Inscrição estadual:", 2, 0)
        self.inscricao_estadual = self.make_entry(form, 3, 0)
        self.make_label(form, "Inscrição municipal:", 2, 1)
        self.inscricao_municipal = self.make_entry(form, 3, 1)
        self.make_label(form, "Telefone:", 2, 2)
        self.telefone = self.make_entry(form, 3, 2)
        self.make_label(form, "Celular:", 2, 3)
        self.celular = self.make_entry(form, 3, 3)
        self.make_label(form, "E-mail:", 4, 0)
        self.email = self.make_entry(form, 5, 0)
        self.make_label(form, "Endereço:", 4, 1)
        self.endereco = self.make_entry(form, 5, 1)
        self.make_label(form, "Número:", 4, 2)
        self.numero = self.make_entry(form, 5, 2)
        self.make_label(form, "Bairro:", 4, 3)
        self.bairro = self.make_entry(form, 5, 3)
        self.make_label(form, "Cidade:", 6, 0)
        self.cidade = self.make_entry(form, 7, 0)
        self.make_label(form, "Estado:", 6, 1)
        self.estado = self.make_entry(form, 7, 1)
        self.make_label(form, "CEP:", 6, 2)
        self.cep = self.make_entry(form, 7, 2)
        self.make_label(form, "País:", 6, 3)
        self.pais = self.make_entry(form, 7, 3)
        self.pais.insert(0, "Brasil")
        self.make_label(form, "Responsável:", 8, 0)
        self.responsavel = self.make_entry(form, 9, 0)
        self.make_label(form, "Status:", 8, 1)
        self.status = self.make_combo(form, 9, 1, ["Ativo", "Inativo"])
        self.make_label(form, "Observações:", 8, 2)
        self.obs = self.make_entry(form, 9, 2, width=45)

        kinds = ttk.LabelFrame(form, text="Tipo de parceiro", padding=6)
        kinds.grid(row=10, column=0, columnspan=4, padx=6, pady=4, sticky="ew")
        for label in ["Cliente", "Fornecedor", "Colaborador", "Prestador de serviço", "Transportador", "Outro"]:
            var = tk.BooleanVar(value=label in ("Cliente", "Fornecedor"))
            self.kind_vars[label] = var
            ttk.Checkbutton(kinds, text=label, variable=var).pack(side="left", padx=4)

        self.save_button = ttk.Button(form, text="Salvar parceiro", style="Primary.TButton", command=self.save)
        self.save_button.grid(row=11, column=0, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Novo parceiro", command=self.clear_form).grid(row=11, column=1, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Exportar XLSX", command=self.export_xlsx).grid(row=11, column=2, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Gerar PDF", command=self.export_pdf).grid(row=11, column=3, padx=6, pady=4, sticky="ew")
        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.load_selected_partner)
        self.bind_double_click(self.tree_frame, self.columns, lambda row: self.open_partner_detail(row.get("codigo")))

    def selected_kinds(self):
        selected = [label for label, var in self.kind_vars.items() if var.get()]
        return "; ".join(selected) if selected else "Outro"

    def clear_form(self):
        self.selected_partner_code = None
        self.clear_entries(
            self.documento, self.nome_razao, self.nome_fantasia,
            self.inscricao_estadual, self.inscricao_municipal,
            self.telefone, self.celular, self.email, self.endereco,
            self.numero, self.bairro, self.cidade, self.estado,
            self.cep, self.responsavel, self.obs
        )
        self.tipo_pessoa.set("Jurídica")
        self.status.set("Ativo")
        self.pais.delete(0, "end")
        self.pais.insert(0, "Brasil")

        for label, var in self.kind_vars.items():
            var.set(label in ("Cliente", "Fornecedor"))

        self.save_button.configure(text="Salvar parceiro")

    def load_selected_partner(self, event=None):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            return

        partner = get_partner_by_code(row.get("codigo"))

        if not partner:
            return

        self.selected_partner_code = partner.get("codigo")
        self.set_combo(self.tipo_pessoa, partner.get("tipo_pessoa"))
        self.set_entry(self.documento, partner.get("documento"))
        self.set_entry(self.nome_razao, partner.get("nome_razao"))
        self.set_entry(self.nome_fantasia, partner.get("nome_fantasia"))
        self.set_entry(self.inscricao_estadual, partner.get("inscricao_estadual"))
        self.set_entry(self.inscricao_municipal, partner.get("inscricao_municipal"))
        self.set_entry(self.telefone, partner.get("telefone"))
        self.set_entry(self.celular, partner.get("celular"))
        self.set_entry(self.email, partner.get("email"))
        self.set_entry(self.endereco, partner.get("endereco"))
        self.set_entry(self.numero, partner.get("numero"))
        self.set_entry(self.bairro, partner.get("bairro"))
        self.set_entry(self.cidade, partner.get("cidade"))
        self.set_entry(self.estado, partner.get("estado"))
        self.set_entry(self.cep, partner.get("cep"))
        self.set_entry(self.pais, partner.get("pais") or "Brasil")
        self.set_entry(self.responsavel, partner.get("responsavel"))
        self.set_combo(self.status, partner.get("status"))
        self.set_entry(self.obs, partner.get("observacoes"))

        kinds = [kind.strip() for kind in str(partner.get("tipo_parceiro", "")).split(";")]

        for label, var in self.kind_vars.items():
            var.set(label in kinds)

        self.save_button.configure(text="Atualizar parceiro")

    def save(self):
        try:
            if self.selected_partner_code:
                codigo = update_partner(
                    self.selected_partner_code, self.tipo_pessoa.get(), self.documento.get(), self.nome_razao.get(), self.nome_fantasia.get(),
                    self.inscricao_estadual.get(), self.inscricao_municipal.get(), self.telefone.get(), self.celular.get(),
                    self.email.get(), self.endereco.get(), self.numero.get(), self.bairro.get(), self.cidade.get(),
                    self.estado.get(), self.cep.get(), self.pais.get(), self.responsavel.get(), self.selected_kinds(),
                    self.status.get(), self.obs.get()
                )
                messagebox.showinfo("Parceiros", f"Parceiro atualizado com sucesso.\nCódigo: {codigo}")
            else:
                codigo = create_partner(
                    self.tipo_pessoa.get(), self.documento.get(), self.nome_razao.get(), self.nome_fantasia.get(),
                    self.inscricao_estadual.get(), self.inscricao_municipal.get(), self.telefone.get(), self.celular.get(),
                    self.email.get(), self.endereco.get(), self.numero.get(), self.bairro.get(), self.cidade.get(),
                    self.estado.get(), self.cep.get(), self.pais.get(), self.responsavel.get(), self.selected_kinds(),
                    self.status.get(), self.obs.get()
                )
                messagebox.showinfo("Parceiros", f"Parceiro cadastrado com sucesso.\nCódigo: {codigo}")

            self.clear_form()
            self.refresh()
        except Exception as error:
            messagebox.showerror("Parceiros", str(error))

    def refresh(self):
        self.fill_tree(self.tree_frame, get_partners(), self.columns)

    def export_xlsx(self):
        export_xlsx("parceiros", self.columns, get_partners())

    def export_pdf(self):
        generate_table_pdf("parceiros", "Relatório de parceiros", self.columns, get_partners())

class CommercialDocumentsTab(BaseFrame):
    def __init__(self, master, title, direction, tipo_codes, partner_kind=None):
        super().__init__(master, padding=10)
        self.title = title
        self.direction = direction
        self.tipo_codes = tipo_codes
        self.partner_kind = partner_kind
        self.selected_document_record_id = None
        self.partner_options = []
        self.columns = ["id", "numero", "tipo_documento", "parceiro", "data_emissao", "data_validade", "prazo_entrega", "data_prevista_entrega", "status", "subtotal", "desconto_total", "acrescimo_total", "total", "fiscal_status"]
        self.payable_columns = ["numero_documento", "parceiro", "descricao", "data_emissao", "data_vencimento", "valor", "status"]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text=self.title, padding=10)
        form.pack(fill="x", pady=(0, 8))
        type_options = [
            option for option in get_document_type_options(self.direction)
            if option.split(" - ")[0] in self.tipo_codes
        ] or [""]
        self.make_label(form, "Tipo:", 0, 0)
        self.tipo = self.make_combo(form, 1, 0, type_options, width=36)
        self.make_label(form, "Parceiro:", 0, 1)
        self.partner = self.make_combo(form, 1, 1, get_partner_options(self.partner_kind), width=45)
        self.partner.configure(state="normal")
        self.partner.bind("<<ComboboxSelected>>", lambda event: self.load_partner_details())
        self.partner.bind("<KeyRelease>", self.filter_partner_options)
        self.partner.bind("<Return>", lambda event: self.load_partner_details())
        self.partner.bind("<FocusOut>", lambda event: self.load_partner_details())
        self.make_label(form, "Emissão:", 0, 2)
        self.data_emissao = self.make_date_entry(form, 1, 2)
        self.data_emissao.insert(0, today_br())
        self.make_label(form, "Validade:", 0, 3)
        self.data_validade = self.make_date_entry(form, 1, 3)
        self.make_label(form, "Prazo entrega:", 2, 0)
        self.prazo_entrega = self.make_date_entry(form, 3, 0)
        self.make_label(form, "Entrega prevista:", 2, 1)
        self.data_prevista_entrega = self.make_date_entry(form, 3, 1)
        self.make_label(form, "Responsável:", 2, 2)
        self.responsavel = self.make_entry(form, 3, 2)
        self.make_label(form, "Status:", 2, 3)
        self.status = self.make_combo(form, 3, 3, self.status_values())
        self.make_label(form, "Condição pagamento:", 4, 0)
        self.condicao = self.make_entry(form, 5, 0)
        self.make_label(form, "Forma pagamento:", 4, 1)
        self.forma = self.make_entry(form, 5, 1)
        self.make_label(form, "Produto:", 4, 2)
        self.product = self.make_combo(form, 5, 2, get_product_options(), width=45)
        self.product.bind("<<ComboboxSelected>>", lambda event: self.load_product_defaults())
        self.make_label(form, "Quantidade:", 4, 3)
        self.quantidade = self.make_entry(form, 5, 3)
        self.make_label(form, "Valor unitário:", 6, 0)
        self.valor_unitario = self.make_entry(form, 7, 0)
        self.make_label(form, "Desconto:", 6, 1)
        self.desconto = self.make_entry(form, 7, 1)
        self.make_label(form, "Acréscimo:", 6, 2)
        self.acrescimo = self.make_entry(form, 7, 2)
        self.make_label(form, "Observações:", 6, 3)
        self.obs = self.make_entry(form, 7, 3)

        action_row = 8

        if self.direction == "compra":
            self.make_label(form, "Vencimento conta:", 8, 0)
            self.payable_due_date = self.make_date_entry(form, 9, 0)
            self.make_label(form, "Status financeiro:", 8, 1)
            self.payable_status = self.make_combo(form, 9, 1, ["Em aberto", "Pago", "Vencido", "Cancelado"])
            action_row = 10
        else:
            self.payable_due_date = None
            self.payable_status = None

        self.save_button = ttk.Button(form, text="Salvar documento", style="Primary.TButton", command=self.save)
        self.save_button.grid(row=action_row, column=0, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Novo documento", command=self.clear_form).grid(row=action_row, column=1, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="PDF do selecionado", command=self.generate_selected_pdf).grid(row=action_row, column=2, padx=6, pady=4, sticky="ew")
        ttk.Button(form, text="Exportar XLSX", command=self.export_xlsx).grid(row=action_row, column=3, padx=6, pady=4, sticky="ew")

        self.partner_info = ttk.Label(self, text="", style="Muted.TLabel", wraplength=980)
        self.partner_info.pack(anchor="w", pady=(0, 4))
        self.product_info = ttk.Label(self, text="", style="Muted.TLabel")
        self.product_info.pack(anchor="w", pady=(0, 8))
        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.tree_frame.tree.bind("<<TreeviewSelect>>", self.load_selected_document)
        self.bind_double_click(self.tree_frame, self.columns, lambda row: self.open_document_detail(row.get("id")))

        if self.direction == "compra":
            payable_frame = ttk.LabelFrame(self, text="Contas a pagar", padding=8)
            payable_frame.pack(fill="both", expand=True, pady=(8, 0))
            self.payable_tree_frame = self.create_tree(payable_frame, self.payable_columns)
            self.payable_tree_frame.pack(fill="both", expand=True)
        else:
            self.payable_tree_frame = None

    def status_values(self):
        if self.tipo_codes == ["FT"]:
            return ["Pendente", "Faturado internamente", "Cancelado"]
        if self.direction == "compra":
            return ["Rascunho", "Em orçamento", "Aguardando aprovação", "Aprovado", "Convertido em ordem de compra", "Aberta", "Enviada ao fornecedor", "Parcialmente recebida", "Recebida", "Cancelado"]
        return ["Rascunho", "Orçamento", "Pedido aberto", "Aguardando faturamento", "Faturado", "Cancelado"]

    def filter_partner_options(self, event=None):
        if event and event.keysym in ("Return", "Escape", "Up", "Down", "Left", "Right", "Tab"):
            return

        typed = self.partner.get().strip().lower()
        options = self.partner_options or get_partner_options(self.partner_kind)

        if typed:
            filtered = [option for option in options if typed in option.lower()]
        else:
            filtered = options

        self.partner["values"] = filtered or options

    def resolve_partner(self, text):
        text = (text or "").strip()

        if not text:
            return None

        partner = get_partner_by_code(option_code(text))

        if partner:
            return partner

        text_lower = text.lower()

        for partner in get_partners(self.partner_kind):
            fields = [
                partner.get("nome_razao", ""),
                partner.get("nome_fantasia", ""),
                partner.get("documento", ""),
                partner.get("codigo", ""),
            ]

            if any(text_lower == str(field).lower() or text_lower in str(field).lower() for field in fields if field):
                return partner

        return None

    def load_partner_details(self):
        text = self.partner.get().strip()

        if not text:
            self.partner_info.configure(text="")
            return

        partner = self.resolve_partner(text)

        if not partner:
            self.partner_info.configure(text="")
            return

        option = self.partner_option(partner.get("codigo"), partner.get("nome_razao"))

        if option:
            self.partner.set(option)

        info = [
            partner.get("nome_razao", ""),
            f"CPF/CNPJ: {partner.get('documento', '')}",
            f"Tipo: {partner.get('tipo_parceiro', '')}",
            f"IE: {partner.get('inscricao_estadual', '')}",
            f"Telefone: {partner.get('telefone', '')}",
            f"E-mail: {partner.get('email', '')}",
            f"Endereço: {partner.get('endereco', '')}, {partner.get('numero', '')} - {partner.get('cidade', '')}/{partner.get('estado', '')}",
        ]
        self.partner_info.configure(text=" | ".join([value for value in info if value and not value.endswith(": ")]))

    def load_product_defaults(self):
        text = self.product.get().strip()

        if not text:
            return

        product = get_product_by_code(text.split(" - ")[0])

        if not product:
            return

        default_price = product.get("valor_custo") if self.direction == "compra" else product.get("valor_venda")
        self.valor_unitario.delete(0, "end")
        self.valor_unitario.insert(0, format_money(default_price))
        self.product_info.configure(text=f"{product.get('codigo')} | {product.get('descricao', '')} | Unidade: {product.get('unidade', '')} | Tipo: {product.get('tipo_produto', '')} | NCM: {product.get('ncm', '')}")

    def clear_form(self):
        self.selected_document_record_id = None
        self.clear_entries(
            self.data_validade, self.prazo_entrega, self.data_prevista_entrega,
            self.responsavel, self.condicao, self.forma, self.quantidade,
            self.valor_unitario, self.desconto, self.acrescimo, self.obs
        )
        self.data_emissao.delete(0, "end")
        self.data_emissao.insert(0, today_br())
        self.partner_info.configure(text="")
        self.product_info.configure(text="")
        self.save_button.configure(text="Salvar documento")

        if self.payable_due_date:
            self.payable_due_date.delete(0, "end")

        if self.payable_status:
            self.payable_status.set("Em aberto")

        type_values = list(self.tipo["values"])
        partner_values = list(self.partner["values"])
        product_values = list(self.product["values"])

        if type_values:
            self.tipo.set(type_values[0])
        if partner_values:
            self.partner.set(partner_values[0])
        if product_values:
            self.product.set(product_values[0])

        status_values = self.status_values()
        if status_values:
            self.status.set(status_values[0])

    def load_selected_document(self, event=None):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            return

        document, items = get_document_by_id(row.get("id"))

        if not document:
            return

        first_item = items[0] if items else {}
        self.selected_document_record_id = document.get("id")
        self.set_combo(self.tipo, self.option_by_code(self.tipo, document.get("tipo_code")))
        self.set_combo(self.partner, self.partner_option(document.get("partner_code"), document.get("parceiro")))
        self.set_entry(self.data_emissao, format_date_br(document.get("data_emissao")))
        self.set_entry(self.data_validade, format_date_br(document.get("data_validade")))
        self.set_entry(self.prazo_entrega, format_date_br(document.get("prazo_entrega")))
        self.set_entry(self.data_prevista_entrega, format_date_br(document.get("data_prevista_entrega")))
        self.set_entry(self.responsavel, document.get("vendedor_responsavel"))
        self.set_combo(self.status, document.get("status"))
        self.set_entry(self.condicao, document.get("condicao_pagamento"))
        self.set_entry(self.forma, document.get("forma_pagamento"))
        self.set_combo(self.product, self.option_by_code(self.product, first_item.get("product_code")))
        self.set_entry(self.quantidade, format_money(first_item.get("quantidade")))
        self.set_entry(self.valor_unitario, format_money(first_item.get("valor_unitario")))
        self.set_entry(self.desconto, format_money(first_item.get("desconto")))
        self.set_entry(self.acrescimo, format_money(first_item.get("acrescimo")))
        self.set_entry(self.obs, document.get("observacoes"))
        self.load_partner_details()
        product = get_product_by_code(first_item.get("product_code"))

        if product:
            self.product_info.configure(text=f"{product.get('codigo')} | {product.get('descricao', '')} | Unidade: {product.get('unidade', '')} | Tipo: {product.get('tipo_produto', '')} | NCM: {product.get('ncm', '')}")

        if self.direction == "compra" and self.payable_due_date and self.payable_status:
            payable = get_account_payable_by_document_id(document.get("id"))

            if payable:
                self.set_entry(self.payable_due_date, format_date_br(payable.get("data_vencimento")))
                self.set_combo(self.payable_status, payable.get("status"))
            else:
                self.set_entry(self.payable_due_date, format_date_br(document.get("data_validade") or document.get("data_prevista_entrega") or document.get("data_emissao")))
                self.set_combo(self.payable_status, "Em aberto")

        self.save_button.configure(text="Atualizar documento")

    def option_by_code(self, combo, code):
        code = str(code or "")

        for option in combo["values"]:
            if option.startswith(f"{code} - ") or option.endswith(f"({code})"):
                return option

        return ""

    def partner_option(self, code, name):
        if not code:
            return ""

        for option in self.partner["values"]:
            if option.endswith(f"({code})"):
                return option

        return f"{name or code} ({code})"

    def save(self):
        try:
            tipo_code = self.tipo.get().split(" - ")[0] if self.tipo.get() else ""
            partner = self.resolve_partner(self.partner.get())
            partner_code = partner.get("codigo") if partner else option_code(self.partner.get())
            product_code = self.product.get().split(" - ")[0] if self.product.get() else ""

            if self.selected_document_record_id:
                document_id = update_document(
                    self.selected_document_record_id, tipo_code, partner_code, self.data_emissao.get(), self.data_validade.get(),
                    self.prazo_entrega.get(), self.data_prevista_entrega.get(), self.responsavel.get(),
                    self.condicao.get(), self.forma.get(), self.status.get(), self.obs.get(),
                    product_code, self.quantidade.get(), self.valor_unitario.get(), self.desconto.get(), self.acrescimo.get()
                )
                message = "Documento atualizado com sucesso."
            else:
                document_id, numero = create_document(
                    tipo_code, partner_code, self.data_emissao.get(), self.data_validade.get(),
                    self.prazo_entrega.get(), self.data_prevista_entrega.get(), self.responsavel.get(),
                    self.condicao.get(), self.forma.get(), self.status.get(), self.obs.get(),
                    product_code, self.quantidade.get(), self.valor_unitario.get(), self.desconto.get(), self.acrescimo.get()
                )
                message = f"Documento salvo com sucesso.\nNúmero: {numero}"

            if messagebox.askyesno(self.title, f"{message}\n\nDeseja gerar o PDF agora?"):
                generate_document_pdf(document_id)

            if self.direction == "compra" and self.payable_due_date and self.payable_status:
                upsert_account_payable_for_document(document_id, self.payable_due_date.get(), self.payable_status.get())

            self.clear_form()
            self.refresh()
        except Exception as error:
            messagebox.showerror(self.title, str(error))

    def refresh(self):
        self.partner_options = get_partner_options(self.partner_kind)
        self.partner["values"] = self.partner_options
        self.product["values"] = get_product_options()
        self.fill_tree(self.tree_frame, get_documents(self.direction, self.tipo_codes), self.columns)

        if self.payable_tree_frame:
            self.fill_tree(self.payable_tree_frame, get_accounts_payable(), self.payable_columns)

    def get_selected_document_id(self):
        row = self.get_selected_row(self.tree_frame, self.columns)

        if not row:
            messagebox.showwarning(self.title, "Selecione um documento na tabela.")
            return None

        return row.get("id")

    def generate_selected_pdf(self):
        document_id = self.get_selected_document_id()
        if document_id:
            generate_document_pdf(document_id)

    def export_xlsx(self):
        export_xlsx(self.title.lower().replace(" ", "_"), self.columns, get_documents(self.direction, self.tipo_codes))

class PurchasesTab(CommercialDocumentsTab):
    def __init__(self, master):
        super().__init__(master, "Compras", "compra", ["OCM", "OC"], "Fornecedor")

class SalesTab(CommercialDocumentsTab):
    def __init__(self, master):
        super().__init__(master, "Vendas", "venda", ["OV", "PV"], "Cliente")

class BillingTab(CommercialDocumentsTab):
    def __init__(self, master):
        super().__init__(master, "Faturamento", "venda", ["FT"], "Cliente")

class BusinessMovementsTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = ["id", "product_code", "numero_documento", "tipo", "quantidade", "data_movimentacao", "parceiro_nome", "observacoes"]
        self.build()
        self.refresh()

    def build(self):
        form = ttk.LabelFrame(self, text="Movimentações empresariais", padding=10)
        form.pack(fill="x", pady=(0, 8))

        for column in range(4):
            form.columnconfigure(column, weight=1, uniform="movement_form")

        self.make_label(form, "Produto:", 0, 0)
        self.product = self.make_combo(form, 1, 0, get_product_options(), width=28)
        self.make_label(form, "Tipo:", 0, 1)
        self.tipo = self.make_combo(form, 1, 1, ["Entrada de compra", "Saída por venda", "Ajuste de estoque", "Devolução de compra", "Devolução de venda", "Cancelamento", "Baixa manual", "Movimentação interna"], width=24)
        self.make_label(form, "Quantidade:", 0, 2)
        self.quantidade = self.make_entry(form, 1, 2, width=16)
        self.make_label(form, "Data:", 0, 3)
        self.data = self.make_date_entry(form, 1, 3, width=16)
        self.data.insert(0, today_br())
        self.make_label(form, "Documento:", 2, 0)
        self.document_number = self.make_entry(form, 3, 0, width=20)
        self.make_label(form, "Parceiro:", 2, 1)
        self.partner = self.make_combo(form, 3, 1, get_partner_options(), width=28)
        self.make_label(form, "Observações:", 2, 2)
        self.obs = self.make_entry(form, 3, 2, width=24)
        ModernButton(form, text="✓  Registrar", command=self.save).grid(row=3, column=3, padx=6, pady=(0, 6), sticky="ew")
        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)

    def save(self):
        try:
            product_code = self.product.get().split(" - ")[0] if self.product.get() else ""
            partner_code = option_code(self.partner.get())
            partner = get_partner_by_code(partner_code) if partner_code else None
            quantity = parse_float_br(self.quantidade.get())
            movement_date = validate_date(self.data.get(), "data", allow_empty=False)

            if not product_code:
                raise ValueError("Selecione um produto.")

            if quantity <= 0:
                raise ValueError("Quantidade deve ser maior que zero.")

            conn = get_conn()
            conn.execute("""
                INSERT INTO business_movements (
                    product_code, numero_documento, tipo, quantidade, data_movimentacao,
                    parceiro_codigo, parceiro_nome, observacoes, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                product_code,
                self.document_number.get().strip(),
                self.tipo.get(),
                quantity,
                movement_date,
                partner_code,
                partner.get("nome_razao", "") if partner else "",
                self.obs.get().strip(),
                now_iso()
            ))
            conn.commit()
            conn.close()
            messagebox.showinfo("Movimentações", "Movimentação registrada com sucesso.")
            self.clear_entries(self.quantidade, self.document_number, self.obs)
            self.refresh()
        except Exception as error:
            messagebox.showerror("Movimentações", str(error))

    def refresh(self):
        self.product["values"] = get_product_options()
        self.partner["values"] = get_partner_options()
        conn = get_conn()
        rows = conn.execute("""
            SELECT *
            FROM business_movements
            ORDER BY data_movimentacao DESC, id DESC
        """).fetchall()
        conn.close()
        movements = []

        for row in rows:
            movement = dict(row)
            movement["_data_movimentacao_iso"] = movement.get("data_movimentacao", "")
            movement["data_movimentacao"] = format_date_br(movement.get("data_movimentacao", ""))
            movements.append(movement)

        self.fill_tree(self.tree_frame, movements, self.columns)

class BusinessReportsTab(BaseFrame):
    def __init__(self, master):
        super().__init__(master, padding=10)
        self.columns = ["relatorio", "status", "xlsx", "pdf"]
        self.build()

    def build(self):
        ttk.Label(self, text="Relatórios gerenciais", font=("Segoe UI", 14, "bold"), style="Title.TLabel").pack(anchor="w", pady=(0, 8))
        button_area = ttk.Frame(self)
        button_area.pack(fill="x", pady=(0, 8))
        primary_buttons = ttk.Frame(button_area)
        primary_buttons.pack(fill="x", pady=(0, 6))
        alert_buttons = ttk.Frame(button_area)
        alert_buttons.pack(fill="x")
        ttk.Button(primary_buttons, text="Produtos XLSX", command=lambda: export_xlsx("produtos", ["codigo", "nome", "categoria", "tipo_produto", "unidade", "ncm", "valor_custo", "valor_venda", "estoque_atual", "estoque_minimo", "status"], get_products())).pack(side="left", padx=4)
        ttk.Button(primary_buttons, text="Parceiros XLSX", command=lambda: export_xlsx("parceiros", ["codigo", "tipo_pessoa", "documento", "nome_razao", "nome_fantasia", "tipo_parceiro", "cidade", "estado", "status"], get_partners())).pack(side="left", padx=4)
        ttk.Button(primary_buttons, text="Compras XLSX", command=lambda: export_xlsx("compras", ["id", "numero", "tipo_documento", "parceiro", "data_emissao", "status", "total"], get_documents("compra", ["OCM", "OC"]))).pack(side="left", padx=4)
        ttk.Button(primary_buttons, text="Vendas XLSX", command=lambda: export_xlsx("vendas", ["id", "numero", "tipo_documento", "parceiro", "data_emissao", "status", "total"], get_documents("venda", ["OV", "PV"]))).pack(side="left", padx=4)
        ttk.Button(primary_buttons, text="Faturamentos XLSX", command=lambda: export_xlsx("faturamentos", ["id", "numero", "tipo_documento", "parceiro", "data_emissao", "status", "total"], get_documents("venda", ["FT"]))).pack(side="left", padx=4)
        ttk.Button(primary_buttons, text="Produtos PDF", command=lambda: generate_table_pdf("produtos", "Relatório de produtos", ["codigo", "nome", "categoria", "tipo_produto", "unidade", "ncm", "valor_venda", "estoque_atual", "status"], get_products())).pack(side="left", padx=4)
        ttk.Button(primary_buttons, text="Parceiros PDF", command=lambda: generate_table_pdf("parceiros", "Relatório de parceiros", ["codigo", "tipo_pessoa", "documento", "nome_razao", "tipo_parceiro", "cidade", "estado", "status"], get_partners())).pack(side="left", padx=4)
        ttk.Button(alert_buttons, text="Pendentes XLSX", command=lambda: export_xlsx("documentos_pendentes", PENDING_DOCUMENT_COLUMNS, get_business_pending_documents_report())).pack(side="left", padx=4)
        ttk.Button(alert_buttons, text="Pendentes PDF", command=lambda: generate_table_pdf("documentos_pendentes", "Documentos pendentes", PENDING_DOCUMENT_COLUMNS, get_business_pending_documents_report())).pack(side="left", padx=4)
        ttk.Button(alert_buttons, text="Estoque baixo XLSX", command=lambda: export_xlsx("estoque_baixo", LOW_STOCK_COLUMNS, get_business_low_stock_report())).pack(side="left", padx=4)
        ttk.Button(alert_buttons, text="Estoque baixo PDF", command=lambda: generate_table_pdf("estoque_baixo", "Estoque baixo", LOW_STOCK_COLUMNS, get_business_low_stock_report())).pack(side="left", padx=4)

        rows = [
            {"relatorio": "Produtos cadastrados", "status": "Disponível", "xlsx": "Sim", "pdf": "Preparado por documento"},
            {"relatorio": "Parceiros cadastrados", "status": "Disponível", "xlsx": "Sim", "pdf": "Preparado"},
            {"relatorio": "Compras por período", "status": "Base criada", "xlsx": "Sim", "pdf": "Documento"},
            {"relatorio": "Vendas por período", "status": "Base criada", "xlsx": "Sim", "pdf": "Documento"},
            {"relatorio": "Faturamento por período", "status": "Base criada", "xlsx": "Sim", "pdf": "Documento"},
            {"relatorio": "Documentos pendentes", "status": "Disponível", "xlsx": "Sim", "pdf": "Sim"},
            {"relatorio": "Estoque baixo", "status": "Disponível", "xlsx": "Sim", "pdf": "Sim"},
        ]
        self.tree_frame = self.create_tree(self, self.columns)
        self.tree_frame.pack(fill="both", expand=True)
        self.fill_tree(self.tree_frame, rows, self.columns)

class ProfileTab(BaseFrame):
    def __init__(self, master, user):
        super().__init__(master, padding=20)
        self.user = user
        self.build()

    def build(self):
        ttk.Label(
            self,
            text="Meu Perfil",
            font=("Segoe UI", 16, "bold")
        ).pack(anchor="w", pady=(0, 12))

        info = [
            ("Usuário", self.user.get("username", "")),
            ("E-mail", self.user.get("email", "")),
            ("Perfil", self.user.get("role", "")),
            ("Status", self.user.get("status", "")),
            ("Código do cliente", self.user.get("client_code", "")),
            ("Criado em", self.user.get("created_at", "")),
        ]

        for label, value in info:
            ttk.Label(
                self,
                text=f"{label}: {value}"
            ).pack(anchor="w", pady=3)

class ProfileEditTab(BaseFrame):
    def __init__(self, master, user):
        super().__init__(master, padding=20)
        self.user = user
        self.photo_source = ""
        self.photo_preview = None
        self.build()

    def build(self):
        ttk.Label(
            self,
            text="Editar perfil",
            font=("Segoe UI", 16, "bold"),
            style="Title.TLabel"
        ).pack(anchor="w", pady=(0, 12))

        form = ttk.LabelFrame(self, text="Dados do perfil", padding=14)
        form.pack(fill="x")

        photo_box = ttk.Frame(form)
        photo_box.grid(row=0, column=0, rowspan=6, padx=(0, 18), sticky="n")
        self.photo_label = ttk.Label(photo_box, text="Sem foto", style="Muted.TLabel")
        self.photo_label.pack(pady=(0, 8))
        ttk.Button(photo_box, text="Escolher foto", command=self.choose_photo).pack(fill="x")

        self.refresh_photo()

        self.make_label(form, "Nome:", 0, 1)
        self.display_name = self.make_entry(form, 1, 1, width=42)
        self.set_entry(self.display_name, self.user.get("display_name") or self.user.get("username"))

        self.make_label(form, "Login:", 2, 1)
        self.username = self.make_entry(form, 3, 1, width=42)
        self.set_entry(self.username, self.user.get("username"))

        self.make_label(form, "E-mail:", 4, 1)
        self.email = self.make_entry(form, 5, 1, width=42)
        self.set_entry(self.email, self.user.get("email"))

        self.make_label(form, "Telefone:", 0, 2)
        self.phone = self.make_entry(form, 1, 2, width=28)
        self.set_entry(self.phone, self.user.get("phone"))

        self.make_label(form, "Documento:", 2, 2)
        self.document = self.make_entry(form, 3, 2, width=28)
        self.set_entry(self.document, self.user.get("document"))

        ModernButton(form, text="✓  Salvar perfil", command=self.save).grid(row=5, column=2, padx=6, pady=(0, 6), sticky="ew")

    def refresh_photo(self):
        self.photo_preview = load_tk_image(self.photo_source or self.user.get("profile_photo"), (96, 96))

        if self.photo_preview:
            self.photo_label.configure(image=self.photo_preview, text="")
        else:
            self.photo_label.configure(image="", text="Sem foto")

    def choose_photo(self):
        path = filedialog.askopenfilename(
            title="Selecionar foto de perfil",
            filetypes=[("Imagens", "*.png;*.jpg;*.jpeg"), ("Todos os arquivos", "*.*")]
        )

        if path:
            self.photo_source = path
            self.refresh_photo()

    def save(self):
        try:
            updated = update_user_profile(
                self.user.get("id"),
                self.username.get(),
                self.email.get(),
                self.display_name.get(),
                self.phone.get(),
                self.document.get(),
                self.photo_source,
                self.user.get("profile_photo", "")
            )

            self.user.clear()
            self.user.update(updated)
            self.photo_source = ""
            self.refresh_photo()
            messagebox.showinfo("Perfil", "Perfil atualizado com sucesso.")
        except Exception as error:
            messagebox.showerror("Perfil", str(error))

class SystemTab(BaseFrame):
    def __init__(self, master, app, current_user):
        super().__init__(master, padding=20)
        self.app = app
        self.current_user = current_user or {}
        self.build()

    def build(self):
        ttk.Label(
            self,
            text="Sistema",
            font=("Segoe UI", 16, "bold"),
            style="Title.TLabel"
        ).pack(anchor="w", pady=(0, 12))

        box = ttk.LabelFrame(self, text="Banco de dados", padding=14)
        box.pack(fill="x")

        ttk.Label(
            box,
            text="Salva uma cópia do banco atual e recria o sistema limpo. Todos os cadastros, documentos, produtos, parceiros, usuários e permissões atuais serão removidos do banco em uso.",
            style="Muted.TLabel",
            wraplength=760
        ).pack(anchor="w", pady=(0, 12))

        ModernButton(
            box,
            text="Apagar banco atual",
            command=self.reset_database,
            variant="danger",
            width=180
        ).pack(anchor="w")

    def reset_database(self):
        if not user_can_access_module(self.current_user, "sistema"):
            messagebox.showerror("Sistema", "Seu perfil não tem permissão para acessar esta ação.")
            return

        confirmed = messagebox.askyesno(
            "Sistema",
            "Esta ação salva um backup e depois apaga o banco atual em uso.\n\nDepois do reset, o sistema volta limpo com o usuário admin padrão.\n\nDeseja continuar?"
        )

        if not confirmed:
            return

        ensure_app_dirs()
        backup_name = f"backup_grion_criteria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
        path = filedialog.asksaveasfilename(
            defaultextension=".db",
            filetypes=[("Banco SQLite", "*.db")],
            initialdir=str(BACKUPS_DIR),
            initialfile=backup_name,
            title="Salvar backup antes de apagar"
        )

        if not path:
            return

        try:
            if Path(path).resolve() == DB_PATH.resolve():
                raise ValueError("Escolha um arquivo de backup diferente do banco em uso.")

            backup_path = save_database_backup(path)
        except Exception as error:
            messagebox.showerror("Backup", f"Não foi possível salvar o backup.\n\n{error}")
            return

        password = simpledialog.askstring(
            "Senha do sistema",
            "Backup salvo. Digite a senha do sistema para apagar o banco atual:",
            show="*",
            parent=self.winfo_toplevel()
        )

        if password != "grion123":
            messagebox.showerror("Sistema", f"Senha incorreta. O banco não foi apagado.\n\nBackup preservado em:\n{backup_path}")
            return

        try:
            reset_database_to_factory()
        except Exception as error:
            messagebox.showerror("Sistema", f"Backup salvo, mas não foi possível resetar o banco.\n\n{error}")
            return

        messagebox.showinfo("Sistema", f"Banco atual apagado e recriado com sucesso.\n\nBackup salvo em:\n{backup_path}")
        self.app.show_login()

class SettingsTab(BaseFrame):
    def __init__(self, master, current_user=None):
        super().__init__(master, padding=20)
        self.current_user = current_user or {}
        self.build()

    def build(self):
        ttk.Label(
            self,
            text="Configurações",
            font=("Segoe UI", 16, "bold")
        ).pack(anchor="w", pady=(0, 12))

        ttk.Button(
            self,
            text="Fazer backup do banco de dados",
            command=self.backup_database
        ).pack(anchor="w", pady=4)

        ttk.Button(
            self,
            text="Exportar clientes CSV",
            command=self.export_clients
        ).pack(anchor="w", pady=4)

        ttk.Button(
            self,
            text="Exportar itens CSV",
            command=self.export_items
        ).pack(anchor="w", pady=4)

        ttk.Button(
            self,
            text="Exportar gastos de EPI CSV",
            command=self.export_epi
        ).pack(anchor="w", pady=4)

        ttk.Button(
            self,
            text="Exportar alertas de validade CSV",
            command=self.export_validity
        ).pack(anchor="w", pady=4)

        actions = ttk.LabelFrame(self, text="Manutenção do sistema", padding=10)
        actions.pack(fill="x", pady=(18, 0))

        ttk.Button(
            actions,
            text="Limpar histórico e dados operacionais",
            command=self.clear_all_data
        ).pack(anchor="w", pady=4)

        ttk.Label(
            actions,
            text="Essa ação mantém usuários, senhas e permissões, mas remove produtos, parceiros, documentos, movimentações e cadastros operacionais.",
            style="Muted.TLabel",
            wraplength=760
        ).pack(anchor="w", pady=(4, 0))

    def backup_database(self):
        ensure_app_dirs()
        path = filedialog.asksaveasfilename(
            defaultextension=".db",
            filetypes=[("Banco SQLite", "*.db")],
            initialdir=str(BACKUPS_DIR),
            initialfile="backup_grion_criteria.db"
        )

        if not path:
            return

        shutil.copy2(DB_PATH, path)
        messagebox.showinfo("Backup", "Backup criado com sucesso.")

    def export_clients(self):
        columns = [
            "codigo",
            "tipo_pessoa",
            "nome_razao",
            "documento",
            "telefone",
            "email",
            "responsavel",
            "username",
            "status"
        ]

        export_csv("clientes", columns, get_clients())

    def export_items(self):
        columns = [
            "codigo",
            "nome",
            "categoria",
            "unidade",
            "saldo",
            "ca",
            "valor_unitario",
            "validade",
            "imagem",
            "image_path",
            "observacoes"
        ]

        export_csv("itens", columns, get_items())

    def export_epi(self):
        columns = [
            "funcionario_codigo",
            "funcionario_nome",
            "item_codigo",
            "item_nome",
            "quantidade",
            "valor_unitario",
            "valor_total",
            "data_movimentacao",
            "service_codigo",
            "destino",
            "observacoes"
        ]

        export_csv("gastos_epi_por_funcionario", columns, get_epi_expense_report())

    def export_validity(self):
        columns = [
            "origem",
            "responsavel",
            "item_codigo",
            "item_nome",
            "quantidade",
            "validade",
            "status",
            "nivel"
        ]

        export_csv("alertas_validade_itens_epi", columns, get_validity_alerts())

    def clear_all_data(self):
        if not self.current_user.get("id"):
            messagebox.showerror("Limpeza", "Não foi possível identificar o usuário logado.")
            return

        confirmed = messagebox.askyesno(
            "Limpar dados",
            "Esta ação remove produtos, parceiros, documentos, faturamentos, movimentações e cadastros operacionais.\n\nUsuários e permissões serão mantidos.\n\nDeseja continuar?"
        )

        if not confirmed:
            return

        password = simpledialog.askstring(
            "Confirmar senha",
            "Digite a senha da conta logada para apagar os dados:",
            show="*",
            parent=self.winfo_toplevel()
        )

        if not password:
            return

        if not verify_user_password(self.current_user.get("id"), password):
            messagebox.showerror("Limpeza", "Senha incorreta. Nenhum dado foi apagado.")
            return

        clear_operational_data()
        messagebox.showinfo("Limpeza", "Histórico e dados operacionais apagados com sucesso.")

class ModernMainFrame(BaseFrame):
    def __init__(self, master, app, user):
        super().__init__(master)
        self.app = app
        self.user = user
        self.content_host = None
        self.topbar_canvas = None
        self.topbar_menu = None
        self.sidebar_canvas = None
        logo_path = LOGO_PATH if LOGO_PATH.exists() else LEGACY_LOGO_PATH
        self.sidebar_logo_image = load_tk_image(str(logo_path), (102, 51)) if logo_path.exists() else None
        self.nav_items = []
        self.current_title = "Dashboard"
        self.active_button_key = None
        self.build()

    def build(self):
        self.configure(style="App.TFrame")
        self.columnconfigure(0, minsize=239)
        self.columnconfigure(1, weight=1)
        self.rowconfigure(0, minsize=76)
        self.rowconfigure(1, weight=1)

        self.build_topbar()
        modules = self.get_modules()

        sidebar = tk.Frame(self, bg=COLOR_SIDEBAR, width=239)
        sidebar.grid(row=1, column=0, sticky="ns")
        sidebar.grid_propagate(False)
        self.build_sidebar(sidebar, modules)

        main = tk.Frame(self, bg=COLOR_BG, bd=0, highlightthickness=0)
        main.grid(row=1, column=1, sticky="nsew")
        main.columnconfigure(0, weight=1)
        main.rowconfigure(0, weight=1)

        self.content_host = tk.Frame(main, bg=COLOR_BG, bd=0, highlightthickness=0)
        self.content_host.grid(row=0, column=0, sticky="nsew")

        if modules:
            first = modules[0]
            self.show_module(first[0], first[1], first[2], *first[3])

    def build_topbar(self):
        self.topbar_canvas = tk.Canvas(self, bg=COLOR_TOPBAR, height=76, highlightthickness=0, bd=0)
        self.topbar_canvas.grid(row=0, column=0, columnspan=2, sticky="ew")
        self.topbar_canvas.bind("<Configure>", self.redraw_topbar)
        self.build_topbar_menu()

    def menu_style(self):
        return {
            "tearoff": 0,
            "bg": COLOR_PANEL,
            "fg": COLOR_TEXT,
            "activebackground": COLOR_NAV_ACTIVE,
            "activeforeground": COLOR_ACCENT,
            "bd": 0,
            "relief": "flat",
            "font": ("Segoe UI", 10),
        }

    def build_topbar_menu(self):
        self.topbar_menu = tk.Menu(self, **self.menu_style())

    def open_topbar_menu(self, event):
        self.topbar_menu.delete(0, "end")
        modules_menu = tk.Menu(self.topbar_menu, **self.menu_style())

        for title, module_code, frame_class, args in self.get_modules():
            modules_menu.add_command(
                label=title,
                command=lambda item=(title, module_code, frame_class, args): self.show_module(item[0], item[1], item[2], *item[3])
            )

        config_menu = tk.Menu(self.topbar_menu, **self.menu_style())
        config_state = "normal" if user_can_access_module(self.user, "configuracoes") else "disabled"
        system_state = "normal" if user_can_access_module(self.user, "sistema") else "disabled"
        profile_state = "normal" if user_can_access_module(self.user, "perfil") else "disabled"

        self.topbar_menu.add_cascade(label="Módulos", menu=modules_menu)
        self.topbar_menu.add_separator()
        config_menu.add_command(label="Configurações", state=config_state, command=self.show_settings)
        config_menu.add_command(label="Sistema", state=system_state, command=self.show_system)
        self.topbar_menu.add_cascade(label="Configurações", menu=config_menu)
        self.topbar_menu.add_separator()
        self.topbar_menu.add_command(label="Editar perfil", state=profile_state, command=self.show_profile_editor)

        try:
            self.topbar_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.topbar_menu.grab_release()

    def redraw_topbar(self, event=None):
        canvas = self.topbar_canvas
        canvas.delete("all")
        pill_width = 330
        pill_height = 48
        width = max(canvas.winfo_width(), 981)
        x1 = 370 if width == 981 else (width - pill_width) / 2
        y1 = 10

        if self.sidebar_logo_image:
            canvas.create_image(14, 12, image=self.sidebar_logo_image, anchor="nw")
        else:
            canvas.create_text(26, 38, text="grION", anchor="w", fill=COLOR_ACCENT, font=("Segoe UI", 15, "bold"))

        draw_round_rect(canvas, x1, y1, x1 + pill_width, y1 + pill_height, 5, COLOR_SURFACE, COLOR_SURFACE, 1)
        canvas.create_text(x1 + pill_width / 2, y1 + 24, text=self.current_title, fill=COLOR_MUTED, font=("Segoe UI", 16))
        canvas.create_line(0, 75, width, 75, fill=COLOR_BORDER, width=1)

    def show_settings(self):
        self.show_module("Configurações", "configuracoes", SettingsTab, self.user)

    def show_system(self):
        self.show_module("Sistema", "sistema", SystemTab, self.app, self.user)

    def show_profile_editor(self):
        self.show_module("Editar perfil", "perfil", ProfileEditTab, self.user)

    def build_sidebar(self, sidebar, modules):
        self.nav_items = modules
        self.sidebar_canvas = tk.Canvas(sidebar, bg=COLOR_SIDEBAR, width=239, highlightthickness=0, bd=0)
        self.sidebar_canvas.pack(fill="both", expand=True)
        self.sidebar_canvas.bind("<Configure>", self.redraw_sidebar)
        self.sidebar_canvas.bind("<Button-1>", self.handle_sidebar_click)
        self.sidebar_canvas.bind("<Motion>", self.handle_sidebar_motion)

    def handle_sidebar_click(self, event):
        if 25 <= event.x <= 61 and 18 <= event.y <= 58:
            self.open_topbar_menu(event)

    def handle_sidebar_motion(self, event):
        self.sidebar_canvas.configure(cursor="hand2" if 25 <= event.x <= 61 and 18 <= event.y <= 58 else "")

    def button_key(self, title, module_code):
        return f"{module_code}:{title}"

    def nav_label(self, title):
        labels = {
            "Faturamento": "Faturam.",
            "Configurações": "Config.",
            "Meu Perfil": "Perfil",
        }
        return labels.get(title, title)

    def redraw_sidebar(self, event=None):
        canvas = self.sidebar_canvas
        canvas.delete("all")
        width = 239
        height = max(canvas.winfo_height(), 421)
        canvas.create_rectangle(0, 0, width, height, fill=COLOR_SIDEBAR, outline="")
        canvas.create_line(width - 1, 0, width - 1, height, fill=COLOR_BORDER, width=1)
        draw_round_rect(canvas, 25, 18, 61, 58, 9, COLOR_BUTTON, COLOR_BORDER, 1)

        for x in (36, 43, 50):
            canvas.create_oval(x - 2, 36, x + 2, 40, fill=COLOR_MUTED, outline="")

        for index, (title, module_code, frame_class, args) in enumerate(self.nav_items):
            key = self.button_key(title, module_code)
            y = 69 + index * 52

            if y + 48 > height:
                continue

            tag = f"nav_{index}"
            active = key == self.active_button_key
            color = COLOR_ACCENT if active else COLOR_MUTED

            if active:
                draw_round_rect(canvas, 13, y, 217, y + 48, 5, COLOR_NAV_ACTIVE, COLOR_NAV_ACTIVE, 1)

            self.draw_nav_icon(canvas, title, module_code, 28, y + 24, color, tag)
            canvas.create_text(58, y + 24, text=self.nav_label(title), anchor="w", fill=color, font=("Segoe UI", 15), tags=(tag,))
            canvas.tag_bind(tag, "<Button-1>", lambda _event, item=(title, module_code, frame_class, args): self.show_module(item[0], item[1], item[2], *item[3]))
            canvas.tag_bind(tag, "<Enter>", lambda _event: canvas.configure(cursor="hand2"))
            canvas.tag_bind(tag, "<Leave>", lambda _event: canvas.configure(cursor=""))

    def draw_nav_icon(self, canvas, title, module_code, x, y, color, tag):
        width = 1

        if module_code == "dashboard":
            for dx in (0, 9):
                for dy in (-9, 0):
                    canvas.create_rectangle(x + dx, y + dy, x + dx + 5, y + dy + 5, outline=color, width=width, tags=(tag,))
            return

        if title == "Clientes":
            canvas.create_oval(x, y - 10, x + 8, y - 2, outline=color, width=width, tags=(tag,))
            canvas.create_arc(x - 3, y - 2, x + 11, y + 13, start=25, extent=130, outline=color, width=width, style="arc", tags=(tag,))
            canvas.create_oval(x + 10, y - 8, x + 17, y - 1, outline=color, width=width, tags=(tag,))
            canvas.create_arc(x + 7, y, x + 20, y + 12, start=25, extent=130, outline=color, width=width, style="arc", tags=(tag,))
            return

        if title == "Estoque":
            canvas.create_polygon(x + 8, y - 11, x + 17, y - 6, x + 17, y + 5, x + 8, y + 11, x - 1, y + 5, x - 1, y - 6, outline=color, fill="", width=width, tags=(tag,))
            canvas.create_line(x - 1, y - 6, x + 8, y, x + 17, y - 6, fill=color, width=width, tags=(tag,))
            canvas.create_line(x + 8, y, x + 8, y + 11, fill=color, width=width, tags=(tag,))
            return

        if title == "Moviment.":
            canvas.create_line(x - 1, y - 5, x + 17, y - 5, fill=color, width=width, tags=(tag,))
            canvas.create_line(x + 13, y - 9, x + 17, y - 5, x + 13, y - 1, fill=color, width=width, tags=(tag,))
            canvas.create_line(x + 17, y + 6, x - 1, y + 6, fill=color, width=width, tags=(tag,))
            canvas.create_line(x + 3, y + 2, x - 1, y + 6, x + 3, y + 10, fill=color, width=width, tags=(tag,))
            return

        if title == "Funcionários":
            canvas.create_oval(x + 1, y - 10, x + 9, y - 2, outline=color, width=width, tags=(tag,))
            canvas.create_arc(x - 2, y - 2, x + 12, y + 13, start=25, extent=130, outline=color, width=width, style="arc", tags=(tag,))
            canvas.create_line(x + 15, y - 4, x + 15, y + 6, fill=color, width=width, tags=(tag,))
            canvas.create_line(x + 10, y + 1, x + 20, y + 1, fill=color, width=width, tags=(tag,))
            return

        if title == "Relatórios":
            canvas.create_line(x - 1, y + 10, x + 19, y + 10, fill=color, width=width, tags=(tag,))
            for dx, top in ((1, 2), (8, -5), (15, -10)):
                canvas.create_line(x + dx, y + 10, x + dx, y + top, fill=color, width=2, tags=(tag,))
            return

        canvas.create_oval(x + 4, y - 4, x + 12, y + 4, fill=color, outline="", tags=(tag,))

    def set_active_button(self, title, module_code):
        self.active_button_key = self.button_key(title, module_code)
        if self.sidebar_canvas:
            self.redraw_sidebar()

    def get_modules(self):
        if self.user.get("role") == "client":
            return [
                ("Dashboard", "dashboard", BusinessDashboardTab, ()),
                ("Vendas", "vendas", SalesTab, ()),
                ("Meu Perfil", "perfil", ProfileTab, (self.user,)),
            ]

        modules = [
            ("Dashboard", "dashboard", BusinessDashboardTab, ()),
            ("Clientes", "parceiros", PartnersTab, ()),
            ("Estoque", "produtos", ProductsTab, ()),
            ("Moviment.", "movimentacoes", BusinessMovementsTab, ()),
            ("Funcionários", "funcionarios", EmployeesTab, ()),
            ("Relatórios", "relatorios", BusinessReportsTab, ()),
            ("Cadastros", "cadastros_aux", AuxiliaryRegistersTab, (self.user,)),
            ("Comercial", "comercial", CommercialPipelineTab, (self.user,)),
            ("Propostas", "propostas", ProposalsTab, (self.user,)),
            ("Contratos", "contratos", ContractsTab, (self.user,)),
            ("Projetos", "projetos", ProjectsTab, (self.user,)),
            ("Escalas", "escalas", ScheduleTab, (self.user,)),
            ("Horas", "horas", TimeTrackingTab, (self.user,)),
            ("RH", "rh", HumanResourcesTab, (self.user,)),
            ("Viagens", "viagens", TravelTab, (self.user,)),
            ("Compras", "compras", PurchasesTab, ()),
            ("Recebimento", "recebimento", ReceivingTab, (self.user,)),
            ("Vendas", "vendas", SalesTab, ()),
            ("Faturamento", "faturamento", BillingTab, ()),
            ("Equipamentos", "equipamentos", EquipmentTab, (self.user,)),
            ("Financeiro", "financeiro", FinancialControlTab, (self.user,)),
            ("Fiscal", "fiscal", FiscalDocumentsTab, (self.user,)),
            ("Operacional", "operacional", OperationalReportsTab, (self.user,)),
            ("Acessos", "acessos", AccessControlTab, (self.user,)),
            ("Aprovação", "acessos", ApprovalTab, ()),
            ("Config.", "configuracoes", SettingsTab, (self.user,)),
        ]

        return [module for module in modules if user_can_access_module(self.user, module[1])]

    def show_module(self, title, module_code, frame_class, *args):
        self.current_title = self.nav_label(title)
        if self.topbar_canvas:
            self.redraw_topbar()
        self.set_active_button(title, module_code)

        for child in self.content_host.winfo_children():
            child.destroy()

        if frame_class is BusinessDashboardTab:
            content = frame_class(self.content_host, *args)
            content.pack(fill="both", expand=True)
            return

        scroll = ScrollableFrame(self.content_host)
        scroll.pack(fill="both", expand=True)
        content = frame_class(scroll.scrollable_frame, *args)
        content.pack(fill="both", expand=True)

class App(tk.Tk):
    def __init__(self):
        super().__init__()

        init_database()

        self.title(APP_NAME)
        self.geometry("981x497")
        self.minsize(981, 497)
        self.configure(background=COLOR_BG)

        self.current_frame = None
        self.logo_image = None

        self.configure_style()
        self.load_logo()
        self.show_login()

    def load_logo(self):
        logo_path = LOGO_PATH if LOGO_PATH.exists() else LEGACY_LOGO_PATH

        if not logo_path.exists():
            self.logo_image = None
            return

        try:
            self.logo_image = tk.PhotoImage(file=str(logo_path))
            self.iconphoto(False, self.logo_image)
        except Exception:
            self.logo_image = None

    def get_logo_image(self):
        return self.logo_image

    def configure_style(self):
        style = ttk.Style()

        try:
            style.theme_use("clam")
        except Exception:
            pass

        style.configure("TFrame", background=COLOR_BG)
        style.configure("App.TFrame", background=COLOR_BG)
        style.configure("Card.TFrame", background=COLOR_PANEL)
        style.configure("Panel.TFrame", background=COLOR_PANEL)
        style.configure("Sidebar.TFrame", background=COLOR_SIDEBAR)
        style.configure("TLabel", background=COLOR_BG, foreground=COLOR_TEXT, font=("Segoe UI", 10))
        style.configure("Card.TLabel", background=COLOR_PANEL, foreground=COLOR_TEXT, font=("Segoe UI", 10))
        style.configure("Muted.TLabel", background=COLOR_BG, foreground=COLOR_MUTED, font=("Segoe UI", 9))
        style.configure("Title.TLabel", background=COLOR_BG, foreground=COLOR_TEXT)
        style.configure("MetricValue.TLabel", background=COLOR_PANEL, foreground=COLOR_ACCENT, font=("Segoe UI", 18, "bold"))
        style.configure("MetricLabel.TLabel", background=COLOR_PANEL, foreground=COLOR_MUTED, font=("Segoe UI", 9))
        style.configure("Sidebar.TLabel", background=COLOR_PANEL, foreground=COLOR_TEXT, font=("Segoe UI", 10))
        style.configure("SidebarTitle.TLabel", background=COLOR_PANEL, foreground=COLOR_CYAN)
        style.configure("SidebarMuted.TLabel", background=COLOR_PANEL, foreground=COLOR_MUTED, font=("Segoe UI", 9))
        style.configure(
            "TButton",
            background=COLOR_BUTTON,
            foreground=COLOR_TEXT,
            bordercolor=COLOR_BORDER,
            lightcolor=COLOR_BUTTON,
            darkcolor=COLOR_BUTTON,
            focusthickness=0,
            focuscolor=COLOR_BUTTON,
            relief="flat",
            font=("Segoe UI", 10, "bold"),
            padding=(14, 8)
        )
        style.map(
            "TButton",
            background=[("pressed", COLOR_BUTTON_HOVER), ("active", COLOR_BUTTON_HOVER)],
            bordercolor=[("pressed", COLOR_ACCENT), ("active", COLOR_ACCENT)],
            foreground=[("pressed", "white"), ("active", "white")]
        )
        style.configure(
            "Primary.TButton",
            background=COLOR_PRIMARY,
            foreground="white",
            bordercolor=COLOR_PRIMARY,
            lightcolor=COLOR_PRIMARY,
            darkcolor=COLOR_PRIMARY,
            focusthickness=0,
            focuscolor=COLOR_PRIMARY,
            relief="flat",
            font=("Segoe UI", 10, "bold"),
            padding=(14, 8)
        )
        style.map(
            "Primary.TButton",
            background=[("pressed", COLOR_PRIMARY_HOVER), ("active", COLOR_PRIMARY_HOVER)],
            bordercolor=[("pressed", COLOR_PRIMARY_HOVER), ("active", COLOR_PRIMARY_HOVER)],
            foreground=[("pressed", "white"), ("active", "white")]
        )
        style.configure("Sidebar.TButton", background=COLOR_PANEL, foreground=COLOR_TEXT, borderwidth=0, padding=(10, 7), anchor="w")
        style.map("Sidebar.TButton", background=[("active", COLOR_ACCENT)], foreground=[("active", "white")])
        style.configure("TLabelframe", background=COLOR_BG, foreground=COLOR_TEXT, bordercolor=COLOR_BORDER, relief="solid")
        style.configure("TLabelframe.Label", background=COLOR_BG, foreground=COLOR_CYAN, font=("Segoe UI", 10, "bold"))
        style.configure("TEntry", fieldbackground="#F8FAFC", foreground="#111827", bordercolor=COLOR_BORDER, lightcolor=COLOR_BORDER, darkcolor=COLOR_BORDER)
        style.configure("TCombobox", fieldbackground="#F8FAFC", foreground="#111827", bordercolor=COLOR_BORDER, lightcolor=COLOR_BORDER, darkcolor=COLOR_BORDER)
        style.configure("TCheckbutton", background=COLOR_BG, foreground=COLOR_TEXT)
        style.configure("Treeview", font=("Segoe UI", 9), rowheight=26, background=COLOR_PANEL, foreground=COLOR_TEXT, fieldbackground=COLOR_PANEL, borderwidth=0)
        style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"), background=COLOR_SURFACE, foreground=COLOR_TEXT, relief="flat")

        style.map(
            "Treeview",
            background=[("selected", COLOR_ACCENT)],
            foreground=[("selected", "white")]
        )

    def clear_screen(self):
        if self.current_frame:
            self.current_frame.destroy()
            self.current_frame = None

    def show_login(self):
        self.clear_screen()
        self.current_frame = LoginFrame(self, self)
        self.current_frame.pack(fill="both", expand=True)

    def show_register(self):
        self.clear_screen()
        self.current_frame = RegisterFrame(self, self)
        self.current_frame.pack(fill="both", expand=True)

    def show_main(self, user):
        self.clear_screen()
        self.current_frame = ModernMainFrame(self, self, user)
        self.current_frame.pack(fill="both", expand=True)

if __name__ == "__main__":
    app = App()
    app.mainloop()


